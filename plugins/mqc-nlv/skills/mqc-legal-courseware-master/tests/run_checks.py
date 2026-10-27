#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""法律课件大师（V7）回归。改完 V7 跑这一个；不过不发。

    python3 tests/run_checks.py            全部
    python3 tests/run_checks.py --update   确认改动是有意的之后，重写成品账（tests/golden/slides.json）

查什么：
  结构校验      认不得的模块 / 外观 / 风格、重叠、出格、太小太大都拒；旧格式（deck@1）自动升级
  网格与连组    相邻同类卡片连成一组（间距收到 4px、内角收小），拆开、换外观、无底都不连
  落点与整理    放不下就不放；整理上推补缝、并排等高、补齐空白，整理两次结果不变
  每个组合      两种风格各拼一遍：颜色只用风格原色、元素不出画布、填好内容就没有问题
  文字          放不下就拒（不缩到 16px 以下）、金句必须写出处、空卡片拦下、转义
  大数字        由同页图表推得出就过，推不出就拒（V6 的数字核对）
  V1–V5 的图    等比缩放、太小就拒、风格不符就拒、位图不收
  导出          PPTX 能打开、页数对、没有位图、颜色在白名单里、字都在、篡改抓得到、同一份说明导两次一样
  编辑器        接口通、只收 .json/.svg、读素材不越界、外来 Host 拒；编辑器的 JS 规则与 Python 逐项一致（node）
  成品账        示例课件每页 SVG 的指纹不悄悄变（V6 升级让 V7 变了，这里会先红）
  兄弟模块      跑完这一整套，V1–V6 的文件一个字节都没变（V7 只读它们）
"""
import base64
import copy
import hashlib
import json
import os
import random
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
import zipfile
import xml.etree.ElementTree as ET

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(HERE, ".."))
sys.path.insert(0, os.path.join(ROOT, "scripts"))
SKILLS = os.path.normpath(os.path.join(ROOT, ".."))
DEMO = os.path.join(ROOT, "examples", "demo")
GOLDEN = os.path.join(HERE, "golden", "slides.json")
GRID_JS = os.path.join(ROOT, "web", "grid.js")
SIBLINGS = ["mqc-litigation-visual-redraw", "mqc-timeline-master", "mqc-legal-relation-master",
            "mqc-chronicle-master", "mqc-trial-confrontation-master", "mqc-legal-chart-master"]

for st in (sys.stdout, sys.stderr):
    try:
        st.reconfigure(encoding="utf-8")
    except Exception:
        pass


def tree_hash(top):
    """一个目录下所有源文件的指纹（略过 __pycache__ 与 .pyc：Python 自己写的缓存不算改动）。"""
    h = hashlib.sha256()
    if not os.path.isdir(top):
        return None
    for dp, dns, fns in os.walk(top):
        dns[:] = sorted(d for d in dns if d != "__pycache__")
        for fn in sorted(fns):
            if fn.endswith(".pyc"):
                continue
            p = os.path.join(dp, fn)
            h.update(os.path.relpath(p, top).encode("utf-8"))
            with open(p, "rb") as f:
                h.update(hashlib.sha256(f.read()).digest())
    return h.hexdigest()


BEFORE = {m: tree_hash(os.path.join(SKILLS, m)) for m in SIBLINGS}
_BRAND_HOME = tempfile.mkdtemp(prefix="nlv7-brand-")
os.environ["NLV_BRAND_DIR"] = _BRAND_HOME           # 测试绝不碰用户本机的公共名片
DEMO_HASH = tree_hash(DEMO)

import paths  # noqa: E402,F401
import deck as D  # noqa: E402
import export as E  # noqa: E402
import grid as G  # noqa: E402
import svgimport  # noqa: E402
import studio  # noqa: E402

RESULTS = []


def check(name, cond, why=""):
    RESULTS.append((name, bool(cond), why))


def raises(fn, exc, contains=""):
    try:
        fn()
    except exc as e:
        return contains in str(e)
    return False


def fp(s):
    return base64.b32encode(hashlib.sha256(s.encode("utf-8")).digest()[:10]).decode().lower()


def demo():
    return D.load(os.path.join(DEMO, "deck.json"))


def page(d, sid):
    """示例课件里按 id 找页（不按位置：示例课件加页时测试不跟着错位）。"""
    return next(s for s in d["slides"] if s["id"] == sid)


def pno(d, sid):
    return [s["id"] for s in d["slides"]].index(sid)


def one(mods, style="奇川风"):
    return {"schema": D.SCHEMA, "title": "t", "style": style, "slides": [{"id": "a", "modules": mods}]}


def M(t, x, y, w, h, i, **kw):
    return dict(id=i, type=t, x=x, y=y, w=w, h=h, **kw)


def colors_in(svg):
    return {c.upper() for c in re.findall(r"#([0-9A-Fa-f]{6})\b", svg)}


def outside(svg):
    """粗查：rect 的 x/y/width/height 都在画布里（容 1px）。"""
    bad = []
    for m in re.finditer(r"<rect\b[^>]*>", svg):
        a = dict(re.findall(r'([\w-]+)="([^"]*)"', m.group(0)))
        try:
            x, y = float(a.get("x", 0)), float(a.get("y", 0))
            w, h = float(a.get("width", 0)), float(a.get("height", 0))
        except ValueError:
            continue
        if x < -1 or y < -1 or x + w > D.W + 1 or y + h > D.H + 1:
            bad.append(m.group(0)[:80])
    return bad


FILL = {
    "conclusion": {"text": "实际施工人只能在欠付范围内向发包人主张"},
    "title": {"text": "名为买卖实为借贷"},
    "text": {"text": "当事人之间并无真实的货物交付。"},
    "bullets": {"items": ["看交付", "看回购"]},
    "stat": {"value": "41", "label": "份全额支持"},
    "step": {"heading": "看交付", "text": "审查是否存在真实的货物交付。"},
    "quote": {"text": "双方的真实意思为借贷。", "source": "示例判决（虚构）"},
    "chart": {"file": "assets/bar-filings.json"},
    "figure": {"file": "assets/small-figure.svg", "from": "V3"},
    "note": {"text": "来源：示例数据"},
    "statute": {"law": "中华人民共和国民法典", "article": "第六百八十条第一款", "text": "禁止高利放贷，借款的利率不得违反国家有关规定。"},
    "holding": {"text": "双方的真实意思为借贷。", "court": "某市中级人民法院", "case_no": "示例案号 A-001"},
    "case": {"case_no": "示例案号 A-001", "court": "某市中级人民法院", "cause": "买卖合同纠纷", "result": "改判", "point": "无真实交付。"},
    "issue": {"heading": "是否存在真实交付", "text": "双方各执一词。"},
    "risk": {"heading": "利息过高部分不受保护", "text": "先核算。"},
    "evidence": {"items": ["银行转账凭证——证明款项已交付", "回购协议——证明到期回购"]},
    "logo": {},
    "profile": {"who": "firm"},
    "contact": {},
}


# ------------------------------------------------------------------ 结构校验
def t_validate():
    d = demo()
    check("示例课件通过校验", D.validate(d) is None)
    for label, mutate, frag in [
        ("schema 不对", lambda x: x.update(schema="x"), "schema"),
        ("风格不认识", lambda x: x.update(style="绿色风"), "style"),
        ("没有页", lambda x: x.update(slides=[]), "至少"),
        ("模块类型不认识", lambda x: x["slides"][0]["modules"][0].update(type="table"), "不认识"),
        ("坐标不是整数", lambda x: x["slides"][0]["modules"][0].update(x=1.5), "整数"),
        ("出了页面", lambda x: x["slides"][0]["modules"][0].update(x=11), "出了页面"),
        ("太小", lambda x: page(x, "s3")["modules"][1].update(w=2), "太小"),
        ("标题太高", lambda x: page(x, "s3")["modules"][0].update(h=4, w=12), "最多"),
        ("外观不认识", lambda x: page(x, "s3")["modules"][3].update(tone="glow"), "外观"),
        ("图表不做重音卡", lambda x: page(x, "s3")["modules"][1].update(tone="accent"), "重音卡"),
        ("出血只给整页图", lambda x: page(x, "s3")["modules"][1].update(bleed=True), "出血"),
        ("模块重叠", lambda x: page(x, "s3")["modules"][3].update(y=2), "重叠"),
        ("模块 id 重复", lambda x: page(x, "s3")["modules"][3].update(id="m2"), "重复"),
        ("页 id 重复", lambda x: x["slides"][1].update(id="s1"), "重复"),
    ]:
        x = copy.deepcopy(d)
        mutate(x)
        check(f"校验拒：{label}", raises(lambda: D.validate(x), D.DeckError, frag))
    nd = D.new_deck("新课件")
    check("新建课件可用：标题照写，第一页是空页（编辑器请律师先选场景）", D.validate(nd) is None and nd["title"] == "新课件"
          and nd["slides"][0]["modules"] == [])
    _, rnd = D.compose(nd, 0, DEMO, {})
    check("空页导出前被拦下，并告诉律师去「图表库」挑图或「诉讼图」点场景", any("还是空的" in p_ and "图表库" in p_ and "诉讼图" in p_ for p_ in rnd["problems"]), rnd["problems"])
    # 旧格式自动升级
    v1 = {"schema": "mqc-nlv/deck@1", "title": "旧", "style": "奇川风", "slides": [
        {"id": "s1", "layout": "cover", "slots": {"title": {"type": "text", "text": "旧课件"},
                                                 "subtitle": {"type": "text", "text": "副题"}}},
        {"id": "s2", "layout": "chart-notes", "slots": {"main": {"type": "chart", "file": "assets/bar-filings.json"},
                                                       "side_title": {"type": "text", "text": "怎么读"},
                                                       "side": {"type": "bullets", "items": ["一", "二"]}}},
        {"id": "s3", "layout": "quote", "slots": {"quote": {"type": "quote", "text": "原文", "source": "某判决"}}}]}
    td = tempfile.mkdtemp()
    try:
        with open(os.path.join(td, "deck.json"), "w", encoding="utf-8") as f:
            json.dump(v1, f, ensure_ascii=False)
        up = D.load(os.path.join(td, "deck.json"))
        check("旧格式 deck@1 自动升级成网格", up["schema"] == D.SCHEMA and all("modules" in s for s in up["slides"]))
        types = [sorted(m["type"] for m in s["modules"]) for s in up["slides"]]
        check("升级后内容都在（标题、图表、要点、金句）", ["note", "title"] == types[0] or "title" in types[0]
              and "chart" in types[1] and "bullets" in types[1] and types[2] == ["quote"], types)
        b = next(m for m in up["slides"][1]["modules"] if m["type"] == "bullets")
        check("升级：侧栏小标题并入要点卡", b.get("heading") == "怎么读" and b.get("items") == ["一", "二"], b)
    finally:
        shutil.rmtree(td, ignore_errors=True)


# ------------------------------------------------------------------ 网格与连组
def t_grid():
    check("网格常数自洽（12 栏 × 6 行刚好铺满 1280×720）",
          G.MX * 2 + G.COLS * G.COL + (G.COLS - 1) * G.GUT == G.W and G.MY * 2 + G.ROWS * G.ROW + (G.ROWS - 1) * G.GUT == G.H)
    row = [M("stat", 3 * i, 1, 3, 2, f"k{i}") for i in range(4)]
    sides, groups = G.fusion(row)
    check("一排四张大数字连成一组", len(groups) == 1 and len(groups[0]) == 4)
    check("连组的边：首 r、中 l+r、尾 l", sides["k0"] == {"r"} and sides["k1"] == {"l", "r"} and sides["k3"] == {"l"}, sides)
    r0 = G.card_rect(row[0], sides["k0"])
    r1 = G.card_rect(row[1], sides["k1"])
    check("连组后组内间距收到 4px", abs(r1[0] - (r0[0] + r0[2]) - G.FUSE_GAP) < 1e-9, (r0, r1))
    check("连组后外缘仍对齐网格", r0[0] == G.px(row[0])[0] and r0[1] == G.px(row[0])[1])
    check("内角收小、外角不变", G.corners(sides["k0"], 12, 4) == (12, 4, 4, 12) and G.corners(sides["k1"], 12, 4) == (4, 4, 4, 4))
    stack = [M("bullets", 0, 0, 4, 2, "a"), M("text", 0, 2, 4, 3, "b")]
    check("上下等宽的正文与要点连成一组", G.fusion(stack)[0]["a"] == {"b"} and G.fusion(stack)[0]["b"] == {"t"})
    for label, mut in [("join:false 拆开", {"join": False}), ("外观不同不连", {"tone": "accent"}), ("无底不连", {"tone": "bare"})]:
        x = copy.deepcopy(row)
        x[1].update(mut)
        check(f"连组：{label}", not G.fusion(x)[0]["k1"], G.fusion(x)[0]["k1"])
    mixed = [M("stat", 0, 1, 3, 2, "s"), M("step", 3, 1, 3, 2, "t")]
    check("不同家族（大数字 / 编号卡）不连", not any(G.fusion(mixed)[0].values()))
    uneven = [M("stat", 0, 1, 3, 2, "s"), M("stat", 3, 1, 3, 3, "t")]
    check("并排但不等高不连", not any(G.fusion(uneven)[0].values()))
    # 落点
    full = [M("chart", 0, 0, 12, 6, "c")]
    check("页满了：放不下就不放", G.place(full, "stat") is None)
    p = G.place([M("title", 0, 0, 12, 1, "t")], "chart")
    check("放图表：默认 8×5 找到空位", p == (0, 1, 8, 5), p)
    p = G.place([M("title", 0, 0, 12, 1, "t"), M("chart", 0, 1, 8, 5, "c")], "chart")
    check("空位不够：缩到最小也放不下就不放（图表最小 6×4）", p is None, p)
    p = G.place([M("title", 0, 0, 12, 1, "t"), M("chart", 0, 1, 8, 5, "c")], "bullets")
    check("要点卡缩到放得下的最大尺寸", p == (8, 1, 4, 3), p)
    # 整理
    messy = [M("title", 0, 0, 12, 1, "t", text="x"), M("stat", 0, 2, 3, 2, "a"), M("stat", 3, 2, 3, 3, "b"),
             M("text", 6, 4, 6, 2, "c")]
    bad = G.check_slide(messy)
    check("整理前的样例本身合法", not bad, bad)
    out, notes = G.tidy(messy)
    check("整理：上推补缝", next(m for m in out if m["id"] == "a")["y"] == 1, out)
    a, b = (next(m for m in out if m["id"] == k) for k in ("a", "b"))
    check("整理：并排同类卡片拉到等高，于是连成一组", a["h"] == b["h"] and G.fusion(out)[0]["a"] == {"r"}, (a, b))
    check("整理：结果合法", not G.check_slide(out))
    out2, notes2 = G.tidy(out)
    check("整理两次结果不变（幂等）", out2 == out and not notes2, notes2)
    check("整理只动位置大小，不动内容", all({k: v for k, v in o.items() if k not in "xywh"} ==
                                          {k: v for k, v in m.items() if k not in "xywh"} for o, m in zip(out, messy)))
    # 组合
    for name, t in G.TEMPLATES.items():
        mods = [dict(m, id=f"m{i}") for i, m in enumerate(t["mods"])]
        check(f"组合「{t['zh']}」本身合法", not G.check_slide(mods), G.check_slide(mods))
    src = [M("title", 0, 0, 12, 1, "t", text="题"), M("chart", 0, 1, 8, 5, "c", file="assets/bar-filings.json"),
           M("bullets", 8, 1, 4, 5, "b", items=["一"]), M("quote", 0, 0, 4, 2, "q")][:3]
    out, left = G.apply_template(src, "two-charts")
    check("换组合：内容按类型搬进新格子，搬不下的放暂存", any(m.get("file") for m in out) and [m["id"] for m in left] == ["b"],
          (out, left))
    check("换组合：保留原 id（编辑器选中不丢）", {"t", "c"} <= {m["id"] for m in out})


# ------------------------------------------------------------------ 每个组合 × 两种风格
SMALL_FIG = ('<svg xmlns="http://www.w3.org/2000/svg" width="600" height="360" viewBox="0 0 600 360">'
             '<rect width="600" height="360" fill="#FFFFFF"/><rect x="40" y="60" width="160" height="70" rx="8" fill="#F3F4F6" stroke="#991B1B"/>'
             '<text x="60" y="102" font-size="20" fill="#1F2933">甲方</text><path d="M200 95 L400 95" stroke="#6B7280"/>'
             '<rect x="400" y="60" width="160" height="70" rx="8" fill="#F3F4F6" stroke="#6B7280"/>'
             '<text x="420" y="102" font-size="20" fill="#1F2933">乙方</text></svg>')


def t_templates(td):
    base = os.path.join(td, "tpl")
    shutil.copytree(DEMO, base)
    with open(os.path.join(base, "assets", "small-figure.svg"), "w", encoding="utf-8") as f:
        f.write(SMALL_FIG)
    for style in D.STYLES:
        allowed = {c.lstrip("#").upper() for c in D.allowed_colors({"style": style})}
        for name, t in G.TEMPLATES.items():
            mods = []
            for i, m in enumerate(t["mods"]):
                mm = dict(m, id=f"m{i}", **FILL[m["type"]])
                if m["type"] == "figure" and style != "奇川风":
                    mm["file"] = None
                mods.append(mm)
            if any(m["type"] == "figure" and not m.get("file") for m in mods):
                continue                                  # 示例图是奇川风，歸藏风下本来就该拒（另有专项）
            dk = one(mods, style)
            D.validate(dk)
            svg, rep = D.compose(dk, 0, base, {})
            extra = colors_in(svg) - allowed - (colors_in(SMALL_FIG) if any(m["type"] == "figure" for m in mods) else set())
            check(f"{style}·{t['zh']}：填好内容就没有问题", not rep["problems"], rep["problems"][:2])
            check(f"{style}·{t['zh']}：只用风格原色", not extra, "、".join(sorted(extra)))
            check(f"{style}·{t['zh']}：不出画布", not outside(svg), outside(svg)[:2])
    svg, rep = D.compose(one([]), 0, DEMO, {})
    check("空页能拼出预览，唯一的问题是「还是空的」", "<svg" in svg and len(rep["problems"]) == 1 and "还是空的" in rep["problems"][0],
          rep["problems"])


# ------------------------------------------------------------------ 文字
def t_text():
    x = one([M("title", 0, 0, 12, 1, "t", text="标题"), M("bullets", 0, 1, 12, 5, "b", items=["字" * 80] * 30)])
    _, rep = D.compose(x, 0, DEMO, {})
    check("要点太多：拒，不缩到 16px 以下", rep["modules"]["b"]["problems"], rep["modules"]["b"])
    x["slides"][0]["modules"][1]["items"] = ["一句" * 25] * 7
    svg, rep = D.compose(x, 0, DEMO, {})
    sizes = [float(s) for s in re.findall(r'font-size="([\d.]+)"[^>]*data-mod="a\.b"', svg)]
    check("要点略多：收字号，仍不小于 16px", not rep["modules"]["b"]["problems"] and sizes and min(sizes) >= D.MIN_BODY - 1e-6,
          (rep["modules"]["b"], sizes[:3]))
    check("要点超过 6 条：提醒拆页", any("6 条" in w for w in rep["modules"]["b"]["warnings"]))
    q = one([M("quote", 1, 1, 10, 4, "q", text="一句话")])
    _, rep = D.compose(q, 0, DEMO, {})
    check("金句没写出处：拒", any("出处" in p for p in rep["problems"]), rep["problems"])
    q["slides"][0]["modules"][0]["source"] = "某判决"
    _, rep = D.compose(q, 0, DEMO, {})
    check("金句写了出处：过，并提示逐字核对", not rep["problems"] and any("逐字" in w for w in rep["warnings"]))
    for t in ("title", "text", "bullets", "stat", "step", "chart", "figure", "note", "statute", "holding", "case", "issue",
              "risk", "evidence"):
        mn = list(G.TYPES[t]["min"])
        mods = [M(t, 0, 0, max(mn[0], 6), mn[1], "e")]
        _, rep = D.compose(one(mods), 0, DEMO, {})
        check(f"空的{G.TYPES[t]['zh']}卡：拦下，不导出空卡片", rep["problems"], rep["modules"])
    x = one([M("title", 0, 0, 12, 1, "t", text='<script>&"x"')])
    svg, _ = D.compose(x, 0, DEMO, {})
    check("文字转义（< & 不破坏 SVG）", "<script>" not in svg and "&lt;script&gt;" in svg)
    x = one([M("title", 0, 0, 12, 1, "t", text="题"), M("step", 0, 1, 4, 5, "b", heading="二"),
             M("step", 4, 1, 4, 5, "a", heading="一")])
    x["slides"][0]["modules"][1]["x"], x["slides"][0]["modules"][2]["x"] = 4, 0
    svg, _ = D.compose(x, 0, DEMO, {})
    na = re.search(r'data-mod="a\.a\.no">(\d+)<', svg).group(1)
    check("编号卡按阅读顺序自动编号（左边的是 01）", na == "01", na)


# ------------------------------------------------------------------ 大数字
def t_stat():
    d = demo()
    s = copy.deepcopy(d)
    m = next(x for x in page(s, "s3")["modules"] if x["type"] == "stat")
    i3 = pno(s, "s3")
    m["from"] = "m2"
    _, rep = D.compose(s, i3, DEMO, {})
    check("大数字由同页图表推得出：过", not rep["modules"][m["id"]]["problems"] and rep["modules"][m["id"]]["notes"])
    m["value"] = "+66%"
    _, rep = D.compose(s, i3, DEMO, {})
    check("大数字推不出：拒（数只许算、不许编）", any("推不出" in p for p in rep["modules"][m["id"]]["problems"]),
          rep["modules"][m["id"]])
    m["from"] = "m9"
    _, rep = D.compose(s, i3, DEMO, {})
    check("大数字 from 指向不存在的图表：拒", rep["modules"][m["id"]]["problems"])
    x = one([M("stat", 0, 0, 3, 2, "k", value="99", label="件")])
    _, rep = D.compose(x, 0, DEMO, {})
    check("大数字没有出处图表：提醒人核", any("手填" in w for w in rep["warnings"]), rep["warnings"])


# ------------------------------------------------------------------ V1–V5 的图
def _svg(w, h, body, color="#991B1B"):
    return (f'<svg xmlns="http://www.w3.org/2000/svg" width="{w}" height="{h}" viewBox="0 0 {w} {h}">'
            f'<rect width="{w}" height="{h}" fill="#FFFFFF"/>{body.replace("COL", color)}</svg>')


def t_figures():
    fig = open(os.path.join(DEMO, "assets", "v3-relation.svg"), encoding="utf-8").read()
    check("认出 V3 示例图是奇川风", svgimport.style_of(fig) == "奇川风")
    els, rep = svgimport.fit(fig, (0, 0, 1280, 720), "奇川风")
    check("V3 图整页放得下", not rep["problems"] and rep["scale"] > 0, rep["problems"])
    n_dash = len(re.findall(r"stroke-dasharray=['\"]", fig))
    check("V3 的虚线（单引号属性）导进来还是虚线：法律上虚线、实线意思不同",
          n_dash > 0 and els.count("stroke-dasharray") == n_dash, (n_dash, els.count("stroke-dasharray")))
    check("属性单引号、双引号都认", svgimport._attrs("""a='1' b="2" c='x"y'""") == {"a": "1", "b": "2", "c": 'x"y'})
    an = rep["anchors"]
    nodes = [x for x in an if x["role"] == "node"]
    rels = [x for x in an if x["role"] == "edge"]
    check("磁吸点：认出 V3 的 9 个主体 / 客体，名称对得上", len(nodes) == 9 and
          {"发包人", "涉案工程", "工程价款"} <= {x["label"].split(" ")[0] for x in nodes}, [x["label"] for x in nodes])
    check("磁吸点：V3 的 11 条关系都认出来，并配上各自的关系标注", len(rels) == 11 and all(x["label"] for x in rels)
          and {"欠付范围内担责", "本案诉请 折价补偿", "承建标的"} <= {x["label"] for x in rels}, [x["label"] for x in rels])
    v1r = ('<svg viewBox="0 0 800 300"><text x="400" y="40" font-size="30">某某关系图</text><g data-role="edges">'
           '<path d="M 100,150 L 300,150" stroke="#4B5563"/><path d="M 100,120 L 100,60 L 700,60 L 700,120" stroke="#991B1B"/></g>'
           '<g data-role="nodes"><g data-role="node" data-id="jia"><rect x="40" y="130" width="60" height="40"/><text x="70" y="155" font-size="16">甲</text></g></g>'
           '<g data-role="edge-labels"><text x="200" y="140" font-size="13">借款</text><text x="400" y="95" font-size="13">连带清偿责任</text></g></svg>')
    _, rv = svgimport.fit(v1r, (0, 0, 800, 300))
    check("磁吸点：V1 edge-labels 组里的标注配给最近的关系线；大字图名不当标注",
          [x["label"] for x in rv["anchors"] if x["role"] == "edge"] == ["借款", "连带清偿责任"], rv["anchors"])
    check("磁吸点：每个都落在格子里、有大小", all(0 <= x["box"][0] and x["box"][0] + x["box"][2] <= 1280 and x["box"][2] > 0 for x in an))
    check("磁吸点：导出的元素带上所属元素的标记", 'data-anchor="node:#1"' in els)
    v1 = ('<svg viewBox="0 0 1000 600"><rect width="1000" height="600" fill="#FFFFFF"/>'
          '<text data-role="title" x="56" y="60" font-size="30">原图自带的大标题</text>'
          '<g data-role="event" data-id="e1"><circle cx="200" cy="300" r="6" fill="#991B1B"/><text x="200" y="330" font-size="12" text-anchor="middle">2015.08.01 签约</text></g>'
          '<g data-role="event" data-id="e2"><circle cx="800" cy="300" r="6"/><text x="800" y="330" font-size="12" text-anchor="middle">2016.05.26 催告</text></g>'
          '<line x1="150" y1="300" x2="850" y2="300" stroke="#333333"/></svg>')
    _, r1 = svgimport.fit(v1, (0, 0, 800, 480))
    e2, r2 = svgimport.fit(v1, (0, 0, 800, 480), trim=True, drop_title=True)
    check("磁吸点：V1 / V2 的 data-id 认作事件", [(x["key"], x["role"]) for x in r1["anchors"]] == [("id:e1", "event"), ("id:e2", "event")], r1["anchors"])
    check("裁边：按内容裁掉原图页边，同一格子里放得更大", r2["scale"] > r1["scale"] * 1.15, (r1["scale"], r2["scale"]))
    check("裁边：字跟着变大", r2["min_font"] > r1["min_font"], (r1["min_font"], r2["min_font"]))
    check("略去原图标题：页面有自己的标题时不重复", "原图自带的大标题" not in e2 and any("标题" in n for n in r2["notes"]))
    check("裁边：图里的内容一个不少", "2015.08.01 签约" in e2 and "2016.05.26 催告" in e2)
    dk = demo()
    ci = next(i for i, s_ in enumerate(dk["slides"]) if any(m["type"] == "chart" for m in s_["modules"]))
    _, crep = D.compose(dk, ci, DEMO, {})
    can = [a_ for m in crep["modules"].values() for a_ in (m.get("anchors") or [])]
    check("磁吸点：V6 图表的每个数据项都是磁吸点，名字带数值和类别", can and all(a_["key"].startswith("id:") for a_ in can)
          and all(" " in a_["label"] for a_ in can if a_["role"] == "bar"), [a_["label"] for a_ in can][:4])
    check("字号缩放后不低于下限", rep["min_font"] is None or rep["min_font"] >= svgimport.MIN_FONT - 1e-6, rep["min_font"])
    ds = re.findall(r' d="([^"]*)"', els)
    check("曲线拍平成折线（PPTX 里可编辑）", ds and not any(re.search(r"[AaCcQqSsTtHhVv]", x) for x in ds), ds[:1])
    _, rep = svgimport.fit(fig, (0, 0, 1280, 720), "歸藏风")
    check("风格不符：拒，不给别的模块改颜色", any("奇川风" in p for p in rep["problems"]))
    tiny = _svg(2400, 1600, '<text x="10" y="40" font-size="14" fill="COL">小字</text>')
    _, rep = svgimport.fit(tiny, (32, 60, 608, 600), "奇川风")
    check("缩得太小、字看不清：拒", rep["problems"], rep)
    # 竖版的 V3 关系图放进横的图卡：拦下时指一条能走通的路（回 V3 加 --shape wide 出横版）；别的图不乱指
    _rel = _svg(630, 1298, '<g data-role="node" data-ref="A"><rect x="10" y="10" width="140" height="48" fill="COL"/>'
                           '<text x="20" y="40" font-size="14" fill="COL">甲公司</text></g>')
    _, rep = svgimport.fit(_rel, (0, 0, 1100, 470), "奇川风")
    check("竖版关系图缩得太小：提示回 V3 加 --shape wide 出横版", any("--shape wide" in p for p in rep["problems"]), rep["problems"])
    _, rep = svgimport.fit(tiny, (32, 60, 608, 600), "奇川风")
    check("不是关系图的小字图：不提 --shape（只给切片、拉大）", rep["problems"] and not any("--shape" in p for p in rep["problems"]), rep["problems"])
    _, rep = svgimport.fit('<svg xmlns="http://www.w3.org/2000/svg"><rect width="5" height="5"/></svg>', (0, 0, 100, 100), "奇川风")
    check("没有宽高：拒", rep["problems"])
    img = _svg(800, 600, '<image href="a.png" x="0" y="0" width="100" height="100"/>'
                         '<text x="10" y="40" font-size="24" fill="COL">字</text>')
    els, rep = svgimport.fit(img, (0, 0, 1280, 720), "奇川风")
    check("位图不收（记在 dropped）", "<image" not in els and rep["dropped"], rep["dropped"])
    arcs = svgimport.path_to_polys("M10 10 a20 20 0 0 1 40 0 q10 10 20 0 t20 0 s10 10 20 0 h10 v10 z")
    check("相对弧线 / 二次 / 简写曲线都能拍平", arcs and len(arcs[0]) >= 2)
    check("认出歸藏风", svgimport.style_of(_svg(800, 600, '<rect width="9" height="9" fill="#002FA7"/>')) == "歸藏风")


# ------------------------------------------------------------------ 法律卡
def t_legal():
    def comp(mods, style="奇川风"):
        return D.compose(one(mods, style), 0, DEMO, {})
    for t, (w, h) in (("statute", (6, 3)), ("holding", (8, 2)), ("case", (4, 5)), ("issue", (4, 3)), ("risk", (4, 2)),
                      ("evidence", (7, 4))):
        for style in D.STYLES:
            svg, rep = comp([M(t, 0, 0, w, h, "a", **FILL[t])], style)
            allowed = {c.lstrip("#").upper() for c in D.allowed_colors({"style": style})}
            check(f"{style}·{G.TYPES[t]['zh']}卡：写好就过、只用风格原色",
                  not rep["problems"] and not (colors_in(svg) - allowed), (rep["problems"], colors_in(svg) - allowed))
    _, rep = comp([M("statute", 0, 0, 6, 3, "a", text="禁止高利放贷。")])
    check("法条卡没写法律名称：拒（照录官方文本，不凭记忆写）", any("法律名称" in p for p in rep["problems"]))
    _, rep = comp([M("statute", 0, 0, 6, 3, "a", **FILL["statute"])])
    check("法条卡：提醒与官方文本逐字核对", any("逐字核对" in w and "民法典" in w for w in rep["warnings"]), rep["warnings"])
    _, rep = comp([M("holding", 0, 0, 8, 2, "a", text="要旨")])
    check("裁判要旨没写出处（法院 / 案号）：拒", any("出处" in p for p in rep["problems"]))
    _, rep = comp([M("case", 0, 0, 4, 5, "a", court="某法院", result="驳回")])
    check("案例卡没写案号：拒（不编）", any("案号" in p for p in rep["problems"]))
    svg, _ = comp([M("issue", 4, 0, 4, 3, "b", heading="二"), M("issue", 0, 0, 4, 3, "a", heading="一")])
    ka = re.search(r'data-mod="a\.a\.kick">([^<]+)<', svg).group(1)
    kb = re.search(r'data-mod="a\.b\.kick">([^<]+)<', svg).group(1)
    check("争议焦点按阅读顺序编号（焦点一、焦点二）", (ka, kb) == ("争议焦点一", "争议焦点二"), (ka, kb))
    sides, _ = G.fusion([M("statute", 0, 0, 6, 2, "a"), M("holding", 0, 2, 6, 3, "b")])
    check("法条与裁判要旨上下等宽：连成一组", sides["a"] == {"b"} and sides["b"] == {"t"})
    _, rep = comp([M("evidence", 0, 0, 5, 2, "a", items=["证据" * 30 + "——" + "目的" * 30] * 6)])
    check("证据太多放不下：拒", any("放不下" in p for p in rep["problems"]))
    for t in ("statute", "holding", "case", "evidence", "logo", "profile", "contact"):
        check(f"{G.TYPES[t]['zh']}不做重音卡", t in G.NO_ACCENT)
    for name in ("statute-read", "case-review", "issues", "trial-prep", "evidence", "timeline", "similar-cases",
                 "firm-intro", "lawyer-intro", "closing", "cover-brand"):
        check(f"组合「{G.TEMPLATES[name]['zh']}」还在（功能不砍）", G.TEMPLATES[name].get("group") in ("排法", "整套课件"))
    scenes = {k: t_ for k, t_ in G.TEMPLATES.items() if t_.get("group") == "场景"}
    check("选图：6 个法律场景，各带一张来自对应模块的图", len(scenes) == 6 and all(
        any(m["type"] in ("chart", "figure") for m in t_["mods"]) for t_ in scenes.values()), list(scenes))
    check("排法：都以图为中心（每个都有图）", all(any(m["type"] in ("chart", "figure") for m in t_["mods"])
                                       for t_ in G.TEMPLATES.values() if t_.get("group") == "排法"))
    check("整套课件的组合（封面、律所介绍、结束页……）都保留在「更多」", {"cover", "section", "closing", "firm-intro"} <=
          {k for k, t_ in G.TEMPLATES.items() if t_.get("group") == "整套课件"})
    check("一句话结论卡：不做重音卡、最多两行高", "conclusion" in G.NO_ACCENT and G.TYPES["conclusion"]["maxh"] == 2)


# ------------------------------------------------------------------ 律所名片
def _png(w, h, rgba=b"\x99\x1b\x1b\xff"):
    import struct
    import zlib
    raw = b"".join(b"\x00" + rgba * w for _ in range(h))
    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xffffffff)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 6, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def t_brand(td):
    import brand as B
    b = B.load(DEMO)
    check("课件目录下的 brand/ 优先（示例律所）", b["_dir"] == os.path.join(DEMO, "brand") and "某某" in b["firm"]["name"])
    empty_dir = os.path.join(td, "nobrand")
    os.makedirs(empty_dir)
    b0 = B.load(empty_dir)
    check("没有课件名片就用本机公共名片（测试里指向临时目录）", b0["_dir"] == _BRAND_HOME and not b0["firm"])
    check("图片尺寸：PNG", B.image_size(_png(40, 12), ".png") == (40, 12))
    check("图片尺寸：读不出就是 None", B.image_size(b"nope", ".png") is None)
    for name, data, frag in [("logo.gif", b"GIF89a", "只收"), ("logo.png", b"x" * 10, "尺寸"),
                             ("logo.svg", b"<html></html>", "SVG"), ("logo.png", b"\x89PNG" + b"0" * (2 * 1024 * 1024), "太大")]:
        check(f"Logo 上传拒：{frag}", raises(lambda: B.save_logo(empty_dir, name, data), B.BrandError, frag))
    check("名片校验：律师要有姓名", raises(lambda: B.validate({"lawyers": [{"id": "a"}]}), B.BrandError, "姓名"))
    check("名片校验：律师 id 不重复", raises(lambda: B.validate({"lawyers": [{"id": "a", "name": "甲"}, {"id": "a", "name": "乙"}]}),
                                      B.BrandError, "重复"))
    proj = os.path.join(td, "brandproj")
    shutil.copytree(DEMO, proj)
    bb = B.load(proj)
    mods, note = B.intro_slide(bb, "firm")
    check("一键生成律所介绍页", {m["type"] for m in mods} >= {"logo", "title", "profile", "contact"} and "某某" in note)
    mods, _ = B.intro_slide(bb, "lawyer", "zhang")
    check("一键生成律师简介页（姓名、职务、代表业绩）", any(m["type"] == "title" and m["text"] == "张某某" and m.get("sub") for m in mods)
          and any(m["type"] == "bullets" and m.get("items") for m in mods))
    check("没有这位律师：拒", raises(lambda: B.intro_slide(bb, "lawyer", "nobody"), B.BrandError))
    mods, _ = B.intro_slide(bb, "firm")
    dk = one([dict(m, id=f"m{i}") for i, m in enumerate(mods)])
    svg1, rep = D.compose(dk, 0, proj, {})
    check("介绍页：名片齐全就能拼", not rep["problems"], rep["problems"])
    bb["firm"]["intro"] = "改过的介绍：专注争议解决。"
    B.save(proj, bb)
    svg2, _ = D.compose(dk, 0, proj, {})
    check("介绍卡实时读名片：改名片，介绍页跟着变", "改过的介绍" in svg2 and "改过的介绍" not in svg1)
    _, rep = D.compose(one([M("logo", 0, 0, 2, 1, "l"), M("contact", 0, 1, 4, 2, "c")]), 0, empty_dir, {})
    check("没有 Logo、没有联系方式：拦下并说去哪里填", any("Logo" in p for p in rep["problems"])
          and any("联系方式" in p or "地址" in p for p in rep["problems"]), rep["problems"])
    bb["firm"].update(address="地" * 30, phone="0" * 20, email="e" * 30 + "@example.com", website="w" * 30)
    B.save(proj, bb)
    _, rep = D.compose(one([M("contact", 0, 0, 4, 1, "c")]), 0, proj, {})
    check("联系方式一行放不下：拒，提示拉高", any("一行放不下" in p for p in rep["problems"]), rep["problems"])


def t_logo_export(td):
    import brand as B
    d = demo()
    out = os.path.join(td, "logo.pptx")
    p, _ = E.build(d, DEMO, out, cache={})
    with zipfile.ZipFile(p) as z:
        media = [n for n in z.namelist() if n.startswith("ppt/media/")]
        n_logo_pages = 0
        for i, s in enumerate(d["slides"]):
            x = z.read(f"ppt/slides/slide{i + 1}.xml").decode("utf-8")
            n_logo_pages += x.count("<p:pic>")
        check("PNG Logo 作为图片放进 PPTX，且只有 Logo", media and all("/" + E.LOGO_MEDIA in n for n in media), media[:3])
        check("页脚 Logo：内容页都有，有 Logo 卡的页不重复", n_logo_pages == len(media))
        ct = z.read("[Content_Types].xml").decode("utf-8")
        check("PNG 已登记内容类型", 'Extension="png"' in ct)
    try:
        from pptx import Presentation
        pics = [sh for s in Presentation(p).slides for sh in s.shapes if sh.shape_type == 13]
        check("python-pptx 认得出 Logo 图片", pics and all(sh.name.startswith("Logo") for sh in pics))
    except ImportError:
        check("python-pptx 认得出 Logo 图片（未装，跳过）", True)
    tam = os.path.join(td, "pic.pptx")
    with zipfile.ZipFile(p) as zi, zipfile.ZipFile(tam, "w", zipfile.ZIP_DEFLATED) as zo:
        for n in zi.namelist():
            data = zi.read(n)
            if n == "ppt/slides/slide3.xml":
                data = data.decode("utf-8").replace("</p:spTree>", '<p:pic><p:nvPicPr><p:cNvPr id="77" name="截图"/><p:cNvPicPr/><p:nvPr/>'
                                                    '</p:nvPicPr><p:blipFill><a:blip r:embed="rIdX"/></p:blipFill><p:spPr/></p:pic></p:spTree>').encode("utf-8")
            zo.writestr(n, data)
    check("审计抓到 Logo 以外的图片（比如截图）", any("位图" in b for b in E.audit(tam, d, DEMO)))
    # SVG Logo：转成形状，颜色算律所品牌色
    proj = os.path.join(td, "svglogo")
    shutil.copytree(DEMO, proj)
    with open(os.path.join(proj, "brand", "logo.svg"), "w", encoding="utf-8") as f:
        f.write('<svg xmlns="http://www.w3.org/2000/svg" width="200" height="60" viewBox="0 0 200 60">'
                '<rect x="0" y="0" width="60" height="60" fill="#0A7B4F"/><text x="70" y="40" font-size="28" fill="#0A7B4F">示例</text></svg>')
    bb = B.load(proj)
    bb["logo"] = "logo.svg"
    B.save(proj, bb)
    p2, _ = E.build(d, proj, os.path.join(td, "svglogo.pptx"), cache={})
    with zipfile.ZipFile(p2) as z:
        check("SVG Logo 转成可编辑形状（PPTX 里没有图片）", not any(n.startswith("ppt/media/") for n in z.namelist()))
        check("SVG Logo 的品牌色不算风格之外", "0A7B4F" in z.read("ppt/slides/slide1.xml").decode("utf-8")
              and E.audit(p2, d, proj) == [])


# ------------------------------------------------------------------ AI 请求
def t_ai(td):
    import ai as A
    proj = os.path.join(td, "aiproj")
    shutil.copytree(DEMO, proj)
    path = os.path.join(proj, "deck.json")
    r = A.add(proj, "module", "s10", "m4", None, "改口语")
    check("记一条请求", r["status"] == "pending" and A.pending(proj)[0]["id"] == r["id"])
    ctx = A.context(proj, r)
    check("给 AI 的上下文：目标卡片、页码、名片摘要、规矩", ctx["target"]["type"] == "bullets" and ctx["page_no"] == pno(demo(), "s10") + 1
          and ctx["brand"]["firm"] and "案号不编" in ctx["rules"])
    for label, args, frag in [("不认识的请求", ("draw", "s10"), "不认识"), ("没有这一页", ("page", "s99"), "没有这一页"),
                              ("没有这张卡", ("module", "s10", "m99"), "没有这张卡"), ("工具没写", ("tool", "s10", "m2"), "V1–V6")]:
        check(f"请求拒：{label}", raises(lambda: A.add(proj, *args), A.AIError, frag))
    before = open(path, encoding="utf-8").read()
    check("结果不合规：拒，课件不动", raises(lambda: A.apply(proj, r["id"], {"module": {"x": 11}}), A.AIError, "校验")
          and open(path, encoding="utf-8").read() == before)
    r2, probs = A.apply(proj, r["id"], {"module": {"items": ["先看合同", "再看扣息"]}, "note": "好了"})
    m = next(x for x in page(D.load(path), "s10")["modules"] if x["id"] == "m4")
    check("交回结果：并进目标卡片，位置不动", m["items"] == ["先看合同", "再看扣息"] and (m["x"], m["y"]) == (6, 1) and r2["status"] == "done")
    check("已完成的请求不能再交", raises(lambda: A.apply(proj, r["id"], {"module": {"items": ["x"]}}), A.AIError, "done"))
    r = A.add(proj, "new-page", "s10", None, None, "砍头息")
    A.apply(proj, r["id"], {"slides": [{"modules": [{"type": "title", "x": 0, "y": 0, "w": 12, "h": 1, "text": "砍头息"}]},
                                       {"modules": [{"type": "title", "x": 0, "y": 0, "w": 12, "h": 1, "text": "第二页"}]}]})
    ids = [s["id"] for s in D.load(path)["slides"]]
    i = ids.index("s10")
    check("加页：插在这一页后面、id 自动编", len(ids) == len(demo()["slides"]) + 2 and len(set(ids)) == len(ids)
          and page(D.load(path), ids[i + 1])["modules"][0]["text"] == "砍头息")
    r = A.add(proj, "page", "s10", None, None, "重做")
    _, probs = A.apply(proj, r["id"], {"modules": [{"type": "title", "x": 0, "y": 0, "w": 12, "h": 1, "text": "利率"},
                                                   {"type": "statute", "x": 0, "y": 1, "w": 6, "h": 3, "law": "民法典"}]})
    check("重做一页：写进去，并把还要改的告诉 AI（法条没写原文）", any("法条" in p for p in probs), probs)
    r = A.add(proj, "tool", "s6", "m2", "V6", "画上诉")
    check("素材文件名只许 ASCII", raises(lambda: A.apply(proj, r["id"], {"assets": {"上诉.json": "{}"}, "module": {"file": "assets/x.json"}}),
                                      A.AIError, "文件名"))
    A.apply(proj, r["id"], {"assets": {"appeals2.json": json.load(open(os.path.join(DEMO, "assets", "sankey-appeals.json"), encoding="utf-8"))},
                            "module": {"file": "assets/appeals2.json"}})
    check("用 V6 出图：图表说明写进 assets/，卡片引用它", os.path.exists(os.path.join(proj, "assets", "appeals2.json"))
          and next(x for x in page(D.load(path), "s6")["modules"] if x["id"] == "m2")["file"] == "assets/appeals2.json")
    r = A.add(proj, "cover", "s1", None, None, "")
    A.set_status(proj, r["id"], "cancelled")
    check("取消的请求不能再交", raises(lambda: A.apply(proj, r["id"], {"modules": []}), A.AIError, "cancelled"))
    rc = subprocess.run([sys.executable, os.path.join(ROOT, "scripts", "ai.py"), proj, "wait", "--timeout", "1"],
                        capture_output=True, text=True, encoding="utf-8")
    check("ai.py wait：没有新请求就按时退出（退出码 4）", rc.returncode == 4, rc.stdout + rc.stderr)
    A.add(proj, "page", "s2", None, None, "章节页换个说法")
    rc = subprocess.run([sys.executable, os.path.join(ROOT, "scripts", "ai.py"), proj, "wait", "--timeout", "5"],
                        capture_output=True, text=True, encoding="utf-8")
    check("ai.py wait：有请求就交出上下文", rc.returncode == 0 and json.loads(rc.stdout)[0]["request"]["prompt"] == "章节页换个说法")
    check("请求队列的暂存文件不留", not any(f.endswith(".nlv-new") for f in os.listdir(os.path.join(proj, "ai"))))

    # ---- 0.8.4：AI 交回的结果还有要改的，要能「改到没有为止」；几方同时写不丢 ----
    pr4 = os.path.join(td, "ai-loop")
    shutil.copytree(DEMO, pr4)
    n0 = len(D.load(os.path.join(pr4, "deck.json"))["slides"])
    rq4 = A.add(pr4, "new-page", "s3", None, None, "加一页")
    page4 = lambda v: {"slides": [{"modules": [  # noqa: E731
        {"type": "title", "x": 0, "y": 0, "w": 12, "h": 1, "text": "收案"},
        {"type": "chart", "x": 0, "y": 1, "w": 8, "h": 5, "file": "assets/bar-filings.json"},
        {"type": "stat", "x": 8, "y": 1, "w": 4, "h": 2, "value": v, "label": "2025 年较 2019 年", "from": "assets/bar-filings.json"}]}]}
    r4, p4 = A.apply(pr4, rq4["id"], page4("+60%"))
    check("AI 交回的大数字推不出：写进去让用户看见，但请求留在「正在做」（原来算「做好了」，再交被拒，没法改到没有为止）",
          p4 and r4["status"] == "working", (r4["status"], p4[:1]))
    try:
        r4b, p4b = A.apply(pr4, rq4["id"], page4("+56%"))
    except A.AIError as e:
        r4b, p4b = {"status": str(e)}, ["拒"]
    d4 = D.load(os.path.join(pr4, "deck.json"))
    check("对同一个请求再交：换掉上次加的那一页，不越加越多；改对了才算做好了",
          r4b["status"] == "done" and not p4b and len(d4["slides"]) == n0 + 1
          and sum(1 for s in d4["slides"] for m in s["modules"] if m.get("value") in ("+56%", "+60%")
                  and m.get("from")) == 1, (r4b.get("status"), len(d4["slides"]), n0))
    ctx4 = A.context(pr4, A.add(pr4, "tool", "s3", "m2", "V6", "换个图"))
    ans4 = ctx4.get("answer") or {}
    cmd4 = ans4.get("apply", "")
    check("给 AI 的上下文自带交回的写法与照抄就能跑的命令（脚本与课件都是真实路径），出图请求带上那个模块 SKILL.md 的路径",
          "slides" not in ans4.get("format", "x") and '"module"' in ans4.get("format", "") and os.path.abspath(A.__file__) in cmd4
          and os.path.abspath(pr4) in cmd4 and os.path.isfile((ctx4.get("tool") or {}).get("skill_md") or ""), ans4)
    pq = os.path.join(td, "ai-race")
    shutil.copytree(DEMO, pq)
    race = ("import sys,os;sys.path.insert(0,sys.argv[1]);import ai as A\n"
            "for i in range(25):\n    A.add(sys.argv[2],'page','s3',None,None,f'{os.getpid()}-{i}')\n")
    rf = os.path.join(td, "race.py")
    with open(rf, "w", encoding="utf-8") as f:
        f.write(race)
    procs = [subprocess.Popen([sys.executable, rf, os.path.join(ROOT, "scripts"), pq], stderr=subprocess.PIPE)
             for _ in range(4)]
    errs = [p.communicate(timeout=120)[1] for p in procs]
    try:
        qq = A.load_q(pq)
    except ValueError as e:                     # 原来实测还会把请求文件写坏（两份 JSON 首尾相接）
        qq, errs = [], [str(e).encode("utf-8")]
    check("四个进程同时各记 25 条请求：100 条一条不丢、编号不重、没有进程崩、请求文件没被写坏（原来只剩 29 条，还会写坏文件）",
          len(qq) == 100 and len({r["id"] for r in qq}) == 100 and all(p.returncode == 0 for p in procs),
          (len(qq), [e.decode("utf-8", "replace")[-120:] for e in errs if e][:1]))


# ------------------------------------------------------------------ 读模板（律所风格第一步）
    # 结果没过校验：同名素材原样换回（开源前第三轮实跑：原来先写后验，别的页在用的图被悄悄换掉）
    import ai as AIq
    pa_ = os.path.join(td, "ai-rollback")
    shutil.copytree(DEMO, pa_)
    old_ = open(os.path.join(pa_, "assets", "bar-filings.json"), "rb").read()
    rq_ = AIq.add(pa_, "page", slide="s3", prompt="重做")
    try:
        AIq.apply(pa_, rq_["id"], {"assets": {"bar-filings.json": {"chart": "bar", "title": "坏"}, "fresh.json": {"chart": "bar"}},
                                   "modules": [{"type": "nosuch", "x": 0, "y": 0, "w": 4, "h": 2}]})
        rej_ = False
    except AIq.AIError:
        rej_ = True
    check("AI 结果没过校验：课件不动，同名的旧素材原样换回，新素材挪开（选图里不出现）",
          rej_ and open(os.path.join(pa_, "assets", "bar-filings.json"), "rb").read() == old_
          and not os.path.exists(os.path.join(pa_, "assets", "fresh.json")))

def t_template_read(td):
    import template_read as TR
    d = demo()
    src = os.path.join(td, "tpl-src.pptx")
    E.build(d, DEMO, src, cache={}, paste=False)                      # 用自己导出的示例课件当一份「模板」：有 Logo、有封面
    out = os.path.join(td, "tpl-read")
    rc = TR.main([src, "-o", out])
    check("读模板：跑得通、写出 tokens.json 与 report.html", rc == 0 and os.path.exists(os.path.join(out, "tokens.json"))
          and os.path.exists(os.path.join(out, "report.html")))
    tk = json.load(open(os.path.join(out, "tokens.json"), encoding="utf-8"))
    check("读模板：16:9、页数对", tk["slide"]["is_16_9"] and len(tk["pages"]) == len(d["slides"]))
    check("读模板：找到反复出现的 Logo，并把图片取出来", tk["logo"] and tk["logo"].get("file") == "logo.png"
          and os.path.exists(os.path.join(out, "logo.png")), tk.get("logo"))
    check("读模板：第一页算封面", tk["pages"][0]["kind"] == "cover")
    check("读模板：可读性检查有对比度数值", tk["checks"] and all(c["ratio"] >= 1 for c in tk["checks"]))
    check("读模板：只写进 -o 目录，不碰原文件", open(src, "rb").read()[:2] == b"PK")
    bad = os.path.join(td, "bad.pptx")
    with open(bad, "wb") as f:
        f.write(b"not a zip")
    check("读模板：不是 .pptx 就说清楚，退出码 1", TR.main([bad, "-o", os.path.join(td, "x")]) == 1)
    lum = TR.contrast("#FFFFFF", "#000000")
    check("对比度计算：黑白 = 21", abs(lum - 21) < 0.01)


# ------------------------------------------------------------------ 律所风格（第二步：模板 → 风格）
FIRM_ACC = "#0F5C6E"                                     # 虚构律所的主色（测试用；真律所的模板不进仓库）


def _fake_template(td):
    """用自己导出的示例课件做一份「律所模板」：把主题主色换成虚构的青色。"""
    src = os.path.join(td, "fake-src.pptx")
    E.build(demo(), DEMO, src, cache={}, paste=False)
    out = os.path.join(td, "fake-template.pptx")
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for n in zi.namelist():
            data = zi.read(n)
            if n.startswith("ppt/theme/theme"):
                x = data.decode("utf-8")
                x = re.sub(r"<a:accent1>.*?</a:accent1>", f'<a:accent1><a:srgbClr val="{FIRM_ACC[1:]}"/></a:accent1>', x, flags=re.S)
                data = x.encode("utf-8")
            zo.writestr(n, data)
    return out


def t_firm(td):
    import brand as B
    import firm_style as FS
    import template_read as TR
    check("PNG 亮度：白字版 Logo 很亮", TR.png_ink(_png(8, 4, b"\xff\xff\xff\xff")) > 0.9)
    check("PNG 亮度：彩色 Logo 暗", TR.png_ink(_png(8, 4)) < 0.3)
    check("PNG 亮度：全透明的读不出", TR.png_ink(_png(8, 4, b"\xff\xff\xff\x00")) is None)
    tpl = _fake_template(td)
    proj = os.path.join(td, "firmproj")
    shutil.copytree(DEMO, proj)
    data = open(tpl, "rb").read()
    for label, args, frag in [("不是 .pptx", (data, None, None, "a.ppt"), ".pptx"), ("不是 zip", (b"nope", None, None, "a.pptx"), "不是"),
                              ("太大", (b"PK" + b"0" * (B.TEMPLATE_MAX + 1), None, None, "a.pptx"), "太大")]:
        check(f"导入模板拒：{label}", raises(lambda: B.import_template(proj, *args), B.BrandError, frag))
    summ = B.import_template(proj, data, "示例律所风", True, "fake-template.pptx")
    check("导入模板：生成律所风格，主色来自模板主题", summ["name"] == "示例律所风" and summ["colors"]["accent"] == FIRM_ACC, summ)
    check("导入模板：只留读出的东西，不留模板本身", os.path.exists(os.path.join(proj, "brand", "template", "report.html"))
          and not any(f.endswith(".pptx") for f in os.listdir(os.path.join(proj, "brand", "template"))))
    check("导入模板：取出 Logo 放进名片", "logo" in summ["logos"] and B.load(proj).get("logo"))
    st, err = B.firm_style(proj)
    check("律所风格文件校验通过", st and not err, err)
    d = demo()
    d["style"] = "示例律所风"
    check("课件可以用律所风格（在课件目录里校验）", D.validate(d, proj) is None)
    check("不在课件目录里校验：律所风格名不认", raises(lambda: D.validate(d), D.DeckError, "style"))
    d2 = dict(d, style="不存在的风")
    check("律所风格名写错：拒，并列出能用的风格", raises(lambda: D.validate(d2, proj), D.DeckError, "自定义"))
    d3 = dict(d, style="自定义")
    check("自定义：课件写「自定义」就用这份自定义风格", D.validate(d3, proj) is None and D.firm_of(d3, proj) is not None)
    check("自定义：没做过自定义风格时写「自定义」就拒，并说怎么做",
          raises(lambda: D.validate(d3, DEMO), D.DeckError, "上传模板"))
    with open(os.path.join(proj, "deck.json"), "w", encoding="utf-8") as f:
        json.dump(d, f, ensure_ascii=False)
    check("自定义：旧课件写的是自定义风格的旧名字，读进来认作「自定义」", D.load(os.path.join(proj, "deck.json"))["style"] == "自定义")
    check("自定义：编辑器读旧课件也认作「自定义」", studio.Project(proj).load()["style"] == "自定义")
    check("自定义：可用风格列表里叫「自定义」，不出现律所名", D.styles_for(proj) == ["奇川风", "歸藏风", "自定义"], D.styles_for(proj))
    import firm_style as FS
    t1 = FS.tokens_from_color("#FFCC00")
    check("自定义·主色：太浅的主色压深到看得清", t1["adjusted"] and FS.contrast(t1["colors"]["roles"]["primary"], "#FFFFFF") >= 4.5, t1)
    check("自定义·主色：生成的风格过得了全部对比度检查", FS.make(t1, "自定义")["name"] == "自定义")
    check("自定义·主色：灰色拒", raises(lambda: FS.tokens_from_color("#777777"), FS.FirmStyleError, "灰"))
    check("自定义·主色：写错格式拒", raises(lambda: FS.tokens_from_color("紫色"), FS.FirmStyleError, "#RRGGBB"))
    check("自定义·截图：从截图里取到主色（不被抗锯齿的浅边带偏）", FS.main_color_of(_shot()) == "#702082", FS.main_color_of(_shot()))
    check("自定义·截图：黑白灰的图拒，并让选主色", raises(lambda: FS.main_color_of(_png(60, 40, b"\x80\x80\x80\xff")), FS.FirmStyleError, "主色"))
    allowed = {c.lstrip("#").upper() for c in D.allowed_colors(d, proj)}
    figs = {c.lstrip("#").upper() for c in colors_in(open(os.path.join(proj, "assets", "v3-relation.svg"), encoding="utf-8").read())}
    bad_pages, warn_fig = [], False
    for i, sl in enumerate(d["slides"]):
        svg, rep_ = D.compose(d, i, proj, {})
        extra = colors_in(svg) - allowed - (figs if any(m["type"] == "figure" for m in sl["modules"]) else set())
        if rep_["problems"] or extra:
            bad_pages.append((sl["id"], rep_["problems"][:1], sorted(extra)[:3]))
        warn_fig |= any("白描" in w for w in rep_["warnings"])
    check("律所风格：整套示例课件都能拼、只用律所风格的颜色", not bad_pages, bad_pages)
    check("律所风格：V1–V5 的图不拒，提醒回原模块出白描版", warn_fig)
    svg, _ = D.compose(d, pno(d, "s1"), proj, {})
    check("律所风格：封面是主色满版", f'fill="{FIRM_ACC}" data-role="paper"' in svg)
    svg, _ = D.compose(d, pno(d, "s3"), proj, {})
    check("律所风格：图表重音换成律所主色，奇川风深红不再出现", FIRM_ACC in svg and "#991B1B" not in svg)
    check("律所风格：图标题与页标题都用粗黑体（无衬线、700）", "方正小标宋" not in svg and 'font-weight="700"' in svg)
    dk = one([M("title", 0, 0, 12, 2, "t", level="cover", text="封面"), M("chart", 0, 2, 8, 4, "c", file="assets/bar-filings.json")])
    dk["style"] = "示例律所风"
    _, rep_ = D.compose(dk, 0, proj, {})
    check("律所风格：深底页上放图表 → 拦下，放到内容页", any("深底" in p for p in rep_["problems"]))
    out = os.path.join(td, "firm.pptx")
    p, _ = E.build(d, proj, out, cache={})
    check("律所风格：导出并回读审计全过", E.audit(p, d, proj) == [])
    # 两版 Logo：深底页用白字版
    B.save_logo(proj, "w.png", _png(40, 12, b"\xff\xff\xff\xff"), "logo_on_dark")
    bb = B.load(proj)
    check("名片里有深底用的 Logo", bb.get("logo_on_dark") and B.logo_info(bb, True)["path"].endswith("logo-on-dark.png")
          and B.logo_info(bb, False)["path"].endswith("logo.png"))
    svg_cover, _ = D.compose(d, pno(d, "s1"), proj, {})
    svg_body, _ = D.compose(d, pno(d, "s3"), proj, {})
    uri_dark = B.data_uri(B.logo_info(bb, True))
    check("深底页用白字版 Logo、浅底页用浅底版", uri_dark in svg_cover and uri_dark not in svg_body)
    # 律所风格文件的校验
    good = json.load(open(os.path.join(proj, "brand", "style.json"), encoding="utf-8"))
    for label, mut, frag in [
        ("主色太浅", lambda x: x["colors"].update(accent="#E0F0F4"), "对比度"),
        ("主色是灰", lambda x: x["colors"].update(accent="#6B6B6B"), "灰"),
        ("缺颜色", lambda x: x["colors"].pop("card"), "card"),
        ("和内置风格同名", lambda x: x.update(name="奇川风"), "同名"),
        ("卡片上的字看不清", lambda x: x["colors"].update(card="#555555"), "看不清"),
        ("封面底色太浅", lambda x: x["colors"].update(cover_bg="#F0F0F0"), ""),
    ]:
        x = copy.deepcopy(good)
        mut(x)
        check(f"律所风格文件拒：{label}", raises(lambda: FS.check(x), FS.FirmStyleError, frag))
    # 坏的风格文件：课件说清楚哪里不行
    with open(os.path.join(proj, "brand", "style.json"), "w", encoding="utf-8") as f:
        json.dump(dict(good, colors=dict(good["colors"], accent="#EEEEEE")), f)
    check("律所风格文件坏了：校验时说清是哪个文件、哪里不行", raises(lambda: D.validate(d, proj), D.DeckError, "style.json"))
    with open(os.path.join(proj, "brand", "style.json"), "w", encoding="utf-8") as f:
        json.dump(good, f, ensure_ascii=False)
    B.disable_firm_style(proj)
    check("停用律所风格：改名留着，不删文件", os.path.exists(os.path.join(proj, "brand", "style.disabled.json"))
          and not os.path.exists(os.path.join(proj, "brand", "style.json")))
    rc = TR.main([tpl, "-o", os.path.join(td, "fs-read")])
    rc2 = FS.main(["make", os.path.join(td, "fs-read", "tokens.json"), "--name", "命令行律所风", "-o", os.path.join(td, "fs", "style.json")])
    rc3 = FS.main(["check", os.path.join(td, "fs", "style.json")])
    check("命令行：读模板 → 生成律所风格 → 校验，三步都通", rc == 0 and rc2 == 0 and rc3 == 0)


# ------------------------------------------------------------------ 导出
def t_export(td):
    d = demo()
    out = os.path.join(td, "exports", "demo.pptx")
    p, rep = E.build(d, DEMO, out, cache={})
    check("示例课件导出", os.path.exists(p) and rep["slides"] == len(d["slides"]), rep)
    check("导出不留暂存文件", not any(f.endswith(".nlv-new") for f in os.listdir(os.path.dirname(out))))
    check("回读审计全过", E.audit(p, d, DEMO) == [])
    with zipfile.ZipFile(p) as z:
        names = z.namelist()
        check("PPTX：除了律所 Logo 没有别的图片", all("/" + E.LOGO_MEDIA in n for n in names if n.startswith("ppt/media/")))
        slides = [n for n in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)]
        check("PPTX 页数对", len(slides) == len(d["slides"]), len(slides))
        sp = sum(z.read(n).decode("utf-8").count("<p:sp") for n in slides)
        check("每页都是可编辑的形状与文字", sp > 20 * len(slides), sp)
        ct = z.read("[Content_Types].xml").decode("utf-8")
        check("每页都在 Content_Types 里登记", all("/" + n in ct for n in slides))
        ids = re.findall(r'<p:sldId [^>]*id="(\d+)"', z.read("ppt/presentation.xml").decode("utf-8"))
        check("页 id 不重复", len(ids) == len(set(ids)) == len(slides))
        s5 = z.read(f"ppt/slides/slide{pno(d, 's5') + 1}.xml").decode("utf-8")
        check("连组的卡片导出成自由形状（四角不同的圆角）", "<a:custGeom" in s5, "")
    try:
        from pptx import Presentation
        prs = Presentation(p)
        check("python-pptx 能打开", len(prs.slides) == len(d["slides"]))
    except ImportError:
        check("python-pptx 能打开（未装，跳过）", True)
    with open(out, "rb") as f:
        old = f.read()
    bad = copy.deepcopy(d)
    bad["style"] = "歸藏风"
    check("有不行的页：拒绝导出", raises(lambda: E.build(bad, DEMO, out, cache={}), E.ExportError, f"第 {pno(d, 's8') + 1} 页"))
    with open(out, "rb") as f:
        check("拒绝导出时旧文件原样不动", f.read() == old)
    tam = os.path.join(td, "tampered.pptx")
    with zipfile.ZipFile(p) as zi, zipfile.ZipFile(tam, "w", zipfile.ZIP_DEFLATED) as zo:
        for n in zi.namelist():
            data = zi.read(n)
            if n == f"ppt/slides/slide{pno(d, 's7') + 1}.xml":
                data = data.decode("utf-8").replace("三点实务提示", "别的字", 1)
                data = re.sub(r'srgbClr val="[0-9A-Fa-f]{6}"', 'srgbClr val="00FF00"', data, count=1).encode("utf-8")
            zo.writestr(n, data)
    bad = E.audit(tam, d, DEMO)
    check("审计抓到外来颜色（绿色）", any("00FF00" in b for b in bad), bad)
    check("审计抓到缺字", any("三点实务提示" in b for b in bad), bad)
    g = copy.deepcopy(d)
    g["style"] = "歸藏风"
    g["slides"] = [s for s in g["slides"] if not any(m["type"] == "figure" for m in s["modules"])]
    p2, _ = E.build(g, DEMO, os.path.join(td, "g.pptx"), cache={})
    check("歸藏风整套导出", E.audit(p2, g, DEMO) == [])
    p3, _ = E.build(d, DEMO, os.path.join(td, "again.pptx"), cache={})
    with zipfile.ZipFile(p) as a, zipfile.ZipFile(p3) as b:
        same = sorted(a.namelist()) == sorted(b.namelist()) and all(
            a.read(n) == b.read(n) for n in a.namelist() if n.startswith("ppt/slides/"))
    check("同一份说明导两次结果一样", same)
    pk = copy.deepcopy(d)
    page(pk, "s3")["parked"] = [{"id": "z", "type": "text", "x": 0, "y": 0, "w": 4, "h": 2, "text": "暂存里的字"}]
    p4, _ = E.build(pk, DEMO, os.path.join(td, "parked.pptx"), cache={})
    with zipfile.ZipFile(p4) as z:
        check("暂存区不导出", not any("暂存里的字" in z.read(n).decode("utf-8") for n in z.namelist() if n.endswith(".xml")))


# ------------------------------------------------------------------ CLI
def t_paste(td):
    """作者定的（2026-10-05）：导出的图表页是给律师复制进自己课件的——整页粘、框选粘都要不变样、不盖母版。"""
    d = demo()
    p, rep = E.build(d, DEMO, os.path.join(td, "paste.pptx"), cache={})
    figs = [(i, m["id"]) for i, s in enumerate(d["slides"]) for m in s["modules"] if m["type"] in ("chart", "figure")]
    with zipfile.ZipFile(p) as z:
        xs = [z.read(f"ppt/slides/slide{i + 1}.xml").decode("utf-8") for i in range(len(d["slides"]))]
    check("粘贴：每页都有页面背景（底色不画成形状）", all("<p:bg>" in x for x in xs))
    full = re.compile(r'<a:off x="0" y="0"\s*/><a:ext cx="12192000" cy="6858000"\s*/>')
    check("粘贴：没有整页大小的底色块（整页粘贴不盖目标课件的 Logo、页脚）",
          not any(full.search(x) for x in xs))
    check("粘贴：每张图都打成一组（点一下选中整张图）",
          figs and all(f'descr="nlv-group:{mid}"' in xs[i] for i, mid in figs), figs[:4])
    check("粘贴：报告里数得出组数", rep["groups"] >= len(figs), rep.get("groups"))
    check("粘贴：不用主题色、不用占位符", not any("schemeClr" in x or "<p:ph" in x for x in xs))
    pn = [f"{i + 1:02d}" for i in range(len(xs))]
    check("粘贴：不画页码（目标课件有自己的页码）",
          not any(f"<a:t>{n}</a:t>" in x for n, x in zip(pn, xs)))
    try:
        from pptx import Presentation
        prs = Presentation(p)
        grp = [sh for s in prs.slides for sh in s.shapes if sh.shape_type == 6]
        check("粘贴：python-pptx 读得出这些组，组里有形状", grp and all(len(g.shapes) >= 2 for g in grp), len(grp))
    except ImportError:
        check("粘贴：python-pptx 读组（未装，跳过）", True)
    s8 = next(i for i, s_ in enumerate(d["slides"]) if any(m["type"] == "figure" and "v3" in (m.get("file") or "") for m in s_["modules"]))
    names = re.findall(r'name="([^"]*)" descr="nlv-anchor:', xs[s8])
    check("逐个讲：关系图里每个主体、每条关系都按名字单独成形状或成组（PowerPoint 里能逐个加动画）",
          len(names) == 20 and "欠付范围内担责" in names and any(n.startswith("发包人") for n in names), names[:6])
    ci = next(i for i, s_ in enumerate(d["slides"]) if any(m["type"] == "chart" for m in s_["modules"]))
    check("逐个讲：统计图每个数据项按「数值 类别」命名", len(re.findall(r'descr="nlv-anchor:id:', xs[ci])) >= 3)
    try:
        from pptx import Presentation
        sl = Presentation(p).slides[s8]
        fig = next(sh for sh in sl.shapes if sh.shape_type == 6 and sh.name.startswith("V1–V5"))
        subs = [sh.name for sh in fig.shapes if sh.shape_type == 6]
        check("逐个讲：python-pptx 读得出图组里的关系子组", "欠付范围内担责" in subs, subs[:5])
    except ImportError:
        check("逐个讲：python-pptx（未装，跳过）", True)
    # 只导出这一页
    i0 = figs[0][0]
    one, r1 = E.build(d, DEMO, os.path.join(td, "one.pptx"), cache={}, pages=[i0])
    with zipfile.ZipFile(one) as z:
        n1 = len([n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)])
    check("只导出这一页：PPTX 只有 1 页", n1 == 1 and r1["slides"] == 1, n1)
    try:
        E.build(d, DEMO, os.path.join(td, "bad.pptx"), cache={}, pages=[999])
        check("只导出这一页：页号不对就拒", False)
    except E.ExportError as e:
        check("只导出这一页：页号不对就拒", "页号" in str(e))
    # 守卫能失败：往一页里塞回整页白底块、拆掉组
    tam = os.path.join(td, "paste-bad.pptx")
    with zipfile.ZipFile(p) as zi, zipfile.ZipFile(tam, "w", zipfile.ZIP_DEFLATED) as zo:
        for n in zi.namelist():
            data = zi.read(n)
            if n == f"ppt/slides/slide{i0 + 1}.xml":
                x = data.decode("utf-8").replace("<p:bg>", "<p:bgX>", 1).replace("</p:bg>", "</p:bgX>", 1)
                x = x.replace(f'descr="nlv-group:{figs[0][1]}"', 'descr="x"')
                x = x.replace("</p:grpSpPr>", '</p:grpSpPr><p:sp><p:nvSpPr><p:cNvPr id="9990" name="白底"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
                              '<p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="12192000" cy="6858000"/></a:xfrm>'
                              '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:sp>', 1)
                data = x.encode("utf-8")
            zo.writestr(n, data)
    b = E.audit(tam, d, DEMO)
    check("守卫能失败：抓到整页底色块", any("整页大小" in x for x in b), b)
    check("守卫能失败：抓到没有页面背景", any("页面背景" in x for x in b), b)
    check("守卫能失败：抓到图没打组", any("一组" in x for x in b), b)


def _shot(w=640, h=360, filt=0):
    """一张假的「模板空白页」截图：白底，左上一块 Logo（第 20–50 行），底下一道页脚线（第 335–340 行）。"""
    import struct
    import zlib
    rows = []
    for y in range(h):
        line = bytearray()
        for x in range(w):
            if 20 <= y < 50 and 20 <= x < 120:
                line += b"\x70\x20\x82"
            elif 335 <= y < 340:
                line += b"\x99\x99\x99"
            else:
                line += b"\xff\xff\xff"
        rows.append(bytes(line))
    raw = b""
    prev = bytes(w * 3)
    for r in rows:
        if filt == 2:                                   # 「上」滤波：检验自带解码器
            raw += b"\x02" + bytes((a - b) & 255 for a, b in zip(r, prev))
        else:
            raw += b"\x00" + r
        prev = r

    def chunk(t_, d):
        return struct.pack(">I", len(d)) + t_ + d + struct.pack(">I", zlib.crc32(t_ + d) & 0xffffffff)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def t_frame(td):
    """作者定的（2026-10-05）：版面两种——铺满全屏 / 预留位置（让出律所模板的 Logo、天头地脚、侧边条）。"""
    import frame as F
    # 不装 Pillow 时走自带的 PNG 解码。CI 装了 Pillow，这条路平时跑不到——所以直接调它，不靠环境。
    for (w_, h_), why in (((60, 40), "图比抽样格小"), ((400, 300), "图比抽样格大")):
        try:
            sw, sh, sg = F._sample_png(_png(w_, h_), 160, 90)
            ok_ = (sw, sh) == (w_, h_) and len(sg) == 90 and all(r_ and len(r_) == 160 and r_[0] == (0x99, 0x1B, 0x1B) for r_ in sg)
        except Exception as e:                                    # noqa: BLE001
            ok_, sg = False, repr(e)
        check(f"截图读取（不装 Pillow）：{why}也每一格都有颜色", ok_, str(sg)[:120])
    g0 = G.geom(None)
    check("版面：铺满全屏的网格就是原来的常数", (g0["MX"], g0["MY"], g0["COL"], g0["ROW"]) == (G.MX, G.MY, G.COL, G.ROW), g0)
    fr = {"mode": "reserve", "top": 121, "right": 61, "bottom": 57, "left": 60}
    g = G.geom(fr)
    check("版面：预留位置后仍是 12 栏 × 6 行，正好铺满内容区",
          abs(g["MX"] + 12 * g["COL"] + 11 * G.GUT - (G.W - 61)) < 1e-6 and abs(g["MY"] + 6 * g["ROW"] + 5 * G.GUT - (G.H - 57)) < 1e-6, g)
    check("版面：整页出血的图只铺满内容区", list(G.px({"bleed": True, "x": 0, "y": 0, "w": 12, "h": 6}, g)) == [60, 121, 1159, 542])
    for bad, why in (({"mode": "x"}, "mode"), (dict(fr, top=-1), "负数"), (dict(fr, top="10"), "不是数"),
                     (dict(fr, top=300, bottom=300), "留太多")):
        check(f"版面：校验拒绝（{why}）", G.check_frame(bad) != [], bad)
    check("版面：铺满全屏不管四边数字", G.check_frame({"mode": "full", "top": -5}) == [])
    d = one([M("title", 0, 0, 12, 1, "t", text="预留位置的一页"), M("chart", 0, 1, 8, 5, "c", file="assets/bar-filings.json"),
             M("bullets", 8, 1, 4, 5, "b", items=["一", "二"])])
    d["frame"] = dict(fr)
    D.validate(d, DEMO)
    bad = dict(d, frame=dict(fr, top=400))
    try:
        D.validate(bad, DEMO)
        check("版面：课件校验拦下留得太多", False)
    except D.DeckError as e:
        check("版面：课件校验拦下留得太多", "预留" in str(e), str(e))
    svg, rep = D.compose(d, 0, DEMO, {})
    check("版面：预留位置这一页拼得出来", not rep["problems"], rep["problems"])
    check("版面：预留位置时不画页码（模板自己有）", 'data-mod="page"' not in svg)
    p, _ = E.build(d, DEMO, os.path.join(td, "frame.pptx"), cache={})
    with zipfile.ZipFile(p) as z:
        x = z.read("ppt/slides/slide1.xml").decode("utf-8")
    offs = [(int(a) / E.EMU, int(b) / E.EMU, int(c) / E.EMU, int(e) / E.EMU) for a, b, c, e in
            re.findall(r'<a:off x="(-?\d+)" y="(-?\d+)"\s*/><a:ext cx="(\d+)" cy="(\d+)"', x)]
    out_ = [o for o in offs if o[0] < 60 - 1 or o[1] < 121 - 1 or o[0] + o[2] > G.W - 61 + 1 or o[1] + o[3] > G.H - 57 + 1]
    check("版面：导出的形状都落在内容区里，预留区留空", offs and not out_, out_[:3])
    # 从模板几何算留白
    tok = {"geometry": {"margin_left": 60, "margin_right": 61, "date": [60, 679, 198, 19], "page_no": [1151, 679, 69, 19]},
           "logo": {"box": [60, 57, 212, 48]}}
    f1 = F.from_tokens(tok)
    check("版面·模板：Logo 在上方 → 上留到 Logo 下沿再加间距", f1["top"] == 57 + 48 + F.GAP, f1)
    check("版面·模板：页码在下方 → 下留到页码上沿再加间距", f1["bottom"] == 720 - 679 + F.GAP, f1)
    check("版面·模板：左右按页边距", (f1["left"], f1["right"]) == (60, 61), f1)
    f2 = F.from_tokens({"geometry": {}})
    check("版面·模板：什么都没读到就用默认页边距", (f2["top"], f2["left"]) == (G.MY, G.MX) and f2["notes"], f2)
    # 从截图算留白：Pillow 与自带 PNG 解码两条路一致
    shot = _shot()
    fa = F.from_image(shot)
    keep = F._sample_pillow
    try:
        F._sample_pillow = lambda *a: None
        fb = F.from_image(_shot(filt=2))
    finally:
        F._sample_pillow = keep
    check("版面·截图：Logo 撑出上留白（约 100 + 16）", 112 <= fa["top"] <= 124, fa)
    check("版面·截图：页脚线撑出下留白（约 50 + 16）", 60 <= fa["bottom"] <= 72, fa)
    check("版面·截图：左右没有设计元素就按默认页边距", (fa["left"], fa["right"]) == (G.MX, G.MX), fa)
    check("版面·截图：没装 Pillow 时自带解码结果一样", all(fa[k] == fb[k] for k in ("top", "right", "bottom", "left")), (fa, fb))
    check("版面·截图：空白页不报「中间有内容」", not any("中间" in n for n in fa["notes"]), fa["notes"])
    try:
        F.from_image(b"not an image")
        check("版面·截图：不是图片就拒", False)
    except F.FrameError:
        check("版面·截图：不是图片就拒", True)
    # 编辑器的 JS 与 Python 算出同一套网格
    node = shutil.which("node")
    if node:
        frames = [None, {"mode": "full"}, fr, {"mode": "reserve", "top": 96, "right": 48, "bottom": 72, "left": 80}]
        js = os.path.join(td, "geom.js")
        with open(js, "w", encoding="utf-8") as f:
            f.write("const G = require(%s); const base = %s; const frames = %s;\n"
                    "process.stdout.write(JSON.stringify(frames.map((fr) => { const g = G.geom(base, fr); "
                    "return [g.MX, g.MY, g.COL, g.ROW, g.BLEED, G.px({grid: g}, {x: 3, y: 2, w: 4, h: 3})]; })));"
                    % (json.dumps(GRID_JS), json.dumps(studio.grid_json()["grid"]), json.dumps(frames)))
        r = subprocess.run([node, js], capture_output=True, text=True, encoding="utf-8")
        py = [[G.geom(f_)["MX"], G.geom(f_)["MY"], G.geom(f_)["COL"], G.geom(f_)["ROW"], G.geom(f_)["BLEED"],
               list(G.px({"x": 3, "y": 2, "w": 4, "h": 3}, G.geom(f_)))] for f_ in frames]
        jsr = json.loads(r.stdout) if r.returncode == 0 else None
        same = jsr is not None and all(abs(a - b) < 1e-9 for pa, ja in zip(py, jsr)
                                       for a, b in zip([*pa[:4], *pa[4], *pa[5]], [*ja[:4], *ja[4], *ja[5]]))
        check("版面：编辑器与 Python 的网格几何逐项一致", same, (py, jsr, r.stderr[-300:]))
    else:
        check("版面：规则对照（未装 node，跳过）", True)


def t_pins_crop(td):
    """作者定的方向（2026-10-07）：磁吸做到图里——读图编号钩住图内元素，切片只取几件事放大。"""
    def page(fig_extra, bullets_extra=None, fw=5):
        d = demo()
        d["slides"] = [{"id": "p1", "modules": [
            M("title", 0, 0, 12, 1, "t", text="实际施工人凭什么向发包人主张工程价款"),
            dict(M("figure", 0, 1, fw, 5, "f", file="assets/v3-relation.svg"), **fig_extra),
            dict(M("bullets", fw, 1, 12 - fw, 5, "b", items=["实际施工人是原告", "诉请折价补偿", "标的 3200 万元", "欠付范围内担责"]),
                 **(bullets_extra or {}))]}]
        return d
    sel = ["node:#1", "node:#9", "node:#4", "edge:#10", "edge:#11"]
    d0 = page({}, fw=8)
    _, r0 = D.compose(d0, 0, DEMO, {})
    check("切片前：关系图放在半页里字太小，被拦下", any("11px" in p_ for p_ in r0["problems"]), r0["problems"])
    sv0, _ = D.compose(d0, 0, DEMO, {})
    check("切片前：字太小也照样画出来（律师要看着图去切片）", "发包人" in sv0)
    check("切片前：导出仍被拦下", raises(lambda: E.build(d0, DEMO, os.path.join(td, "small.pptx"), cache={}), E.ExportError, "11px"))
    d1 = page({"crop": sel, "pins": ["node:#4", "edge:#10", "node:#9", "edge:#11"]}, {"numbered": True})
    D.validate(d1, DEMO)
    svg, r1 = D.compose(d1, 0, DEMO, {})
    f = r1["modules"]["f"]
    check("切片：只取 5 个元素那一块，字过了 11px", not r1["problems"] and any("切片" in n for n in f["notes"]), r1["problems"])
    check("切片：没选中的元素不在页面上", "总承包人" not in svg and "监理单位" not in svg and "实际施工人" in svg)
    check("切片：没选中的关系只剩线头的不画", len([a for a in f["anchors"] if a["role"] == "edge"]) == 2,
          [a["label"] for a in f["anchors"] if a["role"] == "edge"])
    pins = f.get("pins") or []
    check("读图编号：4 个编号依次落在选中的元素上", [p_["n"] for p_ in pins] == [1, 2, 3, 4] and pins[1]["label"].startswith("本案诉请"), pins)
    cb = f["crop"]
    check("读图编号：都在图卡里", all(48 <= p_["box"][0] and p_["box"][0] + p_["box"][2] <= 48 + 5 * 84 + 4 * 16 for p_ in pins), pins)
    check("读图编号：不用重音色（一页只一处重音）", svg.count('fill="#991B1B" data-mod="p1.f"') <= 2)
    check("要点按编号排：条数对得上就不提醒", not r1["modules"]["b"]["warnings"], r1["modules"]["b"]["warnings"])
    d2 = page({"crop": sel, "pins": ["node:#4", "edge:#10"]}, {"numbered": True})
    _, r2 = D.compose(d2, 0, DEMO, {})
    check("要点按编号排：图上 2 个编号、要点 4 条，提醒对不上", any("对不上" in w for w in r2["modules"]["b"]["warnings"]))
    d3 = page({"crop": sel, "pins": ["node:#4", "node:#2"]})
    _, r3 = D.compose(d3, 0, DEMO, {})
    check("读图编号：指向切片外的元素就提醒重新放", any("找不到" in w for w in r3["modules"]["f"]["warnings"]))
    for bad, why in (({"pins": "node:#1"}, "不是列表"), ({"pins": ["node:#1", "node:#1"]}, "重复"),
                     ({"pins": [f"node:#{i}" for i in range(10)]}, "超过 9 个"), ({"crop": []}, "空切片")):
        try:
            D.validate(page(bad), DEMO)
            check(f"读图编号 / 切片：校验拒绝（{why}）", False)
        except D.DeckError:
            check(f"读图编号 / 切片：校验拒绝（{why}）", True)
    p, _ = E.build(d1, DEMO, os.path.join(td, "pins.pptx"), cache={})
    with zipfile.ZipFile(p) as z:
        x = z.read("ppt/slides/slide1.xml").decode("utf-8")
    check("读图编号：导出后每个编号单独成组，命名「编号 n」", re.findall(r'name="(编号 \d)"', x) == ["编号 1", "编号 2", "编号 3", "编号 4"],
          re.findall(r'name="(编号 \d)"', x))
    fi = r1["modules"]["f"].get("fig") or {}
    check("最佳尺寸：报告里给编辑器图的比例与 1:1 时的最小字（拉大拉小实时算字号）",
          fi.get("w") and fi.get("h") and abs(fi["min1"] - 11.0) < 0.6 and fi["pad"] >= 0, fi)
    mods_ = [M("bullets", 0, 2, 4, 3, "a", items=["x"]), dict(M("bullets", 6, 3, 4, 3, "b", items=["y"]), locked=True)]
    out_, _ = G.tidy(mods_)
    lk = next(m_ for m_ in out_ if m_["id"] == "b")
    check("锁定：整理时锁定的卡片一格不动", (lk["x"], lk["y"], lk["w"], lk["h"]) == (6, 3, 4, 3), lk)
    check("锁定：没锁的照常整理", next(m_ for m_ in out_ if m_["id"] == "a")["y"] == 0)
    check("锁定：locked 不是 true/false 就拒", G.check_module(dict(M("text", 0, 0, 3, 2, "z"), locked="yes")) != [])
    dc = one([M("chart", 0, 0, 12, 5, "c", file="assets/bar-filings.json"), M("conclusion", 0, 5, 12, 1, "k", text="收案逐年增长，只有 2021 年回落")])
    svc, rc = D.compose(dc, 0, DEMO, {})
    check("一句话结论：粗体画出来、不出问题", not rc["problems"] and "收案逐年增长" in svc and 'font-weight="700"' in svc, rc["problems"])
    _, re_ = D.compose(one([M("conclusion", 0, 0, 8, 1, "k")]), 0, DEMO, {})
    check("一句话结论：空着就拦", any("结论还是空的" in p_ for p_ in re_["problems"]))
    _, rl = D.compose(one([M("conclusion", 0, 0, 4, 1, "k", text="很长" * 60)]), 0, DEMO, {})
    check("一句话结论：写成一大段就拦（结论只该一两句）", any("一两句" in p_ for p_ in rl["problems"]), rl["problems"])
    with zipfile.ZipFile(p) as z:
        xa = z.read("ppt/slides/slide1.xml").decode("utf-8")
    check("讲课动画：放了 4 个编号的页，点 4 下，每下出编号、它指的元素、同号要点", xa.count('nodeType="clickEffect"') == 4
          and xa.count('nodeType="withEffect"') == 8, (xa.count("clickEffect"), xa.count("withEffect")))
    check("讲课动画：要点每一条各成一组，命名「要点 n」", re.findall(r'name="(要点 \d)"', xa) == ["要点 1", "要点 2", "要点 3", "要点 4"])
    ET.fromstring(xa)
    p2, r2 = E.build(d1, DEMO, os.path.join(td, "noanim.pptx"), cache={}, animate=False)
    with zipfile.ZipFile(p2) as z:
        check("讲课动画：关掉就不写", "<p:timing>" not in z.read("ppt/slides/slide1.xml").decode("utf-8") and r2["clicks"] == 0)
    tam = os.path.join(td, "anim-bad.pptx")
    with zipfile.ZipFile(p) as zi, zipfile.ZipFile(tam, "w", zipfile.ZIP_DEFLATED) as zo:
        for n_ in zi.namelist():
            data = zi.read(n_)
            if n_ == "ppt/slides/slide1.xml":
                data = re.sub(r'spid="\d+"', 'spid="99999"', data.decode("utf-8"), count=1).encode("utf-8")
            zo.writestr(n_, data)
    check("讲课动画：守卫能失败（动画指向不存在的形状就抓出来）", any("讲课动画" in b_ for b_ in E.audit(tam, d1, DEMO)))
    fig_ = {"w": 221.0, "h": 458.0, "pad": 16, "min1": 11.0}
    m3 = [M("title", 0, 0, 12, 1, "t", text="x"), M("figure", 0, 1, 6, 5, "f", file="a.svg"), M("bullets", 6, 1, 6, 5, "b", items=["y"])]
    o3, n3, fx3 = G.fit_figures(m3, {"f": fig_})
    fo, bo = next(m_ for m_ in o3 if m_["id"] == "f"), next(m_ for m_ in o3 if m_["id"] == "b")
    check("协调：整理时又窄又高的图，图卡收到最贴合的宽度，右边的卡片补上空位",
          fo["w"] == 4 and (bo["x"], bo["w"]) == (4, 8) and n3, (fo, bo, n3))
    o5, n5 = G.tidy([M("title", 0, 0, 12, 1, "t", text="x"), M("figure", 0, 1, 4, 5, "f", file="a.svg"),
                     M("statute", 7, 1, 5, 2, "s", law="x", text="y"), M("bullets", 7, 3, 5, 3, "b", items=["z"])])
    check("协调：整理时中间的空栏不留——右边的卡片左靠、再补齐", [(m_["id"], m_["x"], m_["w"]) for m_ in o5 if m_["id"] in "sb"]
          == [("s", 4, 8), ("b", 4, 8)], o5)
    o4, _ = G.tidy(o3, fx3)
    check("协调：收过的图卡，整理的「补齐」不会再把它撑宽", next(m_ for m_ in o4 if m_["id"] == "f")["w"] == 4)
    lk3 = [dict(m_, locked=True) if m_["id"] == "f" else m_ for m_ in m3]
    check("协调：锁定的图卡不收", G.fit_figures(lk3, {"f": fig_})[0][1]["w"] == 6)
    db = one([M("bullets", 0, 0, 6, 5, "b", items=["一条", "两条"])])
    sb, _ = D.compose(db, 0, DEMO, {})
    fsz = [float(v) for v in re.findall(r'font-size="([\d.]+)"', sb)]
    check("协调：要点少、卡片大时字放大（不全挤在左上角），但不超过 24px", fsz and 20 < max(fsz) <= 24, max(fsz) if fsz else None)
    seg = svgimport._clip_seg((0, 0), (10, 10), (2, 2, 6, 6))
    check("切片几何：线段裁到框内", seg == ((2.0, 2.0), (6.0, 6.0)), seg)
    poly = svgimport._clip_poly([(0, 0), (10, 0), (10, 10), (0, 10), (0, 0)], (5, 5, 20, 20))
    check("切片几何：方块裁成框内那一角", sorted(set(poly)) == [(5, 5), (5, 10), (10, 5), (10, 10)], poly)
    # 聚焦（0.8）：整张图还在，只点亮要讲的元素，其余往白里调
    fsel = ["node:#4", "edge:#10"]
    dfc = one([M("title", 0, 0, 12, 1, "t", text="实际施工人凭什么主张工程价款"),
               dict(M("figure", 0, 1, 12, 5, "f", file="assets/v3-relation.svg"), focus=fsel, pins=fsel)])
    D.validate(dfc, DEMO)
    svf, rf = D.compose(dfc, 0, DEMO, {})
    dnf = json.loads(json.dumps(dfc)); del dnf["slides"][0]["modules"][1]["focus"]
    svn, _ = D.compose(dnf, 0, DEMO, {})
    check("聚焦：不出问题，没点亮的元素都还在（整张图都在）", not rf["problems"] and "总承包人" in svf and "监理单位" in svf, rf["problems"])
    hot = lambda sv, k: re.findall(r'fill="(#[0-9A-Fa-f]{6})"[^>]*data-anchor="' + re.escape(k) + '"', sv)
    check("聚焦：点亮的元素颜色不变", hot(svf, "node:#4") == hot(svn, "node:#4") and hot(svf, "node:#4"))
    check("聚焦：没点亮的元素调浅", hot(svf, "node:#2") and hot(svf, "node:#2") != hot(svn, "node:#2")
          and all(c.upper() == svgimport.fade_hex(o).upper() for c, o in zip(hot(svf, "node:#2"), hot(svn, "node:#2"))))
    check("聚焦：图自带的大标题在页面有标题时不重复画（V3 标题没标 data-role 也认得）", "主体与客体并存" not in svf)
    for bad, why in (({"focus": "node:#1"}, "不是列表"), ({"focus": []}, "空列表")):
        try:
            D.validate(page(bad), DEMO)
            check(f"聚焦：校验拒绝（{why}）", False)
        except D.DeckError:
            check(f"聚焦：校验拒绝（{why}）", True)
    pf, rpf = E.build(dfc, DEMO, os.path.join(td, "focus.pptx"), cache={})
    check("聚焦：导出通过审计（调浅的颜色由图自己的颜色推出，审计认得）", os.path.isfile(pf) and rpf["slides"] == 1)
    dch = one([M("chart", 0, 0, 12, 6, "c", file="assets/bar-filings.json", focus=["id:y2021"])])
    svc2, rc2 = D.compose(dch, 0, DEMO, {})
    check("聚焦：V6 图表按 data-id 点亮一根柱子", not rc2["problems"] and svgimport.fade_hex("#991B1B").lower() in svc2.lower(), rc2["problems"])


def t_import_real(td):
    """开源前第二轮实跑（2026-10-09）：把 V1–V5 的真实出图放进 V7。撞到两件事：
    V4 简表的字坐标写在 <tspan> 上，导入时全堆到左上角、表格是空的，检查却不报；A4 页脚把内容撑满整页，字缩到 7px。"""
    import svgimport as S
    svg = ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 400 300" width="400" height="300">'
           '<text font-size="20" fill="#333333"><tspan x="10" y="40">甲方</tspan><tspan x="200" y="40">乙方</tspan></text>'
           '<text x="10" y="100" font-size="20" fill="#333333">第一行<tspan x="10" dy="1.5em">第二行</tspan></text>'
           '<text x="10" y="200" font-size="20" fill="#333333"><tspan>前半</tspan><tspan font-weight="700">后半</tspan></text></svg>')
    out, _ = S.fit(svg, (0, 0, 400, 300))
    pos = {t_: (float(x_), float(y_)) for x_, y_, t_ in re.findall(r'<text x="([\d.]+)" y="([\d.]+)"[^>]*>([^<]*)</text>', out)}
    check("导入：<text> 没写坐标、坐标在 tspan 上时按 tspan 放（V4 简表）",
          "甲方" in pos and "乙方" in pos and pos["乙方"][0] - pos["甲方"][0] > 150 and abs(pos["甲方"][1] - 40) < 1, pos)
    check("导入：tspan 用 dy 分行，em 按字号换算", "第二行" in pos and abs(pos["第二行"][1] - pos["第一行"][1] - 30) < 0.6, pos)
    check("导入：没写位置的 tspan 接在同一行（不另起一段叠在一起）", pos.get("前半后半") == (10.0, 200.0), pos)
    fx = os.path.join(HERE, "fixtures", "v4-brief.svg")       # V4 大事记表大师用它自己的测试案卷出的简表
    v4 = open(fx, encoding="utf-8").read()
    out4, r4 = S.fit(v4, (48, 160, 1184, 500), trim=True)
    nums = [(float(x_), round(float(y_)), t_.strip()) for x_, y_, t_ in
            re.findall(r'<text x="([\d.]+)" y="([\d.]+)"[^>]*>([^<]*)</text>', out4) if t_.strip() in ("1", "2", "3", "4", "5", "6")]
    col = min(x_ for x_, _, _ in nums) if nums else 0               # 序号那一列（最左）；「主要内容」里也有数字
    ys = sorted({y_ for x_, y_, _ in nums if abs(x_ - col) < 2})
    check("导入 V4 简表：六行序号各在各的行上（不再堆在一处）", len(ys) == 6, ys)
    check("导入 V4 简表：表里的内容都在", all(w_ in out4 for w_ in ("签订《商品房买卖合同》", "未办理产权登记", "主要内容")))
    check("导入 V4 简表：A4 页脚（第 n 页）略去，表格放大到字够大", "大事记简表</text>" not in out4
          and not r4["problems"] and r4["min_font"] >= 11 and any("页脚" in n_ for n_ in r4["notes"]), (r4["problems"], r4["min_font"]))
    keep_ = ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1000 1414" width="1000" height="1414">'
             '<rect x="100" y="100" width="800" height="1150" fill="#F2F3F5"/>'
             '<text x="500" y="1300" font-size="20" fill="#333333">第 1 页</text></svg>')
    o5, r5 = S.fit(keep_, (0, 0, 1000, 700), trim=True)
    check("导入：底部的字离上面内容不远时不当页脚（只略去隔着大段空白的「第 n 页」）", "第 1 页" in o5 and not r5.get("foot_dropped"))
    nope = keep_.replace('height="1150"', 'height="600"').replace("第 1 页", "本图由某某制作")
    o6, r6 = S.fit(nope, (0, 0, 1000, 700), trim=True)
    check("导入：底部的字没有「第 n 页」就不动", "本图由某某制作" in o6)
    v2 = open(os.path.join(HERE, "fixtures", "v2-numbered.svg"), encoding="utf-8").read()   # V2 时间轴大师的编号时间轴
    _, r8 = S.fit(v2, S.size_of(v2))
    an8 = r8["anchors"]
    check("导入 V2 编号时间轴：每个圆点都带上了旁边的事件卡片（卡片没标记，按同一水平带 + 引线认）",
          len(an8) == 20 and all(a_["box"][2] > 100 and a_["label"].strip() for a_ in an8),
          [(a_["key"], a_["box"][2], a_["label"][:8]) for a_ in an8[:3]])
    ks = [a_["key"] for a_ in an8[2:5]]
    o9, r9 = S.fit(v2, (0, 0, 1184, 560), trim=True, crop=ks)
    words = [a_["label"].split(" ")[-2] if len(a_["label"].split(" ")) > 2 else a_["label"].split(" ")[0] for a_ in an8[2:5]]
    check("导入 V2：切片选三件事，三张卡片和卡片里的字都在（原来只剩三个圆点、卡片是空边）",
          not r9["problems"] and all(len(w_) >= 4 and w_[:6] in o9 for w_ in words), (r9["problems"], words))
    slivers = [(w_, h_) for w_, h_ in re.findall(r'<rect x="[\d.-]+" y="[\d.-]+" width="([\d.]+)" height="([\d.]+)"', o9)
               if min(float(w_), float(h_)) < 12 and max(float(w_), float(h_)) > 30]      # 横的、竖的细条都算
    check("切片：没选、没整个落在切片里的卡片整个不画，不切成一条条细边", not slivers, slivers[:3])
    lab9 = {a_["key"]: a_["label"] for a_ in an8}
    others = [lab9[k_].split(" ")[1][:6] for k_ in lab9 if k_ not in ks and len(lab9[k_].split(" ")) > 2]
    cut = [w_ for w_ in others if w_ and w_ in o9]
    check("切片：没选的卡片要么整张在、要么不画——字不会被切断一半留在页面上", not cut, cut)
    d9 = one([M("figure", 0, 0, 12, 6, "f", file="v2.svg", pins=[an8[4]["key"]], **{"from": "V2 时间轴大师"})])
    os.makedirs(os.path.join(td, "v2p"), exist_ok=True)
    with open(os.path.join(td, "v2p", "v2.svg"), "w", encoding="utf-8") as f_:
        f_.write(v2)
    _, rp9 = D.compose(d9, 0, os.path.join(td, "v2p"), {})
    pin9 = rp9["modules"]["f"]["pins"][0]["box"]
    a9 = next(a_ for a_ in rp9["modules"]["f"]["anchors"] if a_["key"] == an8[4]["key"])
    dots = [a_["box"] for a_ in rp9["modules"]["f"]["anchors"]]
    card9 = a9.get("card")
    check("读图编号贴在 V2 事件卡片朝外的上角，不压在 V2 自己的编号圆点上",
          card9 and (abs(pin9[0] + 12 - card9[0]) < 8 or abs(pin9[0] + 12 - card9[0] - card9[2]) < 8)
          and abs(pin9[1] + 12 - card9[1]) < 8, (pin9, card9))
    ax = ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 600 400" width="600" height="400">'
          '<g data-role="axis"><text x="100" y="380" font-size="14" fill="#6B7280">2015</text></g>'
          '<g data-role="edges"><line x1="100" y1="50" x2="100" y2="370" stroke="#C3C9D2"/></g></svg>')
    _, r10 = S.fit(ax, (0, 0, 600, 400))
    check("坐标轴组里的年份不当成旁边那条线的名字（V2 甘特图的年份曾被认成关系标注）",
          [a_["label"] for a_ in r10["anchors"]] == [""], [a_["label"] for a_ in r10["anchors"]])
    outside = os.path.join(td, "outside.svg")
    with open(outside, "w", encoding="utf-8") as f_:
        f_.write('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 200 100" width="200" height="100">'
                 '<text x="10" y="50" font-size="20" fill="#333333">别的案子</text></svg>')
    proj_ = os.path.join(td, "proj-out")
    os.makedirs(proj_, exist_ok=True)
    for ref in ("../outside.svg", outside):
        dx = one([M("figure", 0, 0, 12, 6, "f", file=ref, **{"from": "V1 诉讼可视化重画"})])
        sx, rx_ = D.compose(dx, 0, proj_, {})
        check(f"课件文件夹外的素材不读（{'相对路径 ../' if ref.startswith('..') else '绝对路径'}）：免得把别的案件的图带进 PPT",
              "别的案子" not in sx and any("不在这个课件的文件夹里" in p_ for p_ in rx_["problems"]), rx_["problems"])
    # 箭头（开源前第三轮实跑：导出的 PPTX 有箭头，编辑器里的预览没有——关系图的方向是有法律意义的）
    da = demo()
    ia = next(i_ for i_, s_ in enumerate(da["slides"]) if any(m_.get("file") == "assets/v3-relation.svg" for m_ in s_["modules"]))
    sva, _ = D.compose(da, ia, DEMO, {}, paste=True)
    refs = set(re.findall(r'marker-end="url\(#([^)]+)\)"', sva))
    defs_ = set(re.findall(r'<marker id="([^"]+)"', sva))
    check("预览：关系图的箭头都有定义（预览里看得见，和导出的 PPTX 一致）", len(refs) >= 1 and refs <= defs_, (refs, defs_))
    pa, _ = E.build(da, DEMO, os.path.join(td, "arrows.pptx"), cache={}, pages=[ia])
    with zipfile.ZipFile(pa) as z:
        xa_ = z.read("ppt/slides/slide1.xml").decode("utf-8")
    check("导出：箭头定义不进 PPTX，11 个箭头照旧变成线端箭头", len(re.findall(r"<a:(?:tail|head)End", xa_)) == 11,
          len(re.findall(r"<a:(?:tail|head)End", xa_)))
    dfa = one([M("figure", 0, 0, 12, 6, "f", file="assets/v3-relation.svg", focus=["node:#1", "edge:#1"],
                 **{"from": "V3 法律关系图大师"})])
    svf_, _ = D.compose(dfa, 0, DEMO, {})
    check("聚焦：调浅的线配调浅的箭头（不是浅线配深箭头）", re.search(r'marker-end="url\(#[^)]+-f\)"', svf_)
          and re.search(r'<marker id="[^"]+-f"', svf_))
    ds_ = one([M("figure", 0, 0, 12, 6, "f", file="assets\\v3-relation.svg", **{"from": "V3 法律关系图大师"})])
    _, rs_ = D.compose(ds_, 0, DEMO, {})
    check("Windows 写法的路径（assets\\x.svg）在 Mac / Linux 上也认", not any("找不到" in p_ or "不在这个课件" in p_
          for p_ in rs_["problems"]) and rs_["modules"]["f"].get("anchors"), rs_["problems"])
    ds2 = one([M("figure", 0, 0, 12, 6, "f", file="..\\outside.svg", **{"from": "V3 法律关系图大师"})])
    _, rs2 = D.compose(ds2, 0, DEMO, {})
    check("反斜杠写的 ..\\ 照样出不了课件文件夹", any("不在这个课件的文件夹里" in p_ for p_ in rs2["problems"]), rs2["problems"])
    bad = os.path.join(td, "bad-inputs")
    os.makedirs(os.path.join(bad, "assets"), exist_ok=True)
    def load_err(text):
        with open(os.path.join(bad, "deck.json"), "w", encoding="utf-8") as f_:
            f_.write(text)
        try:
            D.load(os.path.join(bad, "deck.json"))
            return ""
        except D.DeckError as e:
            return str(e)
    e1 = load_err('{"schema": "mqc-nlv/deck@2", "title": ')
    check("坏输入：deck.json 不是合法 JSON，说第几行第几列，不甩英文异常", "第 1 行" in e1 and "Expecting" not in e1, e1)
    e2 = load_err('{"schema":"mqc-nlv/deck@2","title":"t","style":"奇川风","slides":"abc"}')
    check("坏输入：slides 不是列表，就说 slides 要是列表（不再把字符串拆成「第 3 页」）", "slides 要是页的列表" in e2 and "第 3 页" not in e2, e2)
    with open(os.path.join(bad, "assets", "bin.svg"), "wb") as f_:
        f_.write(bytes(range(128, 256)) * 8)
    with open(os.path.join(bad, "assets", "empty.svg"), "w", encoding="utf-8") as f_:
        f_.write('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 100" width="100" height="100"></svg>')
    for fn, want, why in (("bin.svg", "不是 SVG 文本", "乱码二进制"), ("empty.svg", "什么都没画", "空的 SVG")):
        _, rb = D.compose(one([M("figure", 0, 0, 12, 6, "f", file="assets/" + fn, **{"from": "V1 诉讼可视化重画"})]), 0, bad, {})
        check(f"坏输入：图是{why}就拦下，说清楚（不再当作「可以」或甩编码异常）", any(want in p_ for p_ in rb["problems"])
              and not any("codec" in p_ for p_ in rb["problems"]), rb["problems"])
    _, r7 = S.fit('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 2000 1000" width="2000" height="1000">'
                  '<text x="10" y="500" font-size="21.94" fill="#333333">小字</text></svg>', (0, 0, 1000, 500))
    check("字太小的提示不自相矛盾（10.97px 写成 10.9，不写成「11.0px（低于 11px）」）",
          any("10.9px" in p_ for p_ in r7["problems"]) and not any("11.0px" in p_ for p_ in r7["problems"]), r7["problems"])


def t_adapt(td):
    """作者定的（2026-10-07）：V1–V5 的幻灯片尺寸出图由 V7 自己兼容，V1–V6 原有定义一个字不改。先做 V5。"""
    import adapt as A
    src = os.path.join(SKILLS, "mqc-trial-confrontation-master", "examples", "matrix-input.json")
    if not os.path.exists(src):
        check("V5 兼容层（没有 V5 样例，跳过）", True)
        return
    proj = os.path.join(td, "v5proj")
    shutil.copytree(DEMO, proj)
    shutil.copy(src, os.path.join(proj, "assets", "v5-case.json"))
    d = demo()
    d["slides"] = [{"id": "v1", "modules": [
        M("title", 0, 0, 12, 1, "t", text="解除合同：设备是否符合约定参数"),
        dict(M("figure", 0, 1, 12, 4, "f", data="assets/v5-case.json"), **{"from": "V5 庭审对抗图大师", "part": 0}),
        M("conclusion", 0, 5, 12, 1, "k", text="两次验收均未通过")]}]
    D.validate(d, proj)
    svg, rep = D.compose(d, 0, proj, {})
    f = rep["modules"]["f"]
    check("V5 兼容层：一个要件按页面宽度出图，放进一页不出问题", not rep["problems"], rep["problems"])
    check("V5 兼容层：字号原样（≥ 11px）", any("最小字 1" in n and float(re.search(r"最小字 ([\d.]+)", n).group(1)) >= 11
                                        for n in f["notes"]), f["notes"])
    check("V5 兼容层：自动只取卡片（没有标题、索引表、图例）", len(f.get("anchors") or []) == 3
          and all(a["key"].startswith("card:P1") for a in f["anchors"]), [a["key"] for a in f.get("anchors") or []])
    check("V5 兼容层：报告里列出各要件，给编辑器选", len(f.get("parts") or []) == 3 and f["parts"][0].startswith("P1"), f.get("parts"))
    d2 = copy.deepcopy(d)
    d2["slides"][0]["modules"][1]["part"] = 7
    _, r2 = D.compose(d2, 0, proj, {})
    check("V5 兼容层：没有这个要件就说清楚", any("没有第 7 个要件" in p_ for p_ in r2["problems"]), r2["problems"])
    d3 = copy.deepcopy(d)
    d3["slides"][0]["modules"][1]["data"] = "assets/nope.json"
    _, r3 = D.compose(d3, 0, proj, {})
    check("V5 兼容层：数据文件找不到就说清楚", any("找不到 V5 数据文件" in p_ for p_ in r3["problems"]))
    bad = copy.deepcopy(d)
    bad["slides"][0]["modules"][1]["part"] = "一"
    check("V5 兼容层：part 不是整数就拒", raises(lambda: D.validate(bad, proj), D.DeckError, "part"))
    p, r = E.build(d, proj, os.path.join(td, "v5.pptx"), cache={})
    check("V5 兼容层：导出过回读审计", r["slides"] == 1 and E.audit(p, d, proj) == [])
    lib = studio.Project(proj).library()
    check("V5 兼容层：素材里认出 V5 数据（3 个要件）", any(x.get("kind") == "v5" and x["parts"] == 3 for x in lib["figures"]))


def t_cli(td):
    py = sys.executable
    cli = os.path.join(ROOT, "scripts", "build.py")
    r = subprocess.run([py, cli, DEMO, "--check"], capture_output=True, text=True, encoding="utf-8")
    check("CLI --check 示例：退出码 0", r.returncode == 0, r.stdout[-300:] + r.stderr[-300:])
    r = subprocess.run([py, cli, DEMO, "--check", "--style", "歸藏风"], capture_output=True, text=True, encoding="utf-8")
    check("CLI --check 风格不符：退出码 2", r.returncode == 2 and f"第{pno(demo(), 's8') + 1:3d} 页" in r.stdout, (r.returncode, r.stdout[-200:]))
    nd = os.path.join(td, "new-deck")
    r = subprocess.run([py, cli, nd, "--new", "--title", "测试课"], capture_output=True, text=True, encoding="utf-8")
    check("CLI --new 建新课件", r.returncode == 0 and os.path.exists(os.path.join(nd, "deck.json")), r.stderr[-300:])
    hint = re.findall(r"python3 (\"[^\"]+\"|\S+\.py)", r.stdout)
    check("CLI --new 打印下一步：打开编辑器与逐页检查的命令，脚本写真实路径，照抄就能跑（原来只说把图放进 assets/）",
          len(hint) == 2 and all(os.path.isfile(h.strip('"')) for h in hint) and "studio.py" in hint[0], r.stdout[-300:])
    r = subprocess.run([py, cli, nd], capture_output=True, text=True, encoding="utf-8")
    check("CLI 导出被拦：只说一次「未导出」（原来印成「未导出：有几页还没准备好，未导出：」）",
          r.returncode == 2 and r.stderr.count("未导出") == 1, r.stderr[:120])
    fut = os.path.join(td, "future-deck")
    shutil.copytree(DEMO, fut)
    dj = json.load(open(os.path.join(fut, "deck.json"), encoding="utf-8"))
    dj["schema"] = "mqc-nlv/deck@3"
    with open(os.path.join(fut, "deck.json"), "w", encoding="utf-8") as f_:
        json.dump(dj, f_, ensure_ascii=False)
    r = subprocess.run([py, cli, fut, "--check"], capture_output=True, text=True, encoding="utf-8")
    import firm_style as FS_
    try:
        FS_.check({"schema": "mqc-nlv/firm-style@2", "name": "x"})
        fs_msg = "没拦"
    except FS_.FirmStyleError as e:
        fs_msg = str(e)
    check("课件或律所风格是更新版本写的：明说「请把插件升级」（学自案图；原来只说 schema 要是 deck@2，像是文件坏了）",
          r.returncode == 1 and "升级" in r.stderr and "没有坏" in r.stderr and "升级" in fs_msg, (r.stderr[:100], fs_msg[:60]))


# ------------------------------------------------------------------ 编辑器的 JS 规则 = Python 规则
def t_js_parity(td):
    node = shutil.which("node")
    if not node:
        check("规则对照（未装 node，跳过）", True)
        return
    html_ = open(os.path.join(HERE, "..", "web", "studio.html"), encoding="utf-8").read()
    fn = re.search(r"function mergeDeck\(base, mine, theirs\) \{.*?\n\}\n", html_, re.S)
    check("编辑器的按页合并函数在（两个窗口同时改）", fn is not None)
    if fn:
        js = fn.group(0) + r"""
const P = (id, t) => ({id, modules: [{id: "t", type: "title", text: t}]});
const base = {title: "x", slides: [P("a", "A0"), P("b", "B0"), P("c", "C0"), P("d", "D0")]};
const mine = {title: "x", slides: [P("a", "A0"), P("n", "新页"), P("b", "B-本窗口"), P("c", "C0")]};          // 改 b、加 n、删 d
const theirs = {title: "y", slides: [P("a", "A-别处"), P("b", "B-别处"), P("c", "C0"), P("d", "D0")]};          // 改 a、b，改标题
const [m, both] = mergeDeck(base, mine, theirs);
console.log(JSON.stringify({ids: m.slides.map(s => s.id), t: m.slides.map(s => s.modules[0].text), title: m.title, both}));
"""
        # 脚本写进文件再跑，不走 -e：命令行参数在中文 Windows 上按区域编码走，脚本里的中文会被改坏
        jsf = os.path.join(td, "merge_parity.js")
        with open(jsf, "w", encoding="utf-8") as f_:
            f_.write(js)
        r_ = subprocess.run([node, jsf], capture_output=True, text=True, timeout=60, encoding="utf-8")
        try:
            o_ = json.loads(r_.stdout.strip().splitlines()[-1])
        except (ValueError, IndexError):
            o_ = {"err": r_.stderr[-300:]}
        check("按页合并：别处改的页用别处的、这里改的用这里的、这里加的页插回原位、这里删的页删掉",
              o_.get("ids") == ["a", "n", "b", "c"] and o_.get("t") == ["A-别处", "新页", "B-本窗口", "C0"], o_)
        check("按页合并：两边都改了同一页（第 2 页）就点名提示；整套标题这里没动，用别处的", o_.get("both") == [2]
              and o_.get("title") == "y", o_)
    rules = studio.grid_json()
    rnd = random.Random(7)
    cases = []
    families = ["stat", "step", "text", "bullets", "chart", "title"]
    for k in range(160):
        mods = []
        for i in range(rnd.randint(1, 7)):
            t = rnd.choice(families)
            mn = G.TYPES[t]["min"]
            w = rnd.randint(mn[0], max(mn[0], 6))
            h = rnd.randint(mn[1], max(mn[1], min(G.TYPES[t].get("maxh", G.ROWS), 3)))
            c = {"id": f"m{i}", "type": t, "w": w, "h": h}
            spot = G.free_spot(mods, w, h, rnd.randint(0, 11), rnd.randint(0, 5))
            if not spot:
                continue
            c["x"], c["y"] = spot
            if rnd.random() < 0.15:
                c["join"] = False
            if rnd.random() < 0.15 and t not in G.NO_ACCENT:
                c["tone"] = "accent"
            mods.append(c)
        cases.append({"mods": mods, "ptype": rnd.choice(list(G.TYPES)), "x0": rnd.randint(0, 11), "y0": rnd.randint(0, 5)})
    for t in G.TEMPLATES.values():
        cases.append({"mods": [dict(m, id=f"m{i}") for i, m in enumerate(t["mods"])], "ptype": "stat", "x0": 0, "y0": 0})
    py_out = []
    for c in cases:
        sides, _ = G.fusion(c["mods"])
        py_out.append({
            "fusion": {k: sorted(v, key="lrtb".index) for k, v in sides.items()},
            "rects": {m["id"]: [round(v, 6) for v in G.card_rect(m, sides[m["id"]])] for m in c["mods"]},
            "corners": {m["id"]: list(G.corners(sides[m["id"]], 12, 4)) for m in c["mods"]},
            "place": list(G.place(c["mods"], c["ptype"], c["x0"], c["y0"]) or []) or None,
        })
    js = os.path.join(td, "parity.js")
    with open(os.path.join(td, "cases.json"), "w", encoding="utf-8") as f:
        json.dump({"rules": rules, "cases": cases}, f, ensure_ascii=False)
    with open(js, "w", encoding="utf-8") as f:
        f.write("""
const G = require(%s); const fs = require("fs");
const {rules, cases} = JSON.parse(fs.readFileSync(%s, "utf8"));
const out = cases.map((c) => {
  const fu = G.fusion(rules, c.mods);
  const rects = {}, corners = {};
  for (const m of c.mods) { rects[m.id] = G.cardRect(rules, m, fu[m.id]).map((v) => Math.round(v * 1e6) / 1e6); corners[m.id] = G.corners(fu[m.id], 12, 4); }
  const p = G.place(rules, c.mods, c.ptype, c.x0, c.y0);
  return {fusion: fu, rects, corners, place: p};
});
/* 吸进一组：三卡并列去掉一张，拖到空出的位置附近，应按 4×5 落下并连上 */
const t = cases[cases.length - 1];
const three = [{id:"t",type:"title",x:0,y:0,w:12,h:1},{id:"a",type:"step",x:0,y:1,w:4,h:5},{id:"b",type:"step",x:4,y:1,w:4,h:5}];
const rs = G.runSnap(rules, three, {id:"n",type:"step",x:8,y:2,w:4,h:3}, null, 2);
const far = G.runSnap(rules, three, {id:"n",type:"step",x:8,y:4,w:4,h:2}, null, 1);
const noFam = G.runSnap(rules, three, {id:"n",type:"chart",x:8,y:1,w:4,h:4}, null, 2);
process.stdout.write(JSON.stringify({out, rs, far, noFam}));
""" % (json.dumps(GRID_JS), json.dumps(os.path.join(td, "cases.json"))))
    r = subprocess.run([node, js], capture_output=True, text=True, encoding="utf-8")
    if r.returncode != 0:
        check("规则对照：node 跑得起来", False, r.stderr[-400:])
        return
    res = json.loads(r.stdout)
    diff = [i for i, (a, b) in enumerate(zip(py_out, res["out"])) if a != b]
    check(f"规则对照：{len(cases)} 个版面上，编辑器的连组 / 外扩 / 圆角 / 落点与 Python 逐项一致", not diff,
          f"第 {diff[:3]} 个不同：py={py_out[diff[0]] if diff else ''} js={res['out'][diff[0]] if diff else ''}")
    rs = res["rs"]
    check("吸进一组：拖到同类卡片旁，按那一组的高落下并连上",
          rs and (rs["x"], rs["y"], rs["w"], rs["h"]) == (8, 1, 4, 5) and G.fusion(
              [M("step", 0, 1, 4, 5, "a"), M("step", 4, 1, 4, 5, "b"), dict(rs, id="n")])[0]["n"] == {"l"}, rs)
    check("吸进一组：超出吸附范围不吸", res["far"] is None, res["far"])
    check("吸进一组：不是同一家族不吸", res["noFam"] is None, res["noFam"])


# ------------------------------------------------------------------ 编辑器接口
def t_prov(td):
    """0.8.4：编辑器里悬停看出处——只读原模块写好的出处（V2 图旁边的溯源索引、V6 逐件记录），不编。"""
    pj = os.path.join(td, "prov")
    shutil.copytree(DEMO, pj)
    shutil.copy(os.path.join(HERE, "fixtures", "v2-numbered.svg"), os.path.join(pj, "assets", "tl.svg"))
    v6ex = os.path.join(ROOT, "..", "mqc-legal-chart-master", "examples", "dataset-causes.json")
    shutil.copy(v6ex, os.path.join(pj, "assets", "causes.json"))
    svg_ = open(os.path.join(pj, "assets", "tl.svg"), encoding="utf-8").read()
    ids_ = sorted(set(re.findall(r'data-id="([^"]+)"', svg_)))
    rows_ = [{"no": i, "head": f"事项{i}", "file": "判决书.pdf", "locator": f"句{i}", "quote": f"第{i}句原文（虚构）"} for i in ids_]
    rows_.append({"no": "999", "head": "图上没有的事项", "file": "x", "locator": "句0", "quote": "不该出现"})
    with open(os.path.join(pj, "assets", "tl-trace.json"), "w", encoding="utf-8") as f_:
        json.dump({"title": "示例・溯源索引", "rows": rows_}, f_, ensure_ascii=False)
    d = D.load(os.path.join(pj, "deck.json"))
    d["slides"][:0] = [
        {"id": "pa", "modules": [{"id": "m1", "type": "figure", "x": 0, "y": 0, "w": 12, "h": 6, "file": "assets/tl.svg", "from": "V2 时间轴大师"}]},
        {"id": "pb", "modules": [{"id": "m1", "type": "chart", "x": 0, "y": 0, "w": 8, "h": 6, "file": "assets/causes.json"}]},
        {"id": "pc", "modules": [{"id": "m1", "type": "chart", "x": 0, "y": 0, "w": 8, "h": 6, "file": "assets/bar-filings.json"},
                                 {"id": "m2", "type": "figure", "x": 8, "y": 0, "w": 4, "h": 6, "file": "assets/v3-relation.svg", "from": "V3"}]}]
    infos = [D.compose(d, i, pj, {})[1]["modules"] for i in range(3)]
    pa, pb, pc = infos[0]["m1"], infos[1]["m1"], infos[2]
    anc_a = {a["key"] for a in pa.get("anchors") or []}
    check("V2 的图旁边有溯源索引：图里每件事都带出处（材料、第几句、原文），图上没有的那条不带进来",
          set((pa.get("prov") or {})) == anc_a and len(anc_a) == len(ids_) and "id:999" not in (pa.get("prov") or {})
          and pa["prov"]["id:1"]["quote"] == "第1句原文（虚构）" and pa["prov"]["id:1"]["src"] == "判决书.pdf 句1",
          (len(pa.get("prov") or {}), len(anc_a)))
    pv = pb.get("prov") or {}
    check("V6 用逐件记录出的图：每一项带「来自哪几条记录、编号」，行号与 V6 的说法一样从 0 数，编号不放进引文",
          set(pv) == {a["key"] for a in pb.get("anchors") or []} and pv.get("id:民间借贷", {}).get("refs", "").startswith("编号：示例-01")
          and "行号 0、" in pv["id:民间借贷"]["src"] and not pv["id:民间借贷"]["quote"], list(pv.items())[:1])
    check("没有逐项出处的图（普通 V6 图表、V3 关系图）：不显示出处，不编", not pc["m1"].get("prov") and not pc["m2"].get("prov"))
    import studio as ST
    lib = ST.Project(pj).library()
    check("图旁边的出处记录（-trace.json）不当成图表列进「已有的图」（原来会列出一张空图表）",
          not any("trace" in c["file"] for c in lib["charts"]) and any(c["file"] == "assets/causes.json" for c in lib["charts"]))
    html_ = open(os.path.join(HERE, "..", "web", "studio.html"), encoding="utf-8").read()
    check("编辑器：鼠标停在图里的元素上浮出出处，选中图卡右侧列出全部",
          'id="provTip"' in html_ and "function showProv(e)" in html_ and "readSection(m) + provSection(m)" in html_)
    # 出处写进演讲者备注（要了才写）
    d2 = D.load(os.path.join(pj, "deck.json"))
    d2["slides"] = [{"id": "pa", "modules": [{"id": "m1", "type": "figure", "x": 0, "y": 0, "w": 12, "h": 6, "file": "assets/tl.svg",
                                              "from": "V2 时间轴大师", "crop": ["id:1", "id:2", "id:3", "id:4"]}]},
                    {"id": "pb", "modules": [{"id": "m1", "type": "chart", "x": 0, "y": 0, "w": 8, "h": 6, "file": "assets/causes.json"}]},
                    {"id": "pc", "modules": [{"id": "m1", "type": "chart", "x": 0, "y": 0, "w": 8, "h": 6, "file": "assets/bar-filings.json"}]}]
    with open(os.path.join(pj, "deck.json"), "w", encoding="utf-8") as f_:
        json.dump(d2, f_, ensure_ascii=False)
    import export as EX_
    import notes as NT_
    pn, rn = EX_.build(D.load(os.path.join(pj, "deck.json")), pj, os.path.join(td, "notes.pptx"), notes=True)
    pp, _ = EX_.build(D.load(os.path.join(pj, "deck.json")), pj, os.path.join(td, "plain.pptx"))
    zn, zp = zipfile.ZipFile(pn), zipfile.ZipFile(pp)
    ns_ = sorted(n for n in zn.namelist() if n.startswith("ppt/notesSlides/notesSlide"))
    t1 = zn.read("ppt/notesSlides/notesSlide1.xml").decode("utf-8") if "ppt/notesSlides/notesSlide1.xml" in zn.namelist() else ""
    check("出处写进演讲者备注：只写有出处的页（V2 时间轴、V6 逐件记录），原文摘录照录；导出后回读核过备注母版、关系、内容类型",
          rn["notes"] == 2 and ns_ == ["ppt/notesSlides/notesSlide1.xml", "ppt/notesSlides/notesSlide2.xml"]
          and "第1句原文（虚构）" in t1 and "notesMasterIdLst" in zn.read("ppt/presentation.xml").decode("utf-8"), (rn.get("notes"), ns_))
    check("不要备注时一个字节都不加：没有备注母版、没有备注页，幻灯片本身与要备注时逐字相同",
          not any("notes" in n for n in zp.namelist())
          and all(zp.read(n) == zn.read(n) for n in zp.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)))
    check("备注的回读能失败：要写的备注页不在，就报出来", any("对不上" in x for x in NT_.audit(zp, {1: ["x"]})))
    # V3 关系图：主体、关系按 data-ref 对上出处；没写出处的明说「没写」，不进备注
    v3svg = ('<svg xmlns="http://www.w3.org/2000/svg" width="600" height="300" viewBox="0 0 600 300">'
             '<rect width="600" height="300" fill="#FFFFFF"/>'
             '<g data-role="node" data-ref="A"><rect data-kind="node" x="40" y="100" width="160" height="80" rx="12" fill="#F7F5F2" stroke="#3D3D3D"/>'
             '<text x="120" y="145" font-size="22" text-anchor="middle" fill="#1A1A1A">甲公司</text></g>'
             '<g data-role="node" data-ref="B"><rect data-kind="node" x="400" y="100" width="160" height="80" rx="12" fill="#F7F5F2" stroke="#3D3D3D"/>'
             '<text x="480" y="145" font-size="22" text-anchor="middle" fill="#1A1A1A">乙公司</text></g>'
             '<path data-kind="edge" data-ref="A→B" d="M200 140 L400 140" fill="none" stroke="#3D3D3D" stroke-width="2"/>'
             '<text x="300" y="130" font-size="20" text-anchor="middle" fill="#1A1A1A">借款</text></svg>')
    with open(os.path.join(pj, "assets", "rel.svg"), "w", encoding="utf-8") as f_:
        f_.write(v3svg)
    with open(os.path.join(pj, "assets", "rel-trace.json"), "w", encoding="utf-8") as f_:
        json.dump({"title": "出处索引", "rows": [{"no": "A", "head": "甲公司", "file": "", "locator": "判决书第2页", "quote": ""},
                                              {"no": "B", "head": "乙公司", "file": "", "locator": "", "quote": ""},
                                              {"no": "A→B", "head": "甲公司 → 乙公司：借款", "file": "", "locator": "借款合同第1条", "quote": ""}]},
                  f_, ensure_ascii=False)
    d3 = D.load(os.path.join(pj, "deck.json"))
    d3["slides"] = [{"id": "pr", "modules": [{"id": "m1", "type": "figure", "x": 0, "y": 0, "w": 12, "h": 6, "file": "assets/rel.svg", "from": "V3 法律关系图大师"}]}]
    rep3 = D.compose(d3, 0, pj, {})[1]["modules"]["m1"]
    pv3 = rep3.get("prov") or {}
    byref = {a.get("ref"): a["key"] for a in rep3.get("anchors") or []}
    check("V3 关系图：主体、关系按图里的 data-ref 对上出处（元素键不变，已放的编号不受影响）；没写出处的明说没写",
          pv3.get(byref.get("A"), {}).get("src") == "判决书第2页" and pv3.get(byref.get("A→B"), {}).get("src") == "借款合同第1条"
          and byref.get("A→B", "").startswith("edge:") and pv3.get(byref.get("B"), {}).get("missing"), (byref, pv3))
    nl = NT_.lines_for({"modules": {"m1": rep3}})
    check("演讲者备注只写有出处的（没写出处的不往备注里写「没写」）", any("判决书第2页" in x for x in nl) and not any(x.startswith("· 乙公司：") for x in nl), nl)


def t_detach(td):
    """0.9.0：由 AI 助手（Agent）代开编辑器。模拟 Agent 的命令工具：新进程组跑命令、管道收输出，命令结束后掐整个进程组。"""
    import signal
    proj = os.path.join(td, "detach")
    shutil.copytree(DEMO, proj)
    cli = os.path.join(ROOT, "scripts", "studio.py")
    op = urllib.request.build_opener(urllib.request.ProxyHandler({}))

    def alive(url):
        try:
            with op.open(url + "api/ping", timeout=3) as r:
                return json.loads(r.read().decode("utf-8"))
        except Exception:
            return None

    def agent(*args, env=None):
        t0 = time.time()
        kw = {"start_new_session": True} if os.name != "nt" else {}
        p = subprocess.Popen([sys.executable, cli, proj, *args], stdout=subprocess.PIPE, stderr=subprocess.PIPE, env=env, **kw)
        try:
            out, err = p.communicate(timeout=60)
        except subprocess.TimeoutExpired:
            p.kill()
            return 99, "", "命令 60 秒没返回（输出管道被占着）", 60
        if os.name != "nt":
            try:
                os.killpg(p.pid, signal.SIGTERM)
            except (ProcessLookupError, PermissionError):
                pass
            time.sleep(1.0)
        return p.returncode, out.decode("utf-8", "replace"), err.decode("utf-8", "replace"), time.time() - t0

    rc, out, err, dt = agent("--detach", "--no-open", "--port", "18790")
    m = re.search(r"http://127\.0\.0\.1:\d+/", out)
    url = m.group(0) if m else ""
    check("--detach：返回 0 并打印网址", rc == 0 and url, (rc, out, err))
    check("--detach：命令马上返回，不占 Agent 的输出管道（Agent 等管道关上才算命令结束）", dt < 20, dt)
    info = alive(url) if url else None
    check("--detach：Agent 掐掉自己的进程组后，编辑器还在，开的是这套课件",
          info and os.path.normcase(info["dir"]) == os.path.normcase(os.path.abspath(proj)), info)
    rc2, out2, err2, _ = agent("--detach", "--no-open")
    check("--detach 再开一次：认出已开着、给原网址，不开第二个", rc2 == 0 and "已经开着" in out2 and url in out2, (out2, err2))
    bad = dict(os.environ, HTTP_PROXY="http://127.0.0.1:9", http_proxy="http://127.0.0.1:9")
    rc3, out3, _, _ = agent("--detach", "--no-open", env=bad)
    check("--detach：设了 HTTP 代理也照样认（连本机不走代理）", rc3 == 0 and "已经开着" in out3, out3)
    rc4, out4, err4, _ = agent("--stop")
    check("--stop：关掉，状态文件删掉", rc4 == 0 and "已关掉" in out4 and alive(url) is None
          and not os.path.exists(os.path.join(proj, studio.STATE)), (out4, err4))
    rc5, out5, _, _ = agent("--stop")
    check("--stop：没开着时不报错", rc5 == 0 and "没有独立运行" in out5, out5)
    with open(os.path.join(proj, studio.STATE), "w", encoding="utf-8") as f:
        json.dump({"url": "http://127.0.0.1:1/", "pid": 1}, f)
    rc6, out6, _, _ = agent("--stop")
    check("--stop：旧状态文件里的进程号不信（不去结束进程 1）", rc6 == 0 and "已关掉" in out6, out6)
    dev = os.path.join(td, "netdev")
    head = "Inter-|   Receive\n face |bytes\n"
    with open(dev, "w", encoding="ascii") as f:
        f.write(head + "    lo: 0 0\n")
    one = studio._net_isolated(dev)
    with open(dev, "w", encoding="ascii") as f:
        f.write(head + "    lo: 0 0\n  eth0: 0 0\n")
    two = studio._net_isolated(dev)
    check("认得出关网沙箱（只剩回环口）；有网卡时不误报",
          (one, two) == ((True, False) if sys.platform.startswith("linux") else (False, False)), (one, two))


def t_lib090(td):
    """0.9.0：图表库（V6 场景库放进 V7）、示例数据拦导出、标题不再伪加粗、AI 做完就停（缪律师实测的几处问题）。"""
    import zipfile
    import ai as A
    from spec import CHARTS
    L = studio.chart_library()
    rs = L["recipes"]
    check("图表库：列出 V6 场景库里现在就能画的场景（100 个以上）", len(rs) >= 100, len(rs))
    check("图表库：每个场景的画法都是 V6 已就绪的、并有示例可放", all(r["form"] in CHARTS for r in rs)
          and {r["form"] for r in rs} <= {f["code"] for f in L["forms"]})
    check("图表库：漏斗是备用画法，排在最后（作者定的）",
          [r["form"] for r in rs if r["form"] == "funnel"] == [r["form"] for r in rs[-sum(1 for r in rs if r["form"] == "funnel"):]])
    check("图表库：同一种画法不扎堆（前 10 个里至少 8 种画法）", len({r["form"] for r in rs[:10]}) >= 8, [r["form"] for r in rs[:10]])
    check("图表库：场景都带名字、要回答的问题、适用场合", all(r["name"] and r["question"] and r["scenes"] for r in rs))
    d = demo()
    bad = [f["code"] for f in L["forms"] if not (studio.chart_thumb(d, DEMO, f["code"], {}) or "").startswith("<svg")]
    check("图表库缩略图：每种画法都画得出来（发射扇、放射汇聚放大画布再画）", not bad, bad)
    # 示例数据：放进来能看，导不出去
    x = copy.deepcopy(d)
    sp = json.load(open(os.path.join(studio.V6_EXAMPLES, studio.form_examples()["bar"][0] + ".json"), encoding="utf-8"))
    sp.pop("annotations", None)
    x["slides"][3]["modules"] = [m for m in x["slides"][3]["modules"] if m["type"] != "chart"] + \
        [{"id": "mx", "type": "chart", "x": 0, "y": 1, "w": 8, "h": 5, "spec": sp, "sample": True, "recipe": rs[0]["id"]}]
    x["slides"][3]["modules"] = [m for m in x["slides"][3]["modules"] if m["type"] in ("title", "chart")]
    D.validate(x, DEMO)
    _, rep = D.compose(x, 3, DEMO, {})
    check("图表库放进来的示例数据：照常画出来，但导出前被拦下", any("示例数据" in p_ for p_ in rep["problems"]), rep["problems"])
    z = copy.deepcopy(x)
    z["slides"][3]["modules"].append({"id": "my", "type": "bullets", "x": 8, "y": 1, "w": 4, "h": 5, "items": ["a"], "sample": True})
    check("示例记号只给图表卡", raises(lambda: D.validate(z, DEMO), D.DeckError, "图表库"))
    proj = os.path.join(td, "lib090")
    shutil.copytree(DEMO, proj)
    path = os.path.join(proj, "deck.json")
    dk = D.load(path)
    dk["slides"][3]["modules"] = x["slides"][3]["modules"]
    with open(path, "w", encoding="utf-8") as f:
        json.dump(dk, f, ensure_ascii=False)
    sid = dk["slides"][3]["id"]
    r = A.add(proj, "tool", sid, "mx", "V6", "用我的数据")
    ctx = A.context(proj, r)
    check("AI 上下文带上图表库场景（要回答的问题、怎么换数据）", ctx.get("recipe", {}).get("id") == rs[0]["id"]
          and "示例数据" in ctx["recipe"]["how"], ctx.get("recipe"))
    check("AI 上下文：做完就停，不要一遍遍 wait（原来写「回到 wait 等下一条」，AI 一直空等不收尾）",
          "做完就停" in ctx["answer"]["then"] and "回到 wait" not in ctx["answer"]["then"], ctx["answer"]["then"])
    sp2 = copy.deepcopy(sp)
    sp2["title"] = sp2["title"]
    A.apply(proj, r["id"], {"module": {"spec": sp2}})
    m = next(q for q in page(D.load(path), sid)["modules"] if q["id"] == "mx")
    check("AI 交回用户的数据：示例记号去掉，场景记号留着", "sample" not in m and m.get("recipe") == rs[0]["id"], m.keys())
    # AI 是否在等：编辑器据此提示
    check("AI 没在 wait 时，编辑器知道", not A.listening(proj))
    pw = subprocess.Popen([sys.executable, os.path.join(ROOT, "scripts", "ai.py"), proj, "wait", "--timeout", "4"],
                          stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    seen = False
    for _ in range(30):
        time.sleep(0.1)
        if A.listening(proj):
            seen = True
            break
    out, _ = pw.communicate(timeout=30)
    check("AI 在 wait 时，编辑器知道；等完了又知道", seen and not A.listening(proj) and pw.returncode == 4)
    check("wait 等不到：告诉 AI 收尾，不要再等", "不要再 wait" in out.decode("utf-8", "replace"), out.decode("utf-8", "replace")[-200:])
    # 标题：方正小标宋本来的字重，不加粗、不描边
    svg, _ = D.compose(d, 3, DEMO, {})
    heavy = [t for t in re.findall(r"<text\b[^>]*>", svg) if D._is_song(re.search(r'font-family="([^"]*)"', t).group(1))
             and re.search(r'font-weight="(600|700|bold)"', t)]
    check("宋体标题不再加粗、不再描边（页标题、卡片标题、V6 图的标题）", not heavy and 'stroke-width="0.3"' not in svg, heavy[:2])
    check("宋体标题打了标题记号，导出 PPT 时据此选字", svg.count('data-title="1"') >= 3)
    f_svg, _ = D.compose(d, next(i for i, s_ in enumerate(d["slides"]) if any(m["type"] == "figure" for m in s_["modules"])), DEMO, {})
    f_heavy = [t for t in re.findall(r"<text\b[^>]*>", f_svg) if re.search(r'font-family="([^"]*)"', t)
               and D._is_song(re.search(r'font-family="([^"]*)"', t).group(1)) and re.search(r'font-weight="(600|700|bold)"', t)]
    check("V1–V5 的图放进来：图里的宋体加粗标题同样改成本来字重（旧图也照改）", not f_heavy and 'data-title="1"' in f_svg, f_heavy[:2])
    q_svg, _ = D.compose(d, next(i for i, s_ in enumerate(d["slides"]) if any(m["type"] == "quote" for m in s_["modules"])), DEMO, {})
    check("本来就是常规字重的宋体（金句）不动：不加标题记号", 'data-title="1"' not in q_svg)
    check("黑体的粗字不动（一句话结论、编号）", D.unbold_titles('<text font-family="\'PingFang SC\',sans-serif" font-weight="700">x</text>')
          == '<text font-family="\'PingFang SC\',sans-serif" font-weight="700">x</text>')
    runs = {}
    for tf in (None, "方正小标宋简体", "<x>"):
        p_, _ = E.build(d, DEMO, os.path.join(td, f"tf{len(tf or '')}.pptx"), cache={}, pages=[3], title_face=tf)
        zz = zipfile.ZipFile(p_)
        xml = "".join(zz.read(n).decode("utf-8") for n in zz.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n))
        runs[tf] = re.findall(r'<a:rPr[^>]*?b="(\d)"[^>]*>(?:(?!</a:rPr>).)*?typeface="([^"]*)".*?<a:t>([^<]*)</a:t>', xml, flags=re.S)
    t0 = [r_ for r_ in runs[None] if r_[2] == "收案逐年增长"]
    t1 = [r_ for r_ in runs["方正小标宋简体"] if r_[2] == "收案逐年增长"]
    check("PPT 标题：没确认装了小标宋，照旧宋体加粗（别人电脑上不变样）", t0 and t0[0][:2] == ("1", "宋体"), t0)
    check("PPT 标题：确认装了小标宋，用它、不加粗", t1 and t1[0][:2] == ("0", "方正小标宋简体"), t1)
    check("PPT：标题以外的字不受影响", [r_ for r_ in runs[None] if r_[2] != "收案逐年增长" and r_[1] != "宋体"]
          == [r_ for r_ in runs["方正小标宋简体"] if r_[2] != "收案逐年增长" and r_[1] not in ("宋体", "方正小标宋简体")])
    tx = [r_ for r_ in runs["<x>"] if r_[2] == "收案逐年增长"]
    check("标题字只认白名单里的名字：别的名字一律当没给", tx and tx[0][:2] == ("1", "宋体"), tx)
    html_ = open(os.path.join(HERE, "..", "web", "studio.html"), encoding="utf-8").read()
    check("编辑器：图表库放在侧栏第一位，按场景、画法、关键词找，点一下放进这一页",
          html_.index('data-tab="lib"') < html_.index('data-tab="pick"') and "function placeRecipe" in html_
          and "data-lint" in html_ and 'id="libScene"' in html_ and 'id="libQ"' in html_)
    check("图表库：先问「这一页想让人看懂什么」，按律师要说明的事分（占几成、钱怎么算、流向哪里……），卡片写问题、例子与核对规则",
          all(r["intent"] for r in rs) and len(L["intents"]) >= 8 and "这一页，想让人看懂什么" in html_ and "核对：" in html_)
    check("编辑器：版式、卡片画成真的小页面和样张（不再是灰方块拼线框）", "function tplThumb" in html_ and "const PREVIEW" in html_)
    check("编辑器：正在打字时 AI 改了课件也先不刷新", "const typing = /INPUT|TEXTAREA/" in html_)
    check("编辑器：AI 没在等时提示回对话里说一声", "回对话里说「处理课件请求」" in html_ and "v.listen" in html_)
    check("编辑器：装了小标宋才让 PPT 用它（量拉丁字母宽度认字体）", "HAS_XBS" in html_ and '"&tf=1"' in html_)


def t_studio(td):
    proj = os.path.join(td, "proj")
    shutil.copytree(DEMO, proj)
    P = studio.Project(proj)
    port = studio.free_port(18760)
    srv = studio.ThreadingHTTPServer(("127.0.0.1", port), studio.make_handler(P))
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{port}"

    def get(path):
        try:
            with urllib.request.urlopen(base + path, timeout=60) as r:
                return r.status, r.read()
        except urllib.error.HTTPError as e:
            return e.code, e.read()

    def post(path, body=b"", headers=None):
        req = urllib.request.Request(base + path, data=body, method="POST", headers=headers or {})
        try:
            with urllib.request.urlopen(req, timeout=120) as r:
                return r.status, json.loads(r.read() or b"{}")
        except urllib.error.HTTPError as e:
            return e.code, json.loads(e.read() or b"{}")

    q = urllib.parse.quote
    try:
        c, b = get("/")
        check("编辑器首页", c == 200 and b"<html" in b.lower())
        c, b = get("/grid.js")
        check("编辑器的规则脚本", c == 200 and b"fusion" in b)
        c, b = get("/api/state")
        st = json.loads(b)
        check("/api/state：课件 + 网格规则", c == 200 and st["deck"]["title"] and st["rules"]["grid"]["COLS"] == G.COLS
              and set(st["rules"]["templates"]) == set(G.TEMPLATES))
        c, b = get("/api/library")
        lib = json.loads(b)
        check("/api/library：素材图表、V6 示例、V1–V5 图都列出", c == 200 and len(lib["charts"]) == 4 and lib["builtin"]
              and lib["figures"])
        c, b = get("/api/slide/0.svg")
        check("/api/slide 预览", c == 200 and b.startswith(b"<svg"))
        _bc, _bb = get("/brand/qichuan-avatar.png")
        _bc2, _bb2 = get("/brand/qichuan-favicon.png")
        _bc3, _bb3 = get("/brand/../studio.py")
        _bc4, _bpage = get("/")
        check("品牌：侧栏左上角是作者头像 + 「新诉讼可视化」，标签页小图标是头像；别的路径不放行",
              _bc == 200 and _bb.startswith(b"\x89PNG") and _bc2 == 200 and _bb2.startswith(b"\x89PNG") and _bc3 != 200
              and "/brand/qichuan-avatar.png" in _bpage.decode("utf-8") and '<span class="bm-name">新诉讼可视化</span>' in _bpage.decode("utf-8")
              and 'rel="icon"' in _bpage.decode("utf-8"))
        check("预览与导出同一个拼页函数", b.decode("utf-8") == D.compose(P.load(), 0, proj, {})[0])
        c, _ = get("/api/slide/99.svg")
        check("没有这一页：404", c == 404)
        c, b = get("/api/report")
        check("/api/report", c == 200 and len(json.loads(b)) == len(P.load()["slides"]))
        # 0.8.4：编辑器只取、只查变了的页（原来改一个字，整套缩略图重拉、整套检查重拼）
        dk = P.load()
        sid = dk["slides"][3]["id"]
        c, b = get(f"/api/slide/0.svg?s={q(sid)}&v=7")
        check("按页 id 取画面：要的路上页被挪了位置，拿到的仍是这一页（页码是它现在的位置）",
              c == 200 and b.decode("utf-8") == D.compose(dk, 3, proj, {})[0])
        c, _ = get("/api/slide/0.svg?s=nosuch")
        check("按页 id 取画面：没有这一页就 404，不拿成别的页", c == 404)
        c, b = get("/api/report?ids=" + q(f"{sid},{dk['slides'][0]['id']},nosuch"))
        part = json.loads(b) if c == 200 else {}
        part = part.get("by_id", {}) if isinstance(part, dict) else {}
        full = json.loads(get("/api/report")[1])
        check("只查这几页：与整套检查里的那几页逐字相同，不认识的 id 不报",
              set(part) == {sid, dk["slides"][0]["id"]} and part[sid] == full[3] and part[dk["slides"][0]["id"]] == full[0])
        html_ = open(os.path.join(HERE, "..", "web", "studio.html"), encoding="utf-8").read()
        check("编辑器记住每页画面、只在课件存好之后去取（存的路上取，会把改之前的样子记在新钥匙下）",
              "const synced = () => !S.dirty && !S.saving && JSON.stringify(S.deck) === S.synced" in html_
              and "if (!synced()) return Promise.resolve" in html_ and '<img alt="" src="/api/slide/' not in html_)
        check("编辑器能缩放与平移（学自案图的画布）：右下角的缩放按钮、Ctrl / ⌘ 滚轮以光标为中心缩放、下限是适应窗口、"
              "放大后页面在中间这一块里滚动（不把整个界面撑高）",
              'id="zoomGroup"' in html_ and "addEventListener(\"wheel\"" in html_ and "S.zoom = Math.max(1," in html_
              and "grid-template-rows:minmax(0,100vh)" in html_ and ".stagewrap { flex:1; min-width:0; min-height:0;" in html_
              and "overflow:auto" in html_.split(".stagewrap {")[1].split("}")[0])
        c, r_ = post("/api/upload?name=" + q("tl-trace.json"), json.dumps({"title": "x", "rows": [{"no": "1", "head": "签约"}]}).encode("utf-8"))
        c2, r2_ = post("/api/upload?name=" + q("bad-trace.json"), b'{"chart": "bar"}')
        check("编辑器里拖进 V2 / V3 图旁边的出处记录（-trace.json）：收下，不再被当成写坏的图表说明拒掉；不像出处记录的照样拒",
              r_.get("ok") and r_.get("sidecar") and os.path.isfile(os.path.join(proj, "assets", "tl-trace.json"))
              and not r2_.get("ok"), (r_, r2_))
        c, _ = get("/api/asset?f=" + q("assets/bar-filings.json"))
        check("读素材 .json", c == 200)
        for bad in ["deck.json", "../../../etc/passwd", "assets/../deck.json", "assets/v3-relation.svg"]:
            c, _ = get("/api/asset?f=" + q(bad))
            check(f"读素材越界拒：{bad}", c == 404)
        c, b = post("/api/upload?name=" + q("新图.json"), json.dumps({"chart": "bar", "title": "t", "data": []}).encode("utf-8"))
        check("上传 V6 说明", c == 200 and b["ok"] and os.path.exists(os.path.join(proj, "assets", "新图.json")))
        for name, body in [("x.exe", b"MZ"), ("x.pptx", b"PK"), ("bad.json", b'{"a":1}')]:
            c, b = post("/api/upload?name=" + q(name), body)
            check(f"上传拒：{name}", c == 200 and not b["ok"])
        post("/api/upload?name=" + q("../escape.json"), b'{"chart":"bar"}')
        check("上传路径穿越：只落在 assets/", not os.path.exists(os.path.join(proj, "escape.json"))
              and not os.path.exists(os.path.join(td, "escape.json")))
        c, b = post("/api/deck", json.dumps({"schema": "x"}).encode("utf-8"))
        check("保存坏的课件说明：拒，原文件不动", c == 200 and not b["ok"] and P.load()["schema"] == D.SCHEMA)
        d = P.load()
        page(d, "s3")["modules"][3]["y"] = 2
        c, b = post("/api/deck", json.dumps(d, ensure_ascii=False).encode("utf-8"))
        check("保存重叠的版面：拒", not b.get("ok") and "重叠" in b.get("error", ""))
        c, b = post("/api/tidy", json.dumps({"modules": [M("stat", 0, 3, 3, 2, "a", value="1")]}).encode("utf-8"))
        check("/api/tidy", b.get("ok") and b["modules"][0]["y"] == 0)
        c, b = post("/api/template", json.dumps({"modules": [], "name": "kpi"}).encode("utf-8"))
        check("/api/template：新页按组合给出空格子", b.get("ok") and len(b["modules"]) == 6 and not b["left"])
        c, b = post("/api/template", json.dumps({"modules": [], "name": "nope"}).encode("utf-8"))
        check("/api/template：没有这个组合就拒", not b.get("ok"))
        d = P.load()
        d["slides"].append({"id": "s99", "modules": [M("title", 0, 0, 12, 1, "m1", text="加的页")]})
        c, b = post("/api/deck", json.dumps(d, ensure_ascii=False).encode("utf-8"))
        check("保存课件说明", b.get("ok") and len(P.load()["slides"]) == len(d["slides"]))
        c, b = post("/api/export")
        check("编辑器导出", b.get("ok") and b.get("slides") == len(d["slides"]), b)
        c, data = get("/api/download?f=" + q(b.get("file", "")))
        check("下载 PPTX", c == 200 and data[:2] == b"PK")
        c, _ = get("/api/download?f=" + q("../deck.json"))
        check("下载越界：拒", c == 404)
        c, b = post("/api/export?page=0")
        check("编辑器：只导出这一页", b.get("ok") and b.get("slides") == 1 and "第1页" in b.get("file", ""), b)
        c, data = get("/api/download?f=" + q(b.get("file", "")))
        check("编辑器：下载这一页", c == 200 and data[:2] == b"PK")
        c, b = post("/api/export?page=abc")
        check("编辑器：页号不对就拒", not b.get("ok"))
        c, b = post("/api/export?page=999")
        check("编辑器：页号超出就拒", not b.get("ok") and "页号" in b.get("error", ""), b)
        d9 = P.load()
        d9["slides"].append({"id": "s98", "modules": [M("conclusion", 0, 0, 8, 1, "k")]})
        post("/api/deck", json.dumps(d9, ensure_ascii=False).encode("utf-8"))
        n9 = len(d9["slides"])
        c, b = post(f"/api/export?page={n9 - 1}")
        check("编辑器：只导出某一页时，报错里的页号是原来的页号", not b.get("ok") and f"第 {n9} 页" in b.get("error", ""), b)
        d9["slides"].pop()
        post("/api/deck", json.dumps(d9, ensure_ascii=False).encode("utf-8"))
        c, b = post("/api/deck", json.dumps(d9, ensure_ascii=False).encode("utf-8"), {"Origin": "https://evil.example"})
        check("编辑器：别的网页发来的请求一律拒（不能改课件）", c == 403, (c, b))
        dd = P.load()
        ci_ = next(i for i, s_ in enumerate(dd["slides"]) if any(m_["type"] == "chart" for m_ in s_["modules"]))
        cm_ = next(m_ for m_ in dd["slides"][ci_]["modules"] if m_["type"] == "chart")
        c, b = post("/api/chart/alts", json.dumps({"slide": ci_, "module": cm_["id"]}).encode("utf-8"))
        al = b.get("alts") or []
        from spec import CHARTS as _V6_CHARTS          # 图种数跟着 V6 走（0.10.0 起十三种），不写死
        check(f"换一种图：V6 自己试，{len(_V6_CHARTS)} 种里标出能换的和当前这种",
              b.get("ok") and len(al) == len(_V6_CHARTS) and sum(x["current"] for x in al) == 1
              and any(x["ok"] and not x["current"] for x in al), [(x["chart"], x["ok"]) for x in al])
        check("换一种图：画不了的说清原因", all(x["why"] for x in al if not x["ok"]))
        sp_ = next(x for x in al if x["ok"] and not x["current"])["spec"]
        dd2 = copy.deepcopy(dd)
        cm2 = next(m_ for m_ in dd2["slides"][ci_]["modules"] if m_["type"] == "chart")
        cm2.pop("file", None)
        cm2["spec"] = sp_
        _, rr = D.compose(dd2, ci_, proj, {})
        check("换一种图：换过去的这一页拼得出来", not rr["problems"], rr["problems"])
        c, b = post("/api/chart/alts", json.dumps({"slide": ci_, "module": "nope"}).encode("utf-8"))
        check("换一种图：不是图表卡就说清楚", not b.get("ok"))
        lib = json.loads(get("/api/library")[1])
        check("编辑器：素材里的 V3 关系图认得出来自 V3", any(f_.get("from") == "V3 法律关系图大师" for f_ in lib["figures"]), lib["figures"][:2])
        c, b = post("/api/frame/read?name=" + q("空白页.png"), _shot())
        check("编辑器：上传截图算出预留位置", b.get("ok") and b["frame"]["mode"] == "reserve" and b["frame"].get("backdrop"), b)
        c, data = get("/api/frame/backdrop")
        check("编辑器：截图存下来垫在页面下面", c == 200 and data[:4] == b"\x89PNG")
        d = P.load()
        d["frame"] = {k: v for k, v in b["frame"].items()}
        c, b2 = post("/api/deck", json.dumps(d, ensure_ascii=False).encode("utf-8"))
        check("编辑器：版面设定随课件保存", b2.get("ok") and P.load().get("frame", {}).get("mode") == "reserve", b2)
        d["frame"] = dict(d["frame"], top=500)
        c, b2 = post("/api/deck", json.dumps(d, ensure_ascii=False).encode("utf-8"))
        check("编辑器：留得太多的版面存不进去", not b2.get("ok"), b2)
        c, b = post("/api/frame/read?name=" + q("x.txt"), b"hello")
        check("编辑器：版面只收模板或截图", not b.get("ok"))
        c, b = post("/api/frame/read?name=" + q("t.pptx"), b"not a zip")
        check("编辑器：假 .pptx 就拒", not b.get("ok"))
        c, b = get("/api/brand")
        br = json.loads(b)
        check("/api/brand：读课件的名片", c == 200 and "某某" in br["brand"]["firm"]["name"] and br["local"])
        c, b = get("/api/brand/logo")
        check("/api/brand/logo", c == 200 and b[:4] == b"\x89PNG")
        bb = br["brand"]
        bb["firm"]["phone"] = "010-1111-2222"
        c, b = post("/api/brand", json.dumps(bb, ensure_ascii=False).encode("utf-8"))
        check("保存名片", b.get("ok") and json.load(open(os.path.join(proj, "brand", "brand.json"), encoding="utf-8"))["firm"]["phone"] == "010-1111-2222")
        c, b = post("/api/brand", json.dumps({"lawyers": [{"id": "x"}]}).encode("utf-8"))
        check("保存坏的名片：拒", not b.get("ok"))
        c, b = post("/api/brand/logo?name=" + q("logo.png"), _png(30, 10))
        check("上传 Logo", b.get("ok") and b.get("logo") == "logo.png")
        c, b = post("/api/brand/logo?name=" + q("logo.exe"), b"MZ")
        check("上传 Logo 拒：类型不对", not b.get("ok"))
        c, b = post("/api/brand/intro", json.dumps({"kind": "firm"}).encode("utf-8"))
        check("/api/brand/intro：律所介绍页", b.get("ok") and any(m["type"] == "profile" for m in b["modules"]))
        c, b = post("/api/ai", json.dumps({"kind": "module", "slide": "s10", "module": "m2", "prompt": "写法条"}).encode("utf-8"))
        check("/api/ai：记请求", b.get("ok") and b["request"]["status"] == "pending")
        rid = b["request"]["id"]
        c, b = get("/api/ai")
        check("/api/ai：列请求", any(r["id"] == rid for r in json.loads(b)["requests"]))
        c, b = post("/api/ai/cancel?id=" + rid)
        check("/api/ai/cancel", b.get("ok") and b["request"]["status"] == "cancelled")
        c, b = post("/api/ai", json.dumps({"kind": "module", "slide": "s99", "module": "m2"}).encode("utf-8"))
        check("/api/ai：没有这一页就拒", not b.get("ok"))
        c, b = post("/api/brand/template?name=" + q("fake.pptx") + "&style=" + q("接口律所风") + "&logos=0",
                    open(_fake_template(td), "rb").read())
        check("/api/brand/template：上传模板生成律所风格", b.get("ok") and b["name"] == "接口律所风" and b["logos"] == [], b)
        c, b = get("/api/state")
        st_ = json.loads(b)
        check("/api/state：风格里多了「自定义」，规则里有它的颜色（不出现律所名）", "自定义" in st_["styles"]
              and st_["rules"]["styles"]["自定义"]["accent"] == FIRM_ACC and "接口律所风" not in st_["rules"]["styles"], st_["styles"])
        c, b = get("/api/brand/template-report")
        check("/api/brand/template-report", c == 200 and "模板读取报告".encode("utf-8") in b)
        d_ = P.load()
        d_["style"] = "接口律所风"
        c, b = post("/api/deck", json.dumps(d_, ensure_ascii=False).encode("utf-8"))
        check("保存用律所风格的课件", b.get("ok"))
        c, b = post("/api/custom?name=" + q("官网截图.png"), _shot())
        check("/api/custom：截图 → 自定义风格，取到主色，顺带算出预留位置",
              b.get("ok") and b["colors"]["accent"] == "#702082" and b["name"] == "自定义" and (b.get("frame") or {}).get("mode") == "reserve", b)
        c, b = post("/api/custom/color", json.dumps({"color": "#1E40AF"}).encode("utf-8"))
        check("/api/custom/color：选主色 → 自定义风格", b.get("ok") and b["colors"]["accent"] == "#1E40AF", b)
        c, b = post("/api/custom/color", json.dumps({"color": "#888888"}).encode("utf-8"))
        check("/api/custom/color：灰色拒", not b.get("ok") and "灰" in b.get("error", ""), b)
        c, b = post("/api/custom?name=" + q("fake.pptx"), open(_fake_template(td), "rb").read())
        check("/api/custom：模板 → 自定义风格，名字一律「自定义」", b.get("ok") and b["name"] == "自定义" and b["logos"] == [], b)
        c, b = post("/api/custom?name=" + q("a.txt"), b"x")
        check("/api/custom：只收模板或图片", not b.get("ok"))
        c, b = post("/api/brand/template?name=" + q("x.pdf"), b"%PDF")
        check("/api/brand/template：不是 .pptx 就拒", not b.get("ok"))
        c, b = post("/api/brand/style/disable")
        check("停用律所风格：课件自动换回奇川风，照常能打开", b.get("ok") and P.load()["style"] == "奇川风")
        c, b = get("/api/version")
        v = json.loads(b)
        check("/api/version：课件、AI 队列、名片的修改时间", c == 200 and v["deck"] and v["ai"] and v["brand"])
        c, b = post("/api/deck", b"{}", {"Host": "evil.example.com"})
        check("外来 Host：拒（防 DNS 重绑定）", c == 403)
        # 两个窗口前后脚保存（开源前第三轮实跑：原来整份覆盖，先存的那处改动被悄悄冲掉）
        cur_ = json.loads(get("/api/version")[1])["deck"]
        dk = P.load()
        c, b = post("/api/deck?base=" + q(str(cur_)), json.dumps(dk, ensure_ascii=False).encode("utf-8"))
        check("保存带着最后看到的版本：版本对得上就存", b.get("ok") and isinstance(b.get("mtime"), str), b)
        c, b2 = post("/api/deck?base=" + q(str(cur_)), json.dumps(dk, ensure_ascii=False).encode("utf-8"))
        check("保存时磁盘已被别处改过：不覆盖，交回现在的课件和版本让编辑器合并", b2.get("conflict") and b2.get("deck")
              and b2.get("mtime") == b["mtime"], {k_: b2.get(k_) for k_ in ("ok", "conflict", "mtime")})
        c, b3 = post("/api/deck", json.dumps(dk, ensure_ascii=False).encode("utf-8"))
        check("不带版本的保存照旧（命令行、旧编辑器）", b3.get("ok"))
        check("课件版本用字符串（纳秒数超过 JS 能精确表示的整数）", isinstance(json.loads(get("/api/version")[1])["deck"], str))
        leaks = []
        for pth in ("/api/state", "/api/library", "/", "/api/version"):   # 读也要核：不然恶意网页借 DNS 重绑定能读走课件内容
            try:
                with urllib.request.urlopen(urllib.request.Request(base + pth, headers={"Host": "evil.example.com"}),
                                            timeout=60) as r:
                    leaks.append((pth, r.status))
            except urllib.error.HTTPError as e:
                if e.code != 403:
                    leaks.append((pth, e.code))
        check("外来 Host 读也拒：课件内容、素材清单、页面都不给（开源前实跑：原来只核了写）", not leaks, leaks)
        check("保存不留暂存文件", not any(f.endswith(".nlv-new") for f in os.listdir(proj)))
    finally:
        srv.shutdown()
        srv.server_close()


# ------------------------------------------------------------------ 成品账
def t_golden(update):
    d = demo()
    now, unstable = {}, []
    for style in D.STYLES:
        x = copy.deepcopy(d)
        x["style"] = style
        for i in range(len(x["slides"])):
            svg, _ = D.compose(x, i, DEMO, {})
            if svg != D.compose(x, i, DEMO, {})[0]:
                unstable.append(f"{style} 第 {i + 1} 页")
            now[f"{style}/{x['slides'][i]['id']}"] = fp(svg)
    check("拼页确定性（每页拼两次一样）", not unstable, unstable)
    if update or not os.path.exists(GOLDEN):
        os.makedirs(os.path.dirname(GOLDEN), exist_ok=True)
        with open(GOLDEN, "w", encoding="utf-8") as f:
            json.dump(now, f, ensure_ascii=False, indent=1, sort_keys=True)
        check("成品账已重写（--update）", True)
        return
    old = json.load(open(GOLDEN, encoding="utf-8"))
    diff = sorted(k for k in set(old) | set(now) if old.get(k) != now.get(k))
    check(f"成品账：示例课件 {len(now)} 页指纹不变", not diff,
          "变了：" + "、".join(diff[:6]) + "。若是有意改动，确认后 --update，并在 CHANGELOG 写明")


def main():
    update = "--update" in sys.argv
    td = tempfile.mkdtemp(prefix="nlv7-test-")
    try:
        t_validate()
        t_grid()
        t_templates(td)
        t_text()
        t_stat()
        t_figures()
        t_legal()
        t_brand(td)
        t_logo_export(td)
        t_ai(td)
        t_template_read(td)
        t_firm(td)
        t_export(td)
        t_paste(td)
        t_frame(td)
        t_pins_crop(td)
        t_import_real(td)
        t_adapt(td)
        t_cli(td)
        t_js_parity(td)
        t_prov(td)
        t_studio(td)
        t_detach(td)
        t_lib090(td)
        t_golden(update)
    finally:
        shutil.rmtree(td, ignore_errors=True)
        shutil.rmtree(_BRAND_HOME, ignore_errors=True)
    after = {m: tree_hash(os.path.join(SKILLS, m)) for m in SIBLINGS}
    for m in SIBLINGS:
        if BEFORE[m] is not None:
            check(f"只读兄弟模块：{m} 一个字节没变", after[m] == BEFORE[m])
    g = tempfile.mkdtemp(prefix="nlv7-guard-")
    try:
        with open(os.path.join(g, "a.py"), "w", encoding="utf-8") as f:
            f.write("x = 1\n")
        h1 = tree_hash(g)
        with open(os.path.join(g, "a.py"), "w", encoding="utf-8") as f:
            f.write("x = 2\n")
        check("兄弟模块守卫能失败（改一个字节指纹就变）", tree_hash(g) != h1)
    finally:
        shutil.rmtree(g, ignore_errors=True)
    check("示例课件目录没被测试改动", tree_hash(DEMO) == DEMO_HASH)
    bad = [(n, w) for n, ok, w in RESULTS if not ok]
    for n, w in bad:
        print(f"  不过　{n}　{str(w)[:400]}")
    print(f"{len(RESULTS) - len(bad)}/{len(RESULTS)} 项通过")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
