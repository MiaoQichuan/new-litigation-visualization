#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""类案编码：一批判决 → 一张编码表 → 一组图。

类案研究是律师做课件、写文章最常见的需求。图全部由同一张编码表统计出来，所以编码表立得住，
观点分歧、结果构成、地域排序、年份趋势、判赔分布、样本筛选都立得住。

分工（与 V2 / V4 / V5 一致）：

  脚本做　摄入（复用大事记表大师 V4 的读取层）；案号、法院、层级、省份、程序、年份这些
          机械字段由正则从原文抽，**不交给模型判断**；统计、出图。
  模型做　按编码方案逐份编码：每个字段取封闭表里的一个值，并附上原文引句作为依据。
  校验做　逐份对账（一份不能少，剔除要写理由）、封闭表、引句照录、金额与引句一致。

  python3 scripts/cases.py ingest <判决书或案例清单...> -w work/
  python3 scripts/cases.py check  work/coding.json work/codes.json -w work/
  python3 scripts/cases.py export work/coding.json work/codes.json -w work/
  python3 scripts/cases.py charts work/coding.json work/codes.json -w work/ [--style 歸藏风]

编码规程见 references/coding-guide.md。
"""
import argparse
import csv
import hashlib
import json
import math
import unicodedata
from decimal import Decimal, InvalidOperation
import os
import re
import sys
import zipfile

import kernel  # noqa: F401  （挂上 V1 路径）

HERE = os.path.dirname(os.path.abspath(__file__))
_V4_VENDOR = os.path.normpath(os.path.join(HERE, "..", "vendor", "v4", "scripts"))
_V4_SIBLING = os.path.normpath(os.path.join(HERE, "..", "..", "mqc-chronicle-master", "scripts"))


def _v4_ok(path):
    """V4 读取层够不够用：有 ingest(paths, workdir)。按能力挑，不只看文件在不在（同 kernel.py 的 V1）。"""
    fp = os.path.join(path, "ingest.py")
    return os.path.exists(fp) and "def ingest(paths, workdir" in open(fp, encoding="utf-8-sig").read()


V4 = next((p for p in (_V4_VENDOR, _V4_SIBLING) if _v4_ok(p)), _V4_SIBLING)


# =============================================================== 封闭表
BUILTIN = {
    "outcome_first": {"label": "一审裁判结果",
                      "options": ["全部支持", "部分支持", "驳回诉讼请求", "驳回起诉", "准予撤诉", "调解结案"]},
    "outcome_appeal": {"label": "二审裁判结果",
                       "options": ["驳回上诉、维持原判", "改判", "部分改判", "发回重审", "准予撤回上诉"]},
    "liability": {"label": "责任形式", "options": ["连带责任", "按份责任", "补充责任", "不承担责任"]},
}
EXCLUDE_REASONS = ["重复", "程序性文书", "与研究主题无关", "超出时间范围", "无法读取", "其他"]
MECHANICAL = ("case_no", "court", "level", "province", "procedure", "year", "doctype")
MECH_ZH = {"case_no": "案号", "court": "法院", "level": "法院层级", "province": "省份",
           "procedure": "程序", "year": "年份", "doctype": "文书类型"}

PROV = {"京": "北京", "津": "天津", "沪": "上海", "渝": "重庆", "冀": "河北", "晋": "山西", "蒙": "内蒙古",
        "辽": "辽宁", "吉": "吉林", "黑": "黑龙江", "苏": "江苏", "浙": "浙江", "皖": "安徽", "闽": "福建",
        "赣": "江西", "鲁": "山东", "豫": "河南", "鄂": "湖北", "湘": "湖南", "粤": "广东", "桂": "广西",
        "琼": "海南", "川": "四川", "黔": "贵州", "贵": "贵州", "云": "云南", "滇": "云南", "藏": "西藏",
        "陕": "陕西", "甘": "甘肃", "青": "青海", "宁": "宁夏", "新": "新疆", "兵": "新疆兵团"}
PROV_NAMES = sorted(set(PROV.values()), key=len, reverse=True)

RE_CASE = re.compile(r"[（(]\s*(\d{4})\s*[）)]\s*(最高法|[一-龥]{1,2})(\d{0,4})\s*"
                     r"([一-龥]{1,4})\s*(\d+)\s*号")
RE_COURT = re.compile(r"([一-龥]{2,30}?(?:人民法院|仲裁委员会))")
CN_DIGIT = {"〇": 0, "零": 0, "○": 0, "Ｏ": 0, "一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6,
            "七": 7, "八": 8, "九": 9, "十": 10}
RE_DATE_CN = re.compile(r"([二〇○零Ｏ一三四五六七八九]{4})年([一二三四五六七八九十]{1,2})月([一二三四五六七八九十]{1,3})日")
RE_DATE_AR = re.compile(r"(\d{4})年(\d{1,2})月(\d{1,2})日")


def _cn_num(s):
    if len(s) == 4:                                 # 年：逐字
        return int("".join(str(CN_DIGIT[c]) for c in s))
    if s == "十":
        return 10
    if s.startswith("十"):
        return 10 + CN_DIGIT[s[1]]
    if s.endswith("十"):
        return CN_DIGIT[s[0]] * 10
    if "十" in s:
        a, b = s.split("十")
        return CN_DIGIT[a] * 10 + CN_DIGIT[b]
    return CN_DIGIT[s]


def mechanical(text, name=""):
    """从原文抽机械字段。抽不出的留空，交给编码时用 override 补（须附引句）。"""
    head = text[:600]
    out = {k: None for k in MECHANICAL}
    m = RE_CASE.search(head) or RE_CASE.search(text)
    if m:
        out["case_no"] = re.sub(r"\s+", "", m.group(0)).replace("(", "（").replace(")", "）")
        code, kind = m.group(2), m.group(4)
        if code == "最高法":
            out["province"] = "最高法"
        elif code[0] in PROV:
            out["province"] = PROV[code[0]]
        out["procedure"] = ("一审" if kind.endswith("初") else "二审" if kind.endswith("终")
                            else "再审" if ("再" in kind or "申" in kind) else "执行" if "执" in kind else None)
        out["year"] = int(m.group(1))
    c = RE_COURT.search(head)
    if c:
        court = c.group(1)
        out["court"] = court
        out["level"] = ("最高" if court.startswith("最高人民法院") else "高级" if "高级人民法院" in court
                        else "中级" if ("中级人民法院" in court or re.search(r"(知识产权|金融|海事)法院", court))
                        else "仲裁" if "仲裁委员会" in court else "基层")
        if not out["province"]:
            for pn in PROV_NAMES:
                if court.startswith(pn):
                    out["province"] = pn
                    break
    head_tight = re.sub(r"\s+", "", head)            # 文书标题常排成「民 事 判 决 书」
    for kw in ("判决书", "裁定书", "调解书", "决定书", "裁决书"):
        if kw in head_tight:
            out["doctype"] = kw
            break
    tail = text[-400:]
    d = None
    for mm in RE_DATE_CN.finditer(tail):
        try:
            d = (_cn_num(mm.group(1)), _cn_num(mm.group(2)), _cn_num(mm.group(3)))
        except (KeyError, ValueError):
            pass
    if not d:
        for mm in RE_DATE_AR.finditer(tail):
            d = tuple(int(x) for x in mm.groups())
    if d:
        out["date"] = "%04d-%02d-%02d" % d
        out["year"] = d[0]                           # 裁判年度优先于立案年度
    return out


# =============================================================== 摄入
LIST_EXT = (".csv", ".xlsx")
LIST_HINT = {"case_no": ("案号",), "court": ("法院", "审理法院", "仲裁机构"), "date": ("裁判日期", "日期", "审结日期"),
             "title": ("标题", "案件名称", "名称"), "text": ("全文", "正文", "裁判要旨", "要旨", "摘要")}


def _read_list(path):
    if path.lower().endswith(".csv"):
        # 中文 Windows 上 Excel「另存为 CSV（逗号分隔）」默认是 GBK，不是 UTF-8：原来直接甩 UnicodeDecodeError。
        # 先按 UTF-8（带不带 BOM 都行）读，读不了再按 GB18030（GBK 的超集）读，并说一声。
        try:
            with open(path, encoding="utf-8-sig", newline="") as f:
                rows = list(csv.reader(f))
        except UnicodeDecodeError:
            try:
                with open(path, encoding="gb18030", newline="") as f:
                    rows = list(csv.reader(f))
            except UnicodeDecodeError:
                raise SystemExit(f"读不了案例清单 {path}：既不是 UTF-8 也不是 GBK。请在 Excel 里另存为「CSV UTF-8（逗号分隔）」")
            print(f"  提示：{os.path.basename(path)} 是 GBK 编码（Excel 默认另存的 CSV），已按 GBK 读入")
    else:
        try:
            import openpyxl
        except ImportError:
            raise SystemExit("读 .xlsx 案例清单需要 openpyxl（pip install openpyxl）；或先另存为 .csv")
        ws = openpyxl.load_workbook(path, read_only=True, data_only=True).active
        rows = [["" if v is None else str(v) for v in r] for r in ws.iter_rows(values_only=True)]
    # 记下每一行在源文件里的真实行号：空行跳过后再用 rows.index 找，两行内容相同时第二行会被记成第一行
    rows = [(i + 1, r) for i, r in enumerate(rows) if any(str(c).strip() for c in r)]
    if len(rows) < 2:
        raise SystemExit(f"{os.path.basename(path)}：案例清单至少要有表头和一行数据")
    return rows[0][1], rows[1:]


def ingest(paths, work):
    """文书走 V4 的读取层（格式分派、扫描件待读图、对账判据不变）；案例清单按行读。"""
    missing = [p for p in paths if not os.path.exists(p)]
    if missing:                                # 路径打错：原来会报成「探测出 0 页，文件可能损坏」，误导人去查文件
        raise SystemExit(f"找不到这些文件：{'、'.join(missing)}。请核对路径（中文与空格都可以，记得加引号）")
    os.makedirs(work, exist_ok=True)
    docs, sources, todo = [], {}, []
    files = [p for p in paths if not p.lower().endswith(LIST_EXT)]
    lists = [p for p in paths if p.lower().endswith(LIST_EXT)]
    if files:
        if not _v4_ok(V4):
            raise SystemExit("读判决书文件要用 V4 大事记表大师（mqc-chronicle-master）的读取层，没有找到。"
                             "请按插件完整安装；或先把案例导出成 .csv / .xlsx 清单再读")
        sys.path.insert(0, V4)
        import ingest as v4
        materials, srcs, td = v4.ingest(files, os.path.join(work, "raster"))
        for m in materials:
            did = "D%03d" % (len(docs) + 1)
            doc = {"did": did, "name": m["name"], "kind": "file", "medium": m.get("medium"), "pages": m.get("pages")}
            if m["mid"] in srcs:
                sources[did] = srcs[m["mid"]]
                doc.update(mechanical(srcs[m["mid"]], m["name"]))
            else:
                doc["needs_reading"] = True
            docs.append(doc)
        for t in td:
            todo.append(dict(t, did=next(d["did"] for d in docs if d["name"] == t["name"])))
    for lp in lists:
        head, rows = _read_list(lp)
        col = {}
        for key, hints in LIST_HINT.items():
            for i, h in enumerate(head):
                if any(x in str(h) for x in hints):
                    col.setdefault(key, i)
        for lineno, r in rows:
            did = "D%03d" % (len(docs) + 1)
            text = "\n".join(f"{head[i]}：{r[i]}" for i in range(min(len(head), len(r))) if str(r[i]).strip())
            sources[did] = text
            doc = {"did": did, "name": f"{os.path.basename(lp)} 第 {lineno} 行", "kind": "row", "row": lineno}
            if len(r) < len(head):
                doc["short_row"] = f"这一行只有 {len(r)} 列，表头有 {len(head)} 列"
            mech = mechanical(text)
            cell = lambda key: str(r[col[key]]).strip() if key in col and col[key] < len(r) else ""
            if cell("court"):                    # 短行缺这一列时不越界（外部测试 R2-058：缺中间列会 IndexError）
                mech["court"] = cell("court")
            doc.update(mech)
            docs.append(doc)
    # 重复分三级（外部评审 F10）：内容完全相同 → 重复；案号与文书类型都相同 → 疑似重复；
    # 只有案号相同（同一案件的判决与裁定、补正裁定）→ 同案材料，不算重复，由研究者决定统计单位。
    by_hash, by_key, by_case = {}, {}, {}
    for d in docs:
        txt = sources.get(d["did"])
        if txt:
            d["sha256"] = hashlib.sha256(_ws(txt).encode("utf-8")).hexdigest()[:16]
            if d["sha256"] in by_hash:
                d["duplicate_of"], d["dup_kind"] = by_hash[d["sha256"]], "内容相同"
                continue
            by_hash[d["sha256"]] = d["did"]
        cn = d.get("case_no")
        if not cn:
            continue
        key = (cn, d.get("doctype") or "")
        if key in by_key:
            d["duplicate_of"], d["dup_kind"] = by_key[key], "案号与文书类型相同"
        elif cn in by_case:
            d["same_case_as"] = by_case[cn]
        by_key.setdefault(key, d["did"])
        by_case.setdefault(cn, d["did"])
    json.dump(docs, open(os.path.join(work, "docs.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    json.dump(sources, open(os.path.join(work, "sources.json"), "w", encoding="utf-8"), ensure_ascii=False)
    return docs, todo


# =============================================================== 校验
_PUNCT = re.compile(r"[\s，。、；：！？“”‘’（）()《》【】\[\]〔〕,.;:!?\"'—\-·…]+")


def _ws(s):
    """只去空白、统一全角半角：换行、分栏、OCR 拆开的空格不算改字；标点一个都不能动。"""
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", str(s)))


def _norm(s):
    return _PUNCT.sub("", str(s))


def quote_ok(q, src, mode="strict"):
    """引句照录：按「……」分段，每段须**连续**出现在原文里，且段与段先后有序。只许删（用省略号跳过），不许改。

    strict（默认）：只忽略空白与全角半角之差，标点必须一致——0.8.0 连标点也一并去掉再比，
    把「，」改成「。」、把「不予支持」前的逗号挪个位置都能过，弱于「一字不改」的承诺（外部评审 F04）。
    ocr：扫描件转写稿可选，忽略标点，校验结果里会标出这一份用的是宽松比对。"""
    f = _ws if mode == "strict" else lambda x: _norm(_ws(x))
    segs = [s for s in re.split(r"…+|\.{3,}", q) if f(s)]
    if not segs:
        return False
    hay, pos = f(src), 0
    for s in segs:
        i = hay.find(f(s), pos)
        if i < 0:
            return False
        pos = i + len(f(s))
    return True


COURT_MARKERS = ("本院认为", "本院查明", "经审理查明", "本院经审理认定", "本院经审查认为", "判决如下", "裁定如下")


def quote_start(q, src, mode="strict"):
    """引句第一段在原文（规整后）里的位置；找不到返回 None。"""
    f = _ws if mode == "strict" else lambda x: _norm(_ws(x))
    segs = [s_ for s_ in re.split(r"…+|\.{3,}", q) if f(s_)]
    if not segs:
        return None
    i = f(src).find(f(segs[0]))
    return i if i >= 0 else None


def court_part_start(src, mode="strict"):
    """法院说理与裁判部分从哪里开始：第一个「本院认为 / 经审理查明 / 判决如下」等出现的位置。"""
    f = _ws if mode == "strict" else lambda x: _norm(_ws(x))
    hay = f(src)
    pos = [hay.find(f(mk)) for mk in COURT_MARKERS]
    pos = [p_ for p_ in pos if p_ >= 0]
    return min(pos) if pos else None


UNIT = {"元": 1, "万元": 10_000, "亿元": 100_000_000, "万": 10_000, "亿": 100_000_000}
RE_MONEY = re.compile(r"(\d[\d,，]*(?:\.\d+)?)\s*(亿元|万元|元|万|亿)")
# 千分位只认标准写法（每三位一组）：「12,34 元」不能被拼成 1234 元认下来（外部测试 R2-041）
RE_GROUPED = re.compile(r"\d{1,3}(?:[,，]\d{3})+(?:\.\d+)?|\d+(?:\.\d+)?")


RE_APPROX_BEFORE = re.compile(r"(约|大约|近|将近|逾|超过|不足|约为)\s*$")
RE_APPROX_AFTER = re.compile(r"^\s*(左右|余|多|以上|上下)")


def _dec(x):
    try:
        d = Decimal(str(x))
    except (InvalidOperation, ValueError):
        return None
    return d if d.is_finite() else None


def amount_in(q, value, unit, approx=False):
    """引句里有没有这个金额：按单位换到元后**精确**相等（Decimal，不设容差）。
    0.8.0 用 0.5% 的容差，120.5 万元能对上引句里的 120 万元（外部评审 F03）。
    引句里是约数（「约 120 万元」「120 万余元」）时，编码必须写 approx: true，否则不认。"""
    if isinstance(value, bool):
        return False
    v = _dec(value)
    if v is None or unit not in UNIT:
        return False
    want = v * UNIT[unit]
    for mt in RE_MONEY.finditer(q):
        if not RE_GROUPED.fullmatch(mt.group(1)):
            continue
        got = _dec(mt.group(1).replace(",", "").replace("，", ""))
        if got is None or got * UNIT[mt.group(2)] != want:
            continue
        rough = RE_APPROX_BEFORE.search(q[max(0, mt.start() - 3):mt.start()]) or \
            RE_APPROX_AFTER.search(q[mt.end():mt.end() + 3])
        if rough and not approx:
            continue
        return True
    return False


def amount_is_approx(q, value, unit):
    return amount_in(q, value, unit, approx=True) and not amount_in(q, value, unit, approx=False)


def load_scheme(path):
    sc = json.load(open(path, encoding="utf-8-sig"))
    errs = []
    for f in ("topic", "scope", "fields"):
        if not sc.get(f):
            errs.append(f"编码方案缺 {f}")
    keys = set()
    fields = []
    for i, f in enumerate(sc.get("fields", [])):
        f = dict(f)
        if "builtin" in f:
            b = BUILTIN.get(f["builtin"])
            if not b:
                errs.append(f"fields[{i}]：没有叫 {f['builtin']!r} 的内置封闭表（可选：{'、'.join(BUILTIN)}）")
                continue
            f.setdefault("label", b["label"])
            f["options"] = b["options"]
            f.setdefault("quote", True)
        if not f.get("key") or not f.get("label"):
            errs.append(f"fields[{i}] 要有 key 与 label")
            continue
        if f["key"] in keys or f["key"] in MECHANICAL:
            errs.append(f"字段 key {f['key']!r} 重复或与机械字段同名")
        keys.add(f["key"])
        if f.get("type") == "amount":
            if f.get("unit") not in UNIT:
                errs.append(f"金额字段 {f['key']} 要写 unit：元 / 万元 / 亿元")
            f["quote"] = True
        elif not f.get("options") or len(f["options"]) < 2:
            errs.append(f"字段 {f['key']} 要有封闭的 options（至少两项）；开放填写的字段统计不出图")
        fields.append(f)
    if errs:
        raise SystemExit("编码方案有问题：\n  " + "\n  ".join(errs))
    sc["fields"] = fields
    return sc


def _court_role(did, f, q, src, mode):
    """C7 证据角色（可选，字段写 "quote_from": "court" 时才查）：引句必须出自法院说理或裁判部分，
    不能是当事人的诉称、辩称（外部测试 R2-078：「原告主张」不能当成「法院认定」）。
    判据是位置：引句在原文里的位置不早于第一个「本院认为 / 经审理查明 / 判决如下」。
    这是机械判据，不是语义证明：法院在说理里转述当事人意见，位置也在后面，仍要人看。"""
    start = court_part_start(src, mode)
    if start is None:
        return [f"C7 证据角色：{did}「{f['label']}」要求引法院的话，可原文里找不到「本院认为」「经审理查明」"
                f"「判决如下」等法院部分的起点（是不是只有起诉状、清单摘要？）"]
    f_ = _ws if mode == "strict" else lambda x: _norm(_ws(x))
    segs = [s_ for s_ in re.split(r"…+|\.{3,}", q) if f_(s_)]
    hay, pos, after_ok = f_(src), start, True
    for sg in segs:                            # 从法院部分的起点往后，能按顺序找到每一段才算
        i = hay.find(f_(sg), pos)
        if i < 0:
            after_ok = False
            break
        pos = i + len(f_(sg))
    if not after_ok:
        return [f"C7 证据角色：{did}「{f['label']}」的引句在法院说理之前，是当事人的诉称或辩称，"
                f"不能当作法院认定；请改引「本院认为」之后的原文"]
    return []


def check(scheme, codes, docs, sources):
    """返回 (问题列表, 统计)。问题列表为空才能导出与出图。"""
    bad = []
    by = {}
    for k, c in enumerate(codes):
        did = c.get("did")
        if did in by:
            bad.append(f"C1 对账：{did} 编了两次")
        by[did] = c
    ids = [d["did"] for d in docs]
    for did in ids:
        if did not in by:
            bad.append(f"C1 对账：{did}（{next(d['name'] for d in docs if d['did'] == did)}）没有编码。一份都不能少；不要的请写剔除理由")
    for did in by:
        if did not in ids:
            bad.append(f"C1 对账：编码里有 {did}，材料里没有这一份")
    dmap = {d["did"]: d for d in docs}
    for did, c in by.items():
        if did not in dmap:
            continue
        d, src = dmap[did], sources.get(did, "")
        if d.get("needs_reading") and not src:
            if c.get("include", True) or c.get("reason") != "无法读取":
                bad.append(f"C1 对账：{did} 是扫描件，还没有转写稿；要么先读图转写，要么剔除并写「无法读取」")
            continue
        if not c.get("include", True):
            r = c.get("reason")
            if r not in EXCLUDE_REASONS:
                bad.append(f"C2 剔除：{did} 的理由 {r!r} 不在封闭表里（{'、'.join(EXCLUDE_REASONS)}）")
            elif r == "其他" and not c.get("note"):
                bad.append(f"C2 剔除：{did} 理由是「其他」，要在 note 里写明")
            elif r == "重复" and not d.get("duplicate_of") and not c.get("note"):
                bad.append(f"C2 剔除：{did} 标为重复，但案号没有与别的文书相同；请在 note 里写明与哪一份重复")
            continue
        dup = d.get("duplicate_of")
        if dup:
            # 同一组重复（都指向同一个主本）至多纳入一份：主本剔除、留一份副本可以；
            # 三份相同、剔一留二不行（外部测试 R2-057：0.9.0 只比了主本，两份副本都纳入也放过了）
            group = [x for x in docs if x["did"] == dup or x.get("duplicate_of") == dup]
            kept = [x["did"] for x in group if by.get(x["did"], {}).get("include", True) and x["did"] in by]
            if len(kept) > 1 and did == kept[1]:
                bad.append(f"C2 剔除：{'、'.join(kept)} 是同一组重复（{d.get('dup_kind', '案号相同')}），"
                           f"至多纳入一份；其余剔除并写「重复」")
        qmode = "ocr" if (d.get("needs_reading") or scheme.get("quote_mode") == "ocr") else "strict"
        vals = c.get("values", {})
        for k in vals:
            if k in MECHANICAL:
                bad.append(f"C6 机械字段：{did} 的 {MECH_ZH[k]} 由脚本从原文抽取，不能在 values 里改；确需更正请用 override 并附引句")
        for k, ov in (c.get("override") or {}).items():
            if k not in MECHANICAL:
                bad.append(f"C6 机械字段：override 只用于 {'、'.join(MECH_ZH.values())}")
            elif not isinstance(ov, dict) or not ov.get("q") or not quote_ok(ov["q"], src, "strict"):
                bad.append(f"C6 机械字段：{did} 更正 {MECH_ZH[k]} 要附一句原文引句，且引句须照录")
        for f in scheme["fields"]:
            key = f["key"]
            if key not in vals:
                bad.append(f"C3 缺项：{did} 没有编「{f['label']}」")
                continue
            ent = vals[key]
            v, q = (ent.get("v"), ent.get("q")) if isinstance(ent, dict) else (ent, None)
            if f.get("type") == "amount":
                if v is None:
                    if not ent.get("absent"):
                        bad.append(f"C3 缺项：{did}「{f['label']}」没有数；原文确实没有请写 absent: true")
                    continue
                if not isinstance(v, (int, float)) or isinstance(v, bool) or not math.isfinite(v):
                    bad.append(f"C3 金额：{did}「{f['label']}」要是有限的数字（不能是 NaN、无穷、真假值）")
                    continue
                if not q or not quote_ok(q, src, qmode):
                    bad.append(f"C4 引句：{did}「{f['label']}」的引句不在原文里（只许用省略号跳过，不许改字、改标点）")
                elif f.get("quote_from") == "court" and _court_role(did, f, q, src, qmode):
                    bad += _court_role(did, f, q, src, qmode)
                elif not amount_in(q, v, f["unit"], approx=bool(ent.get("approx"))):
                    if amount_is_approx(q, v, f["unit"]):
                        bad.append(f"C5 金额：{did}「{f['label']}」引句里是约数，编码要写 approx: true（约数会列进请核）")
                    else:
                        bad.append(f"C5 金额：{did}「{f['label']}」写的是 {v}{f['unit']}，引句里找不到这个数"
                                   f"（已按单位精确换算，不设容差）")
                continue
            if v not in f["options"]:
                bad.append(f"C3 封闭表：{did}「{f['label']}」的值 {v!r} 不在选项里（{'、'.join(f['options'])}）")
                continue
            need = f.get("quote", True) and v not in f.get("no_quote_for", [])
            if need and (not q or not quote_ok(q, src, qmode)):
                bad.append(f"C4 引句：{did}「{f['label']}」= {v}，要附原文引句，且引句须照录（只许用省略号跳过，标点也不能改）")
            elif need and f.get("quote_from") == "court":
                bad += _court_role(did, f, q, src, qmode)
    incl = [did for did, c in by.items() if c.get("include", True) and did in dmap]
    amount_keys = [f["key"] for f in scheme["fields"] if f.get("type") == "amount"]
    absent = [(did, k) for did in incl for k in amount_keys
              if isinstance(by[did].get("values", {}).get(k), dict) and by[did]["values"][k].get("absent")]
    stats = {"total": len(docs), "included": len(incl),
             "excluded": {r: sum(1 for c in by.values() if not c.get("include", True) and c.get("reason") == r)
                          for r in EXCLUDE_REASONS},
             # 请核：这些不是错，但脚本证明不了，要人看一眼
             "review": [f"{did}「{k}」声明原文没有金额（absent）：脚本证明不了「没有」，请核" for did, k in absent]
             + [f"{d['did']} 与 {d['same_case_as']} 是同一案号的不同文书（{d.get('doctype') or '类型未识别'}）："
                f"按案件统计时只计一次，按文书统计时各计一次，请在编码方案里写明统计单位"
                for d in docs if d.get("same_case_as") and d["did"] in incl and d["same_case_as"] in incl]
             + [f"{did}「{k}」是约数（approx），请核" for did in incl for k in amount_keys
                if isinstance(by[did].get("values", {}).get(k), dict) and by[did]["values"][k].get("approx")]
             + [f"{d['did']} 是扫描件转写稿，引句按宽松比对（忽略标点），请核" for d in docs
                if d.get("needs_reading") and d["did"] in incl]
             + ([f"编码方案设了 quote_mode: ocr：全部 {len(incl)} 份的引句都按宽松比对（忽略标点），"
                 f"「照录」只核到字、没核到标点，请核"] if scheme.get("quote_mode") == "ocr" else [])}
    return bad, stats


# =============================================================== 表
def table(scheme, codes, docs):
    by = {c["did"]: c for c in codes}
    rows = []
    for d in docs:
        c = by.get(d["did"], {})
        mech = {k: d.get(k) for k in MECHANICAL}
        for k, ov in (c.get("override") or {}).items():
            mech[k] = ov.get("v")
        row = {"did": d["did"], "文件": d["name"], **{MECH_ZH[k]: mech[k] for k in MECHANICAL},
               "裁判日期": d.get("date"), "纳入": "是" if c.get("include", True) else "否",
               "剔除理由": "" if c.get("include", True) else c.get("reason", "")}
        for f in scheme["fields"]:
            ent = (c.get("values") or {}).get(f["key"])
            v, q = (ent.get("v"), ent.get("q")) if isinstance(ent, dict) else (ent, None)
            row[f["label"]] = v
            row[f["label"] + "（依据）"] = q or ""
        rows.append(row)
    return rows


def export(scheme, rows, work):
    head = list(rows[0].keys())
    with open(os.path.join(work, "cases.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(head)
        for r in rows:
            w.writerow(["" if r[h] is None else r[h] for h in head])
    out = [os.path.join(work, "cases.csv")]
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        return out
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "编码表"
    ws.append(head)
    for r in rows:
        ws.append(["" if r[h] is None else r[h] for h in head])
    hdr = PatternFill("solid", fgColor="F3F4F6")
    for c in ws[1]:
        c.font = Font(bold=True, color="1F2933")
        c.fill = hdr
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, h in enumerate(head, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 40 if "依据" in h else max(8, min(26, len(h) * 2 + 4))
    ws.freeze_panes = "C2"
    info = wb.create_sheet("说明")
    for line in [["研究主题", scheme["topic"]], ["样本口径", scheme["scope"]],
                 ["机械字段", "案号、法院、层级、省份、程序、年份、文书类型由脚本从原文抽取"],
                 ["编码字段", "、".join(f["label"] for f in scheme["fields"])],
                 ["依据", "每个编码值附原文引句；引句只许用省略号跳过，不许改字"]]:
        info.append(line)
    info.column_dimensions["A"].width = 12
    info.column_dimensions["B"].width = 80
    p = os.path.join(work, "cases.xlsx")
    wb.save(p)
    return out + [p]


# =============================================================== 出图
def _pct_title(field, opt, share, n):
    return f"{field}：{share:g}% 为「{opt}」" if share < 100 else f"{field}：全部 {n} 份为「{opt}」"


def _shares(counts, n):
    """百分比到一位小数，最大余数法保证合计正好 100。"""
    raw = [c / n * 1000 for c in counts]
    base = [int(r) for r in raw]
    rest = 1000 - sum(base)
    for i in sorted(range(len(raw)), key=lambda i: (-(raw[i] - base[i]), i))[:rest]:
        base[i] += 1
    return [b / 10 for b in base]


def chart_specs(scheme, rows, stats):
    inc = [r for r in rows if r["纳入"] == "是"]
    n = len(inc)
    scope = f"{scheme['scope']} · 类案编码"
    src = f"类案编码表（{scheme['topic']}），逐份案号见 cases.xlsx"
    specs = {}
    # 样本筛选
    steps = [{"id": f"ex{i}", "label": f"剔除：{r}", "value": -c}
             for i, (r, c) in enumerate(stats["excluded"].items()) if c]
    if steps:
        specs["sample"] = {"chart": "waterfall", "title": f"{stats['total']} 份文书中纳入 {n} 份",
                           "scope": scheme["scope"], "unit": "份", "n": n, "source": src, "emphasis": "end",
                           "data": {"start": {"label": "收集文书", "value": stats["total"]},
                                    "steps": steps, "end": {"label": "纳入研究", "value": n}}}
    for f in scheme["fields"]:
        lab = f["label"]
        if f.get("type") == "amount":
            vals = [r[lab] for r in inc if isinstance(r[lab], (int, float))]
            if len(vals) >= 3:
                specs[f"{f['key']}_strip"] = {
                    "chart": "strip", "title": f"「{lab}」的分布（{len(vals)} 件）", "scope": scope,
                    "unit": f["unit"], "n": len(vals), "source": src,
                    "data": {"groups": [{"id": "all", "label": "纳入判决", "values": vals}]}}
            continue
        counts = [sum(1 for r in inc if r[lab] == o) for o in f["options"]]
        opts = [(o, c) for o, c in zip(f["options"], counts) if c]
        if not opts:
            continue
        top = max(opts, key=lambda t: t[1])
        if len(opts) <= 4:
            if n <= 100:
                data = {"measure": "count", "items": [{"id": f"o{i}", "label": o, "value": c} for i, (o, c) in enumerate(opts)]}
            else:
                sh = _shares([c for _, c in opts], n)
                data = {"measure": "percent", "items": [{"id": f"o{i}", "label": o, "value": s} for i, ((o, _), s) in enumerate(zip(opts, sh))]}
            share = round(top[1] / n * 100)
            specs[f["key"]] = {"chart": "waffle", "title": _pct_title(lab, top[0], share, n), "scope": scope,
                               "unit": "份判决", "n": n, "source": src,
                               "emphasis": f"o{[o for o, _ in opts].index(top[0])}", "data": data}
        else:
            specs[f["key"]] = {"chart": "hbar", "title": f"「{top[0]}」最多，{top[1]} 份", "scope": scope,
                               "unit": "份", "n": n, "source": src, "emphasis": "o0",
                               "data": {"items": [{"id": f"o{i}", "label": o, "value": c}
                                                  for i, (o, c) in enumerate(sorted(opts, key=lambda t: -t[1]))]}}
    # 机械字段：省份、层级、年份
    for key, chart in (("省份", "hbar"), ("法院层级", "hbar"), ("年份", "bar")):
        cnt = {}
        for r in inc:
            if r[key] is not None:
                cnt[r[key]] = cnt.get(r[key], 0) + 1
        if len(cnt) < 2:
            continue
        items = sorted(cnt.items(), key=(lambda t: t[0]) if key == "年份" else (lambda t: -t[1]))
        top = max(items, key=lambda t: t[1])
        ties = [k for k, v in items if v == top[1]]
        suffix = "年" if key == "年份" else ""
        if len(ties) > 1:                     # 并列第一不能只点一个名、只标一处重音
            title, emph = f"{key}：{len(ties)} 个并列最多，各 {top[1]} 份", None
        else:
            title, emph = f"{top[0]}{suffix}最多，{top[1]} 份", f"k{[k for k, _ in items].index(top[0])}"
        specs[{"省份": "province", "法院层级": "level", "年份": "year"}[key]] = {
            "chart": chart, "title": title,
            "scope": scope, "unit": "份", "n": n, "source": src, "order": "given",
            "emphasis": emph,
            "data": {"items": [{"id": f"k{i}", "label": str(k), "value": c} for i, (k, c) in enumerate(items)]}}
    return specs


# =============================================================== 命令行
def _load_work(work):
    docs = json.load(open(os.path.join(work, "docs.json"), encoding="utf-8-sig"))
    sources = json.load(open(os.path.join(work, "sources.json"), encoding="utf-8-sig"))
    tdir = os.path.join(work, "transcripts")
    if os.path.isdir(tdir):
        for f in os.listdir(tdir):
            if f.endswith(".txt"):
                did = f[:-4]
                sources[did] = open(os.path.join(tdir, f), encoding="utf-8-sig").read()
                for d in docs:
                    if d["did"] == did and d.get("needs_reading"):
                        d.update(mechanical(sources[did]))
                        d["needs_reading"] = False
                        d["medium"] = "image"
    return docs, sources


def main(argv=None):
    ap = argparse.ArgumentParser(description="类案编码：一批判决 → 编码表 → 图")
    sub = ap.add_subparsers(dest="cmd", required=True)
    a1 = sub.add_parser("ingest"); a1.add_argument("files", nargs="+"); a1.add_argument("-w", "--work", required=True)
    for name in ("check", "export", "charts"):
        s = sub.add_parser(name)
        s.add_argument("scheme"); s.add_argument("codes"); s.add_argument("-w", "--work", required=True)
        if name == "charts":
            s.add_argument("--style"); s.add_argument("--formats", default="svg,pptx")
    a = ap.parse_args(argv)
    if a.cmd == "ingest":
        docs, todo = ingest(a.files, a.work)
        print(f"读入 {len(docs)} 份（文书 {sum(1 for d in docs if d['kind'] == 'file')}，清单行 {sum(1 for d in docs if d['kind'] == 'row')}）")
        miss = [d for d in docs if not d.get("needs_reading") and (not d.get("case_no") or not d.get("court"))]
        dup = [d for d in docs if d.get("duplicate_of")]
        if todo:
            print(f"扫描件 {len(todo)} 份待读图（约 {sum(t['est_tokens'] for t in todo)} token）："
                  + "、".join(f"{t['did']} {t['name']}" for t in todo))
            print("  读完把转写稿存为 work/transcripts/<did>.txt；不读就在编码里剔除并写「无法读取」")
        for d in miss:
            print(f"  {d['did']} {d['name']}：没抽到{'案号' if not d.get('case_no') else '法院'}，编码时用 override 补（附引句）")
        for d in dup:
            print(f"  {d['did']} 与 {d['duplicate_of']} {d.get('dup_kind', '案号相同')}，至多纳入一份")
        for d in docs:
            if d.get("same_case_as"):
                print(f"  {d['did']} 与 {d['same_case_as']} 同一案号、文书不同：不是重复，统计单位由编码方案定")
            if d.get("short_row"):
                print(f"  {d['did']} {d['name']}：{d['short_row']}，请核对清单")
        print(f"机械字段已写入 {os.path.join(a.work, 'docs.json')}；下一步按 references/coding-guide.md 写编码方案与编码")
        return 0
    scheme = load_scheme(a.scheme)
    codes = json.load(open(a.codes, encoding="utf-8-sig"))
    docs, sources = _load_work(a.work)
    bad, stats = check(scheme, codes, docs, sources)
    if bad:
        print(f"编码没过（{len(bad)} 处）：", file=sys.stderr)
        for b in bad:
            print("  " + b, file=sys.stderr)
        return 2
    print(f"编码通过：共 {stats['total']} 份，纳入 {stats['included']} 份，剔除 "
          + "、".join(f"{r} {c}" for r, c in stats["excluded"].items() if c) if any(stats["excluded"].values())
          else f"编码通过：共 {stats['total']} 份，全部纳入")
    for r in stats.get("review", []):
        print("请核　" + r)
    if a.cmd == "check":
        return 0
    rows = table(scheme, codes, docs)
    files = export(scheme, rows, a.work)
    if a.cmd == "export":
        print("母表：" + "、".join(files))
        return 0
    import make
    specs = chart_specs(scheme, rows, stats)
    outdir = os.path.join(a.work, "charts")
    os.makedirs(outdir, exist_ok=True)
    made, failed = [], []
    for k, sp in specs.items():
        p = os.path.join(outdir, f"{k}.json")
        json.dump(dict(sp, schema_version=1), open(p, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        args = [p, "-o", os.path.join(outdir, k), "--formats", a.formats] + (["--style", a.style] if a.style else [])
        import io
        import contextlib
        buf, err = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(buf), contextlib.redirect_stderr(err):
            rc = make.main(args)
        (made if rc == 0 else failed).append((k, err.getvalue().strip()))
    print("母表：" + "、".join(files))
    print(f"出图 {len(made)} 张：" + "、".join(k for k, _ in made) + f"（{outdir}）")
    for k, e in failed:
        print(f"  {k} 未出图：{e.splitlines()[0] if e else ''}")
    print("图的标题由脚本从编码表算出，是事实陈述；要改成更有观点的结论，改 charts/<名>.json 的 title 再出图，"
          "标题里的数仍须由数据推得出")
    return 0


if __name__ == "__main__":
    # Windows 中文系统的控制台默认 GBK：不固定成 UTF-8，外部 AI（按 UTF-8 读）看到的是乱码，
    # 遇到 GBK 里没有的字（生僻字人名）还会在图已画好之后崩掉。写法与其余模块一致。
    for _stream in (sys.stdout, sys.stderr):
        try:
            _stream.reconfigure(encoding="utf-8")
        except Exception:                               # noqa: BLE001
            pass
    try:                                   # 常见输入错误（路径、JSON、编码）说人话，不甩堆栈
        __import__("_cli_errors").install()
    except ImportError:
        pass
    sys.exit(main())
